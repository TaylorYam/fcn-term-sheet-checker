"""核對結果檔：一批說明書核對後唯一的 Excel 產出（Issue #72）。

- 「回填後」：一律照匯入格式（設定檔 [import_format]，比照公司原始待建檔表）在新的活頁簿重新產生（Issue #203）：
  第 1 列標題與留下的列裡最晚的交易日、第 2 列固定欄位的表頭、第 3 列起是整份通過（或人工放行）且已回填的列，
  順序照參考條件表；可以直接匯入資料庫。每格依欄名取參考條件表同名欄（同一標準欄位的其他欄名、fallbacks 的欄也算）
  回填後的值與顯示格式，`checked_column` 一律打勾，表上沒有的欄留空白；欄寬取同名欄。參考條件表的其他欄、
  其他工作表、公式、格式化條件、資料驗證、篩選、合併儲存格都不帶入，所以什麼樣的參考條件表都產生得出來。
- 「錯誤清單」：每份沒通過的說明書一列，依 PDF 輸入順序：TDCC Code、PDF 檔名、錯訊（多條以換行分隔）。

檔名 `<參考條件表檔名>_核對結果_<YYYYMMDD-HHMMSS>.xlsx`，寫到指定資料夾，不覆蓋既有檔案。
"""

from __future__ import annotations

import datetime as dt
import io
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from pathlib import Path

from openpyxl.styles import Alignment
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.dimensions import ColumnDimension
from openpyxl.worksheet.worksheet import Worksheet

from .config import ReferenceFormat
from .identification import Identification, name_code
from .ingestion import write_new
from .orders.reference import SheetLayout, header_names, locate, normalize_cell

FILLED_SHEET = "回填後"
ERROR_SHEET = "錯誤清單"
ERROR_COLUMNS = (("TDCC Code", 16), ("PDF 檔名", 40), ("錯訊", 100))  # 表頭與欄寬


@dataclass(frozen=True)
class ErrorRow:
    """錯誤清單的一列。"""

    tdcc_code: str | None
    pdf_name: str
    message: str

    @classmethod
    def of(cls, identification: Identification, messages: Sequence[str]) -> ErrorRow:
        """一份 PDF 的錯誤清單列（TDCC Code 見 `tdcc_code`）。錯訊多條以換行分隔。"""
        return cls(tdcc_code(identification), identification.pdf.name, "\n".join(messages))


def tdcc_code(identification: Identification) -> str | None:
    """錯誤清單（與回報包）的 TDCC Code：辨識結果的商品代號（說明書封面；投資人須知為檔名）；
    取不到時用檔名前 12 碼（12 位數字才算），否則為 None。"""
    return identification.product_code or name_code(identification.pdf)


def output_path(out_dir: Path, reference_sheet: Path, now: dt.datetime) -> Path:
    return out_dir / f"{reference_sheet.stem}_核對結果_{now:%Y%m%d-%H%M%S}.xlsx"


def save(wb: Workbook, rfmt: ReferenceFormat, keep: Iterable[int], errors: Sequence[ErrorRow], out: Path) -> Path:
    """寫出核對結果檔並回傳路徑；檔案已存在（不覆蓋）或無法寫入時丟出 IngestionError。

    `wb` 是已回填的參考條件表（讀值開啟）；`keep` 是要留在「回填後」的資料列號。
    """
    write_new(out, _build(wb, rfmt, keep, errors), "核對結果檔")
    return out


def _build(wb: Workbook, rfmt: ReferenceFormat, keep: Iterable[int], errors: Sequence[ErrorRow]) -> bytes:
    layout = locate(wb, rfmt)
    out = Workbook()
    ws = out.active
    ws.title = FILLED_SHEET
    _filled_sheet(ws, wb[layout.sheet], layout, rfmt, sorted(set(keep)))
    _error_sheet(out.create_sheet(ERROR_SHEET), errors)
    buf = io.BytesIO()
    out.save(buf)
    return buf.getvalue()


def _filled_sheet(ws: Worksheet, src: Worksheet, layout: SheetLayout, rfmt: ReferenceFormat, rows: list[int]) -> None:
    """照匯入格式寫「回填後」：每格取參考條件表 `rows` 各列對應欄的值與顯示格式。"""
    imp = rfmt.import_format
    by_header: dict[str, int] = {}  # 參考條件表欄名 → 欄號（重複時取第一個）
    for col in range(1, src.max_column + 1):
        header = normalize_cell(src.cell(layout.header_row, col).value)
        if header is not None:
            by_header.setdefault(str(header), col)
    names = header_names(rfmt)
    sources = [None if h == imp.checked_column else _source(h, by_header, rfmt, names) for h in imp.columns]

    trade = next((by_header[h] for h in names["trade_date"] if h in by_header), None)
    dates = [src.cell(r, trade) for r in rows if trade and isinstance(src.cell(r, trade).value, dt.datetime)]
    ws.cell(1, 1, imp.title)
    if dates:
        latest = max(dates, key=lambda c: c.value)
        ws.cell(1, 2, latest.value).number_format = latest.number_format
    for k, header in enumerate(imp.columns, start=1):
        ws.cell(imp.header_row, k, header)
    for r_out, r_src in enumerate(rows, start=imp.header_row + 1):
        for k, (header, col) in enumerate(zip(imp.columns, sources, strict=True), start=1):
            if header == imp.checked_column:
                ws.cell(r_out, k, rfmt.checked_value)
            elif col is not None:
                cell = src.cell(r_src, col)
                ws.cell(r_out, k, cell.value).number_format = cell.number_format
    dims = _dim_columns(src, src.max_column)
    for k, col in enumerate(sources, start=1):
        if col is not None and col in dims and dims[col].width:
            ws.column_dimensions[get_column_letter(k)].width = dims[col].width


def _source(header: str, by_header: dict[str, int], rfmt: ReferenceFormat, names: dict[str, list[str]]) -> int | None:
    """匯入格式某欄取參考條件表的哪一欄：同名欄、同一標準欄位的其他欄名（例：計價幣別 ← 承作幣別）、fallbacks；都沒有時 None。"""
    std = rfmt.columns.get(header)
    fallback = rfmt.import_format.fallbacks.get(header)
    for name in (header, *(names[std] if std else ()), *((fallback,) if fallback else ())):
        if name in by_header:
            return by_header[name]
    return None


def _dim_columns(ws: Worksheet, last: int) -> dict[int, ColumnDimension]:
    """欄號 → 欄設定（欄寬、隱藏等）；一筆設定涵蓋多欄時展開，只到第 `last` 欄。"""
    out = {}
    for dim in ws.column_dimensions.values():
        lo = dim.min or column_index_from_string(dim.index)
        for c in range(lo, min(dim.max or lo, last) + 1):
            out[c] = dim
    return out


def _error_sheet(ws: Worksheet, errors: Sequence[ErrorRow]) -> None:
    ws.append([header for header, _ in ERROR_COLUMNS])
    for e in errors:
        ws.append((e.tdcc_code, e.pdf_name, e.message))
    for k, (_, width) in enumerate(ERROR_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(k)].width = width
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
