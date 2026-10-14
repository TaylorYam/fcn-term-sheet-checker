"""回報包：一批核對結果中所有未通過的商品，交給維護者重現和判斷的整包資料（Issue #204）。

PANEL「匯出回報包…」呼叫 `export_bundle`，不需要先儲存核對結果，也不寫核對紀錄。資料夾結構是平的，
維護者拿到後可以直接用 `fcn-batch 參考條件表_未通過.xlsx PDF/*.pdf` 原樣重跑：

- `回報總表.xlsx`：「回報」一條錯訊一列（上手、TDCC Code、PDF 檔名、錯訊、我的判斷、備註），「我的判斷」是
  誤報／真錯／不確定的下拉選單，由作業人員填；「摘要」是這批各類別的份數。
- `參考條件表_未通過.xlsx`：只有這些商品的列，只留匯入格式的欄、核對欄與發行機構（其他欄拿掉），版面位置照原表。
- `核對紀錄_未通過.json`：核對紀錄中這些商品的部分，格式相同。
- `PDF/`：這些商品的說明書與投資人須知（同商品兩份都附，就算只有一份未通過）。檔名是辨識的依據，不改名；同名的不同檔案
  （例：多份對到同一列）另放 `PDF/同名<n>/`，重跑時一起帶入（`PDF/*.pdf PDF/同名*/*.pdf`），摘要會註明。

「未通過」以錯誤清單為準（沒通過、也沒人工放行）；已知不支援的（未支援上手、不支援樣態）不列入，只算進摘要（Issue #212）。
回報包含真實交易資料，只能放 Private Repo，不能放進本 Repo。
"""

from __future__ import annotations

import datetime as dt
import io
import os
import shutil
import tempfile
import time
from collections.abc import Sequence
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.datavalidation import DataValidation

from . import reporting
from .batch import BatchItem, BatchOutcome, Category
from .config import ReferenceFormat
from .ingestion import IngestionError
from .orders.reference import header_columns, locate
from .result_file import tdcc_code
from .schema import overall_status

REPORT_FILE = "回報總表.xlsx"
REFERENCE_FILE = "參考條件表_未通過.xlsx"
RECORD_FILE = "核對紀錄_未通過.json"
PDF_DIR = "PDF"
REPORT_SHEET = "回報"
SUMMARY_SHEET = "摘要"
REPORT_COLUMNS = (("上手", 10), ("TDCC Code", 16), ("PDF 檔名", 32), ("錯訊", 90), ("我的判斷", 12), ("備註", 40))
JUDGEMENT_COLUMN = "我的判斷"
JUDGEMENTS = ("誤報", "真錯", "不確定")
UNSUPPORTED_LABEL = "未支援上手（已知不支援，未列入回報包）"
VARIANT_LABEL = "不支援樣態（已知不支援，未列入回報包）"
FAILED = (Category.MISMATCH, Category.REVIEW, Category.ERROR)  # 摘要的「未通過」（不支援樣態另計）
SAME_NAME_DIR = "同名"
RENAME_ATTEMPTS = 5  # 改名成回報包時，防毒或索引服務可能還抓著剛寫好的檔案：稍等重試


def _known_unsupported(item: BatchItem) -> bool:
    """已知不支援、不是規則誤報：未支援上手或不支援樣態。"""
    return item.unsupported or item.identification.unsupported_variant


def reportable(outcome: BatchOutcome) -> tuple[BatchItem, ...]:
    """要列入回報包的 PDF（依輸入順序）：錯誤清單上的，扣掉已知不支援的。"""
    return tuple(i for i in outcome.items if not i.fillable and not _known_unsupported(i))


def bundle_path(out_dir: Path, reference_sheet: Path, now: dt.datetime) -> Path:
    return out_dir / f"{reference_sheet.stem}_回報包_{now:%Y%m%d-%H%M%S}"


def export_bundle(outcome: BatchOutcome, out_dir: Path, *, root: Path, now: dt.datetime | None = None) -> Path:
    """確認來源都與核對時相同後，在 out_dir 建立回報包資料夾並回傳路徑。

    沒有未通過的商品（nothing_to_report）、來源已變更（source_changed）、資料夾已存在（output_exists，不覆蓋）
    或無法寫入（output_unwritable）時丟出 IngestionError，不留下半成品。root 是根目錄（核對紀錄的程式 commit 取自它）。
    """
    failing = reportable(outcome)
    rfmt = outcome.reference_format
    if rfmt is None or not failing:  # 整批錯誤時沒有核對項目，也就沒有未通過的商品
        raise IngestionError("nothing_to_report", "這批沒有未通過的商品，不需要匯出回報包")
    if outcome.snapshot is None or not outcome.snapshot.still_valid():
        raise IngestionError(
            "source_changed", "參考條件表、說明書或設定檔在核對後已變更或無法讀取，請重新核對後再匯出回報包"
        )
    now = now or dt.datetime.now()
    members = _members(outcome, failing)
    pdfs = _pdf_copies(members)
    files = {
        REPORT_FILE: _report_book(outcome, failing, pdfs, now),
        REFERENCE_FILE: _reference_extract(outcome.reference_sheet, rfmt, members),
        RECORD_FILE: reporting.record_bytes(_record(outcome, members), Path(root), now),
    }
    target = bundle_path(Path(out_dir), outcome.reference_sheet, now)
    _write_folder(target, files, pdfs)
    return target


def _members(outcome: BatchOutcome, failing: Sequence[BatchItem]) -> tuple[BatchItem, ...]:
    """未通過的 PDF 加上同商品的另一份，依輸入順序。"""
    wanted = {id(i) for i in failing} | {id(i.partner) for i in failing if i.partner is not None}
    return tuple(i for i in outcome.items if id(i) in wanted)


def _pdf_copies(members: Sequence[BatchItem]) -> list[tuple[Path, Path]]:
    """(回報包內的相對路徑, 來源)。檔名不改（辨識要用）；同名的不同檔案放到 PDF/同名<n>/；讀不到的檔案不附。"""
    out: list[tuple[Path, Path]] = []
    used: set[str] = set()
    sources: set[Path] = set()
    for item in members:
        src = item.term_sheet
        if src in sources or not src.is_file():
            continue
        sources.add(src)
        rel, n = Path(PDF_DIR) / src.name, 1
        while str(rel) in used:
            n += 1
            rel = Path(PDF_DIR) / f"{SAME_NAME_DIR}{n}" / src.name
        used.add(str(rel))
        out.append((rel, src))
    return out


def _write_folder(target: Path, files: dict[str, bytes], pdfs: Sequence[tuple[Path, Path]]) -> None:
    """先寫進同資料夾的暫存資料夾，全部寫完才改名成回報包；失敗時移除暫存資料夾，錯訊一律寫回報包的路徑。"""
    try:
        target.parent.mkdir(parents=True, exist_ok=True)
    except OSError as e:
        raise IngestionError("output_unwritable", f"無法建立回報包的資料夾 {target.parent}：{e}") from e
    if target.exists():
        raise _exists(target)
    try:
        tmp = Path(tempfile.mkdtemp(prefix=f"{target.name}_暫存_", dir=target.parent))
    except OSError as e:
        raise IngestionError("output_unwritable", f"無法寫入回報包 {target}：{e}") from e
    try:
        for name, data in files.items():
            (tmp / name).write_bytes(data)
        for rel, src in pdfs:
            (tmp / rel).parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(src, tmp / rel)
        _rename(tmp, target)
    except FileExistsError as e:
        raise _exists(target) from e
    except OSError as e:
        raise IngestionError("output_unwritable", f"無法寫入回報包 {target}：{e}") from e
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def _exists(target: Path) -> IngestionError:
    return IngestionError("output_exists", f"回報包 {target.name} 已經存在，不覆蓋")


def _rename(tmp: Path, target: Path) -> None:
    for attempt in range(RENAME_ATTEMPTS):
        try:
            os.rename(tmp, target)
            return
        except PermissionError:
            if attempt == RENAME_ATTEMPTS - 1:
                raise
            time.sleep(0.2)


def _xlsx(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _report_book(
    outcome: BatchOutcome, failing: Sequence[BatchItem], pdfs: Sequence[tuple[Path, Path]], now: dt.datetime
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = REPORT_SHEET
    ws.append([header for header, _ in REPORT_COLUMNS])
    for item in failing:
        code = tdcc_code(item.identification)
        for message in item.problem_messages or (item.status_label,):
            ws.append((item.issuer, code, item.term_sheet.name, message, None, None))
    judgement = get_column_letter([h for h, _ in REPORT_COLUMNS].index(JUDGEMENT_COLUMN) + 1)
    dv = DataValidation(type="list", formula1='"' + ",".join(JUDGEMENTS) + '"', allow_blank=True)
    dv.error, dv.errorTitle, dv.showErrorMessage = "請從清單選擇：" + "、".join(JUDGEMENTS), JUDGEMENT_COLUMN, True
    ws.add_data_validation(dv)
    dv.add(f"{judgement}2:{judgement}{ws.max_row}")
    for k, (_, width) in enumerate(REPORT_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(k)].width = width
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    _summary_sheet(wb.create_sheet(SUMMARY_SHEET), outcome, failing, pdfs, now)
    return _xlsx(wb)


def _summary_sheet(
    ws, outcome: BatchOutcome, failing: Sequence[BatchItem], pdfs: Sequence[tuple[Path, Path]], now: dt.datetime
) -> None:
    """各類別的份數（以 PDF 計）；不支援樣態另列一列、不算進需人工覆核；「未通過」= 不一致＋需人工覆核＋執行錯誤。"""
    items, variants = outcome.items, [i for i in outcome.items if i.identification.unsupported_variant]
    counted = [i for i in items if not i.identification.unsupported_variant]
    ws.append(("參考條件表", outcome.reference_sheet.name))
    ws.append(("匯出時間", f"{now:%Y-%m-%d %H:%M:%S}"))
    ws.append(("PDF 總數", len(items)))
    for c in Category:
        ws.append((UNSUPPORTED_LABEL if c == Category.UNSUPPORTED else c.value, sum(i.category == c for i in counted)))
        if c == Category.UNSUPPORTED:
            ws.append((VARIANT_LABEL, len(variants)))
    ws.append(("未通過", sum(i.category in FAILED for i in counted)))
    ws.append(("列入回報包", len(failing)))
    if any(rel.parent.name.startswith(SAME_NAME_DIR) for rel, _ in pdfs):
        ws.append(("注意", f"有同名的不同 PDF，另放在 {PDF_DIR}/{SAME_NAME_DIR}<n>/，重跑時一起帶入"))
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 40


def _reference_extract(reference_sheet: Path, rfmt: ReferenceFormat, members: Sequence[BatchItem]) -> bytes:
    """參考條件表只留這些商品的列與認得的欄；只有第一張工作表，表頭以上各列與資料列的位置照原表版面，值與數字格式照原表。"""
    try:
        wb = openpyxl.load_workbook(reference_sheet, data_only=True)
    except Exception as e:
        raise IngestionError("reference_unreadable", f"參考條件表無法開啟（可能已損毀或不是 Excel）：{e}") from e
    layout = locate(wb, rfmt)
    src = wb[layout.sheet]
    known = {*rfmt.import_format.columns, *rfmt.columns, rfmt.issuer_column}
    cols = sorted(found[0] for header, found in header_columns(src, layout).items() if header in known)
    rows = sorted({i.reference_row for i in members if i.reference_row is not None})
    wb = Workbook()
    ws = wb.active
    ws.title = layout.sheet
    for r_out, r_src in enumerate([*range(1, layout.first_data_row), *rows], start=1):
        for c_out, c_src in enumerate(cols, start=1):
            cell = src.cell(r_src, c_src)
            ws.cell(r_out, c_out, cell.value).number_format = cell.number_format
    for c_out, c_src in enumerate(cols, start=1):
        dim = src.column_dimensions.get(get_column_letter(c_src))
        if dim is not None and dim.width:
            ws.column_dimensions[get_column_letter(c_out)].width = dim.width
    return _xlsx(wb)


def _record(outcome: BatchOutcome, members: Sequence[BatchItem]) -> reporting.BatchRecord:
    """核對紀錄中這些 PDF 的部分；回報包不回填（filled 一律 False），也沒有核對結果檔。"""
    return reporting.BatchRecord(
        status=overall_status([i.status for i in members]),
        result_file=None,
        metadata=outcome.metadata,
        errors=tuple(outcome.errors),
        items=tuple(
            reporting.RecordItem(i.identification, i.report, False, i.released, i.not_filled_reason) for i in members
        ),
    )
