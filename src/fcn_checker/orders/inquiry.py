"""上手詢價表 adapter：依格式設定把 Excel 轉成標準欄位。未知欄名回報，不自行推測。"""

from __future__ import annotations

import datetime as dt
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter

from ..config import OrderFormat
from ..ingestion import IngestionError
from ..schema import OrderValue


@dataclass
class OrderRecord:
    product_code: OrderValue
    fields: dict[str, OrderValue]
    unknown_columns: list[OrderValue] = field(default_factory=list)  # value = 欄名
    missing_columns: list[str] = field(default_factory=list)  # 設定有、檔案沒有的 Excel 欄名
    duplicate_columns: list[OrderValue] = field(default_factory=list)  # 重複出現的欄名（value = 欄名）
    source: str = "詢價表"  # 條件來源名稱，用在核對訊息


def normalize_cell(v: Any) -> Any:
    """Excel 值標準化：數字轉 Decimal 並四捨五入到 9 位清除浮點尾數；日期轉 date；空字串轉 None。"""
    if v is None:
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, int):
        return Decimal(v)
    if isinstance(v, float):
        return Decimal(repr(round(v, 9)))
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, str):
        s = v.strip()
        return s or None
    return v


def load_inquiry(path: Path, fmt: OrderFormat) -> OrderRecord:
    if not path.is_file():
        raise IngestionError("order_not_found", f"找不到詢價表檔案：{path.name}")
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    except Exception as e:
        raise IngestionError("order_unreadable", f"詢價表無法開啟（可能已損毀或不是 Excel）：{e}") from e
    if fmt.sheet not in wb.sheetnames:
        raise IngestionError("order_sheet_missing", f"詢價表缺少工作表「{fmt.sheet}」")
    ws = wb[fmt.sheet]

    def ref(cell: str) -> str:
        return f"{fmt.sheet}!{cell}"

    raw_code = ws[fmt.product_code_cell].value
    if isinstance(raw_code, float) and raw_code.is_integer():
        raw_code = int(raw_code)
    code = str(raw_code).strip() if raw_code is not None else None
    product_code = OrderValue(code or None, ref(fmt.product_code_cell))

    fields: dict[str, OrderValue] = {}
    unknown: list[OrderValue] = []
    duplicate: list[OrderValue] = []
    seen: set[str] = set()
    for col in range(1, ws.max_column + 1):
        header = normalize_cell(ws.cell(row=fmt.header_row, column=col).value)
        if header is None:
            continue
        header = str(header)
        letter = get_column_letter(col)
        if header in seen:
            duplicate.append(OrderValue(header, ref(f"{letter}{fmt.header_row}")))
        elif header in fmt.columns:
            std = fmt.columns[header]
            value = normalize_cell(ws.cell(row=fmt.data_row, column=col).value)
            fields[std] = OrderValue(value, ref(f"{letter}{fmt.data_row}"))
            seen.add(header)
        elif header in fmt.ignored:
            seen.add(header)
        else:
            unknown.append(OrderValue(header, ref(f"{letter}{fmt.header_row}")))
            seen.add(header)
    missing = [h for h in fmt.columns if h not in seen]
    return OrderRecord(product_code, fields, unknown, missing, duplicate)
