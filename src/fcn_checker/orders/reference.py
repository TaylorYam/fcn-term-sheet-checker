"""參考條件表 adapter：多列表格，每列一檔商品（docs/order-formats/reference-sheet.md）。

依格式設定把各列轉成標準欄位（`OrderRecord`），同時記下每個欄位的 Excel 欄名與儲存格位置，供核對結果與回填使用。
標準欄位名怎麼組（`underlying_{i}`、`underlying_{i}_{價格}`、`autocall_date_{n}`）只在這裡定義，規則用型別化的讀法。
只讀第一張工作表，表頭列依配對欄（TDCC Code）找；只檢查必要欄（缺少或重複 → 整批錯誤），其他欄一律略過（Issue #199）。
"""

from __future__ import annotations

import datetime as dt
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from ..config import ReferenceFormat
from ..ingestion import IngestionError
from ..schema import OrderValue

SOURCE = "參考條件表"
OFFLINE_SHEET = "樣本清單"  # 離線建的列（`OrderRecord.offline`）儲存格位置用的工作表名稱
UNDERLYING_SLOTS = 5  # UL_1～UL_5
PRICE_FIELDS = {"initial": "initial_price", "strike": "strike_price", "ki": "ki_price", "ko": "ko_price"}


def underlying_key(i: int) -> str:
    """第 i 檔標的代號的標準欄位（UL_i）。"""
    return f"underlying_{i}"


def price_key(i: int, col: str) -> str:
    """第 i 檔標的某價格的標準欄位；`col` 為 initial／strike／ki／ko（UL_i_進場價 等）。"""
    return f"underlying_{i}_{PRICE_FIELDS[col]}"


def autocall_date_key(n: int) -> str:
    """第 n 期比價日的標準欄位（比價日_n）。"""
    return f"autocall_date_{n}"


def normalize_cell(v: Any) -> Any:
    """Excel 值標準化：數字轉 Decimal 並四捨五入到 9 位清除浮點尾數；日期轉 date；空字串轉 None。"""
    if v is None:
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, int):
        return Decimal(v)
    if isinstance(v, float):
        return Decimal(repr(round(v, 9)))
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, str):
        s = v.strip()
        return s or None
    return v


def header_names(fmt: ReferenceFormat) -> dict[str, list[str]]:
    """標準欄位 → 格式設定的全部 Excel 欄名，依宣告順序（例：currency → 承作幣別、計價幣別）。"""
    out: dict[str, list[str]] = {}
    for header, std in fmt.columns.items():
        out.setdefault(std, []).append(header)
    return out


def headers_of(fmt: ReferenceFormat) -> dict[str, str]:
    """標準欄位 → 格式設定的 Excel 欄名（同一標準欄位對到多個欄名時取先宣告的）。"""
    return {std: names[0] for std, names in header_names(fmt).items()}


@dataclass
class OrderRecord:
    """參考條件表的一列（核對規則的條件來源）：標準欄位 → 值、Excel 欄名與儲存格，一處完成。

    `fields` 只有這張表有的欄；`headers` 是格式設定宣告的全部欄名（表上缺的欄也查得到欄名，回填規則據此判斷）。
    必要欄缺少或重複時讀檔就整批錯誤，所以表上缺的只會是不是必要欄的欄（TS、IIS）。
    """

    row: int  # Excel 列號
    product_code: OrderValue
    issuer: OrderValue | None  # 表上沒有發行機構欄時為 None（不是必要欄）
    fields: dict[str, OrderValue]  # 標準欄位 → 值與來源儲存格（`OrderValue.column` 為 Excel 欄名）
    cells: dict[str, str]  # 標準欄位 → 儲存格位置（例：F4）
    headers: Mapping[str, str] = field(default_factory=dict)  # 標準欄位 → 格式設定的 Excel 欄名
    source: str = SOURCE  # 條件來源名稱，用在核對訊息

    def get(self, key: str) -> OrderValue | None:
        """標準欄位的值與來源；表上沒有這欄時為 None（值空白時是 value 為 None 的 OrderValue）。"""
        return self.fields.get(key)

    def has(self, key: str) -> bool:
        """表上有這一欄。"""
        return key in self.cells

    def cell(self, key: str) -> str | None:
        """儲存格位置（例：F4）；表上沒有這欄時為 None。"""
        return self.cells.get(key)

    def header(self, key: str) -> str | None:
        """格式設定給這個標準欄位的 Excel 欄名；格式設定沒有宣告時為 None。"""
        return self.headers.get(key)

    def underlying(self, i: int) -> OrderValue | None:
        """第 i 檔標的代號（UL_i）。"""
        return self.get(underlying_key(i))

    def price(self, i: int, col: str) -> OrderValue | None:
        """第 i 檔標的某價格；`col` 為 initial／strike／ki／ko。"""
        return self.get(price_key(i, col))

    def autocall_date(self, n: int) -> OrderValue | None:
        """第 n 期比價日（比價日_n）。"""
        return self.get(autocall_date_key(n))

    @property
    def underlying_slots(self) -> int:
        """格式設定有幾檔標的的價格欄（UL_1～UL_n）。"""
        suffix = f"_{PRICE_FIELDS['initial']}"
        return sum(1 for key in self.headers if key.startswith("underlying_") and key.endswith(suffix))

    @classmethod
    def offline(
        cls,
        fmt: ReferenceFormat,
        values: Mapping[str, Any],
        *,
        row: int = 4,
        issuer: Any = None,
        sheet: str = OFFLINE_SHEET,
    ) -> OrderRecord:
        """不經 Excel 直接以標準欄位值建一列（測試與離線用）：有格式設定的每個標準欄位（取先宣告的欄名），
        沒給值的欄是空白格；儲存格依宣告順序從 A 欄起編號。"""
        fields: dict[str, OrderValue] = {}
        cells: dict[str, str] = {}
        code = OrderValue(None, f"{sheet}!?{row}", fmt.key_column)
        for col, (std, header) in enumerate(headers_of(fmt).items(), start=1):
            cell = f"{get_column_letter(col)}{row}"
            value = _code(values.get(std)) if header == fmt.key_column else values.get(std)
            fields[std] = OrderValue(value, f"{sheet}!{cell}", header)
            cells[std] = cell
            if header == fmt.key_column:
                code = OrderValue(value, f"{sheet}!{cell}", fmt.key_column)
        return cls(
            row,
            code,
            OrderValue(issuer, f"{sheet}!?{row}", fmt.issuer_column),
            fields,
            cells,
            headers_of(fmt),
        )


@dataclass
class ReferenceSheet:
    sheet: str  # 讀的工作表名稱（第一張工作表）
    rows: list[OrderRecord]

    def find(self, product_code: str) -> list[OrderRecord]:
        return [r for r in self.rows if r.product_code.value == product_code]


def _code(v: object) -> str | None:
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip() if v is not None else ""
    return s or None


@dataclass(frozen=True)
class SheetLayout:
    """參考條件表在活頁簿裡的位置：第一張工作表；表頭是前幾列裡有配對欄（TDCC Code）欄名的那一列，下一列起是資料。"""

    sheet: str
    header_row: int

    @property
    def first_data_row(self) -> int:
        return self.header_row + 1


def locate(wb: Workbook, fmt: ReferenceFormat) -> SheetLayout:
    """找出參考條件表的工作表與表頭列；找不到或不只一列像表頭時丟出 IngestionError（整批不核對）。"""
    ws = wb.worksheets[0]
    rows = [
        r
        for r, values in enumerate(ws.iter_rows(max_row=fmt.header_search_rows, values_only=True), start=1)
        if any(normalize_cell(v) == fmt.key_column for v in values)
    ]
    where = f"參考條件表第一張工作表「{ws.title}」"
    if not rows:
        raise IngestionError(
            "reference_header_missing",
            f"{where}的前 {fmt.header_search_rows} 列找不到「{fmt.key_column}」欄名，無法確定表頭在哪一列",
        )
    if len(rows) > 1:
        listed = "、".join(f"第 {r} 列" for r in rows)
        raise IngestionError(
            "reference_header_ambiguous", f"{where}的{listed}都有「{fmt.key_column}」欄名，無法確定哪一列是表頭"
        )
    return SheetLayout(ws.title, rows[0])


def header_columns(ws: Worksheet, layout: SheetLayout) -> dict[str, list[int]]:
    """表頭的欄名 → 欄號（由左到右；同名欄有多個欄號，無表頭的欄略過）。讀檔、核對結果檔與回報包共用。"""
    out: dict[str, list[int]] = {}
    for col in range(1, ws.max_column + 1):
        header = normalize_cell(ws.cell(row=layout.header_row, column=col).value)
        if header is not None:
            out.setdefault(str(header), []).append(col)
    return out


def load_reference_sheet(path: Path, fmt: ReferenceFormat) -> ReferenceSheet:
    if not path.is_file():
        raise IngestionError("reference_not_found", f"找不到參考條件表檔案：{path.name}")
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:
        raise IngestionError("reference_unreadable", f"參考條件表無法開啟（可能已損毀或不是 Excel）：{e}") from e
    layout = locate(wb, fmt)
    ws = wb[layout.sheet]

    def ref(cell: str) -> str:
        return f"{layout.sheet}!{cell}"

    by_header = header_columns(ws, layout)
    problems = _column_problems(fmt, by_header)
    if problems:
        raise IngestionError(
            "reference_columns_invalid", f"參考條件表「{layout.sheet}」{'；'.join(problems)}，整批不核對"
        )
    std_cols = {fmt.columns[h]: cols[0] for h, cols in by_header.items() if h in fmt.columns}
    headers = {cols[0]: h for h, cols in by_header.items()}
    key_col = by_header[fmt.key_column][0]
    issuer_col = by_header[fmt.issuer_column][0] if fmt.issuer_column in by_header else None

    declared = headers_of(fmt)
    rows: list[OrderRecord] = []
    for r in range(layout.first_data_row, ws.max_row + 1):
        code = _code(ws.cell(row=r, column=key_col).value)
        if code is None:
            continue  # 空白列
        issuer = None
        if issuer_col:
            issuer_v = normalize_cell(ws.cell(row=r, column=issuer_col).value)
            issuer = OrderValue(issuer_v, ref(f"{get_column_letter(issuer_col)}{r}"), fmt.issuer_column)
        fields: dict[str, OrderValue] = {}
        cells: dict[str, str] = {}
        for std, col in std_cols.items():
            cell = f"{get_column_letter(col)}{r}"
            raw = ws.cell(row=r, column=col).value
            value = _code(raw) if col == key_col else normalize_cell(raw)
            fields[std] = OrderValue(value, ref(cell), headers[col])
            cells[std] = cell
        rows.append(
            OrderRecord(
                r,
                OrderValue(code, ref(f"{get_column_letter(key_col)}{r}"), fmt.key_column),
                issuer,
                fields,
                cells,
                declared,
            )
        )
    return ReferenceSheet(layout.sheet, rows)


def _column_problems(fmt: ReferenceFormat, by_header: Mapping[str, Sequence[int]]) -> list[str]:
    """表頭（欄名 → 欄號）的必要欄問題：缺少必要欄、會讀到的欄（[columns] 的欄與發行機構）重複；沒有問題時為空。

    同一標準欄位的欄名（例：承作幣別、計價幣別）合起來算一欄：有一個就不缺，出現兩次以上就算重複。
    不在格式設定的欄一律不管。依匯入格式的欄位順序列出（不在匯入格式的 TS、IIS、發行機構排最後）。"""
    groups = [*header_names(fmt).values(), [fmt.issuer_column]]  # 每組是同一欄的各種欄名
    rank = {h: i for i, h in enumerate(fmt.import_format.columns)}
    ordered = sorted(groups, key=lambda hs: min(rank.get(h, len(rank)) for h in hs))
    optional = {*fmt.optional_columns, fmt.issuer_column}
    missing = [
        hs[0] + (f"（或 {'、'.join(hs[1:])}）" if len(hs) > 1 else "")
        for hs in ordered
        if not set(hs) <= optional and not any(h in by_header for h in hs)
    ]
    repeated = [
        "／".join(h for h in hs if h in by_header) for hs in ordered if sum(len(by_header.get(h, ())) for h in hs) > 1
    ]
    return [
        *([f"缺少必要欄：{'、'.join(missing)}"] if missing else []),
        *([f"欄名重複：{'、'.join(repeated)}"] if repeated else []),
    ]
