"""讀參考條件表：只檢查必要欄，多出的欄、欄的順序、表頭列與工作表名稱都不限（Issue #199）。

測試切點：`load_reference_sheet`（Excel → 各列 `OrderRecord`；表本身有問題時丟出 `IngestionError`，整批不核對）。
"""

from __future__ import annotations

from decimal import Decimal

import openpyxl
import pytest

from fcn_checker.ingestion import IngestionError
from fcn_checker.orders.reference import load_reference_sheet
from harness import CONFIG
from reference_synth import REFERENCE_HEADERS, build_reference_sheet, make_row

FMT = CONFIG.reference_format
CODE = "029100000001"

# 仿公司原始待建檔表：沒有 TS、IIS、庫存狀態、當日比價、發行機構；幣別欄叫「計價幣別」；多了分公司等欄，順序也不同
_DROPPED = {"TS", "IIS", "庫存狀態", "當日比價", "發行機構", "承作幣別"}
ORIGINAL_HEADERS = [
    "TS/IIS完成檢核",
    "交易日期",
    *(h for h in REFERENCE_HEADERS if h not in _DROPPED and h != "單位面額"),
    "計價幣別",
    "單位面額",
    "分公司_1",
    "分公司_1_承作金額",
    "下一比價日",
]


def _row(**overrides):
    row = make_row("BARC", {"product_code": CODE, "denomination": 10000, "ko_pct": 100.0}, **overrides)
    row.update({"計價幣別": "USD", "TS/IIS完成檢核": "待檢核", "分公司_1": "城中", "分公司_1_承作金額": 340000})
    return row


def test_original_layout_is_read_by_column_name(tmp_path):
    path = build_reference_sheet(
        tmp_path / "原始.xlsx", [_row()], ORIGINAL_HEADERS, title="ELN待建檔_測試", header_row=2
    )
    sheet = load_reference_sheet(path, FMT)

    [row] = sheet.rows
    assert (row.row, row.product_code.value) == (3, CODE)
    assert (row.get("currency").value, row.get("currency").column) == ("USD", "計價幣別")
    assert row.get("denomination").value == Decimal(10000)
    assert row.get("ko_pct").source == "ELN待建檔_測試!" + row.cell("ko_pct")
    assert row.issuer is None, "表上沒有發行機構欄"
    assert not row.has("checked_term_sheet") and not row.has("checked_iis")


def _error(path) -> IngestionError:
    with pytest.raises(IngestionError) as e:
        load_reference_sheet(path, FMT)
    return e.value


def test_missing_required_columns_stop_the_whole_batch_and_are_listed_together(tmp_path):
    headers = [h for h in ORIGINAL_HEADERS if h not in ("KO(%)", "比價日_3", "計價幣別")]
    e = _error(build_reference_sheet(tmp_path / "缺欄.xlsx", [_row()], headers, header_row=2))
    assert e.reason_code == "reference_columns_invalid"
    assert "缺少必要欄：承作幣別（或 計價幣別）、比價日_3、KO(%)" in str(e)


def test_duplicate_columns_that_are_read_stop_the_whole_batch(tmp_path):
    headers = [*ORIGINAL_HEADERS, "KO(%)", "承作幣別", "TS", "TS", "分公司_1"]
    e = _error(build_reference_sheet(tmp_path / "重複.xlsx", [_row()], headers, header_row=2))
    assert e.reason_code == "reference_columns_invalid"
    assert "欄名重複：承作幣別／計價幣別、KO(%)、TS" in str(e)
    assert "分公司_1" not in str(e), "不讀的欄重複不提示"


def test_header_must_be_one_row_within_the_first_ten(tmp_path):
    late = _error(build_reference_sheet(tmp_path / "太下面.xlsx", [_row()], ORIGINAL_HEADERS, header_row=11))
    assert late.reason_code == "reference_header_missing" and "前 10 列找不到「TDCC Code」" in str(late)

    path = build_reference_sheet(tmp_path / "兩列.xlsx", [_row()], ORIGINAL_HEADERS, header_row=2)
    wb = openpyxl.load_workbook(path)
    wb.active["A1"] = "TDCC Code"
    wb.save(path)
    two = _error(path)
    assert two.reason_code == "reference_header_ambiguous" and "第 1 列、第 2 列" in str(two)


def test_only_the_first_worksheet_is_read(tmp_path):
    path = build_reference_sheet(tmp_path / "兩張.xlsx", [_row()], ORIGINAL_HEADERS, header_row=2)
    wb = openpyxl.load_workbook(path)
    wb.move_sheet(wb.create_sheet("說明"), offset=-wb.index(wb["說明"]))  # 新的空白工作表放到最左邊
    wb.save(path)
    assert _error(path).reason_code == "reference_header_missing"
