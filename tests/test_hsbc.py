"""HSBC 黑箱回歸：合成資料經批量入口（預覽＋核對）、CLI、PANEL；不直接測 parser。"""

from __future__ import annotations

import datetime as dt
from decimal import Decimal

import fitz
import openpyxl
import pytest

from fcn_checker.backfill import BackfillAction
from fcn_checker.batch import failed_batch
from fcn_checker.cli import main
from fcn_checker.ingestion import IngestionError
from fcn_checker.issuers import HSBC, REGISTRY
from fcn_checker.panel_workflow import PanelSession
from fcn_checker.schema import CheckReport
from fcn_checker.schema import CheckStatus as S
from harness import CONFIG, REVIEW_STANDARD, ROOT, STANDARD, check_all, cli_root, load_record, with_iis
from hsbc_synth import Spec, build_pdf
from pdf_writer import Edit
from reference_synth import REFERENCE_HEADERS, as_headers, build_reference_sheet, reference_row


def run_check(pdf, excel, registry=REGISTRY):
    """單份測試仍經公開批量入口，不直接呼叫 parser／規則；參考條件表本身有問題時同 CLI 記成整批錯誤。"""
    try:
        out = check_all(excel, [pdf], CONFIG.with_registry(registry))
    except IngestionError as e:
        out = failed_batch(excel, e)
    return out.items[0].report if out.items else CheckReport(out.status, None, list(out.errors), [])


def only(report, rule_id, field=None):
    out = [x for x in report.results if x.rule_id == rule_id and (field is None or x.field == field)]
    assert len(out) == 1, [(x.rule_id, x.field) for x in report.results if x.rule_id == rule_id]
    return out[0]


def sheet_for(tmp_path, s, overrides=None):
    """與 `s` 一致的參考條件表（正式版面，一列）；overrides 以標準欄位名覆寫。"""
    return build_reference_sheet(tmp_path / "order.xlsx", [reference_row(s, **as_headers(overrides or {}))])


def check(tmp_path, spec=None, overrides=None, edits=()):
    s = spec or Spec()
    return run_check(
        build_pdf(tmp_path / f"{s.product_code}_TS.pdf", s, edits=edits), sheet_for(tmp_path, s, overrides)
    )


@pytest.mark.parametrize(
    "obs,memory,ki,count",
    [
        ("D", True, "AM", 2),
        ("D", True, "none", 2),
        ("D", False, "none", 1),
        ("P", False, "none", 1),
        ("P", False, "AM", 2),
        ("P", True, "AM", 2),
        ("D", False, "D", 2),
    ],
)
def test_supported_types_pass(tmp_path, obs, memory, ki, count):
    r = check(tmp_path, Spec(ko_obs=obs, memory=memory, ki=ki, count=count))
    assert r.template == HSBC.template_id
    assert r.status == S.PASS, [
        (x.rule_id, x.field, x.reason_code) for x in r.results if x.status not in (S.PASS, S.NOT_APPLICABLE)
    ]
    assert all(x.document_evidence for x in r.results if x.rule_id.startswith(("doc.", "derive.", "schedule.")))
    assert {x["rule_id"] for x in r.not_covered} == {x["rule_id"] for x in HSBC.not_covered}


@pytest.mark.parametrize(
    "field,value,rule",
    [
        ("isin", "XS1999900002", "backfill.isin"),
        ("denomination", 20000, "field.denomination"),
        ("trade_date", "2031-01-07", "field.trade_date"),
        ("issue_date", dt.date(2031, 1, 14), "backfill.issue_date"),
        ("final_valuation_date", "2031-07-07", "field.final_valuation_date"),
        ("maturity_date", "2031-07-10", "field.maturity_date"),
        ("ko_pct", 101, "field.ko_pct"),
        ("strike_pct", 71, "field.strike_pct"),
        ("ki_pct", 61, "field.ki_pct"),
        ("coupon_pa_pct", 13, "field.coupon_pa_pct"),
        ("tenor_months", 7, "field.tenor_months"),
        ("first_callable_period", 3, "field.first_callable_period"),
        ("currency", "JPY", "field.currency"),
        ("ko_memory", "N", "field.ko_memory"),
        ("ko_observation", "P", "field.ko_observation"),
        ("ki_type", "D", "field.ki_type"),
        ("underlying_1", "ZZ9 UW", "field.underlyings"),
        ("underlying_1", "ZZ1 UN", "field.underlyings"),
        ("underlying_1_initial_price", 101, "field.underlying_prices"),
        ("underlying_1_strike_price", 71, "field.underlying_prices"),
        ("underlying_1_ki_price", 61, "field.underlying_prices"),
        ("underlying_1_ko_price", 101, "field.underlying_prices"),
        ("underlying_5", "ZZ5 UW", "field.underlyings"),
        ("autocall_date_3", "2030-04-07", "backfill.compare_dates"),
        ("autocall_date_2", "-", "backfill.compare_dates"),
        ("autocall_date_6", "2030-07-08", "backfill.compare_dates"),
    ],
)
def test_order_difference_is_reported(tmp_path, field, value, rule):
    r = check(tmp_path, overrides={field: value})
    assert any(x.rule_id == rule and x.status in (S.MISMATCH, S.REVIEW_REQUIRED) for x in r.results)


@pytest.mark.parametrize(
    "old,new,rule",
    [
        ("最低交易金額：美元10,000元", "最低交易金額：美元20,000元", "field.min_amounts"),
        ("商品開始受理申購日：2030年1月7日", "商品開始受理申購日：2030年1月8日", "doc.subscription_start_date"),
        ("(最終版)刊印日期：2030年1月7日", "(最終版)刊印日期：2030年1月9日", "doc.print_date"),
        (
            "[受託或銷售機構]審查通過之日期：2026年6月11日",
            "[受託或銷售機構]審查通過之日期：2026年6月12日",
            "standard.approval_date",
        ),
        ("0%~5%", "0%~6%", "standard.fees"),
        (
            "固定配息金額=美元10,000×1.0000%=美元100.00",
            "固定配息金額=美元10,000×1.0000%=美元101.00",
            "doc.scenario_calculations",
        ),
        ("6個計息期間配息金額共為美元600.00", "5個計息期間配息金額共為美元600.00", "doc.scenario_parameters"),
        (
            "到期贖回金額為美元10,000×100%=美元10,000.00",
            "到期贖回金額為美元10,000×100%=美元10,001.00",
            "doc.scenario_calculations",
        ),
        (
            "平均年化報酬率(以簡單平均年化報酬率之方式計算)為12.00%",
            "平均年化報酬率(以簡單平均年化報酬率之方式計算)為13.00%",
            "doc.scenario_general_annualized",
        ),
        ("自動提前到期價格為期初股價×100%", "自動提前到期價格為期初股價×101%", "doc.price_header_pct"),
        ("70.0000", "71.0000", "derive.prices"),
        ("配息期數=6", "配息期數=5", "doc.coupon_periods"),
        ("發行價格：100%", "發行價格：99%", "standard.issue_price"),
        (
            "商品天期為6個月期，每單位面額為美元10,000元",
            "商品天期為7個月期，每單位面額為美元10,000元",
            "doc.scenario_parameters",
        ),
    ],
)
def test_document_difference_is_reported(tmp_path, old, new, rule):
    r = check(tmp_path, edits=[Edit(old, new, "ts")])
    assert any(x.rule_id == rule and x.status in (S.MISMATCH, S.REVIEW_REQUIRED) for x in r.results)


@pytest.mark.parametrize(
    "field,value",
    [
        ("ko_observation", "Q"),
        ("ko_memory", "X"),
        ("ki_type", "M"),
        ("coupon_pa_pct", "?"),
        ("underlying_1", "-"),
    ],
)
def test_unknown_or_missing_order_value_requires_review(tmp_path, field, value):
    r = check(tmp_path, overrides={field: value})
    assert any(x.status == S.REVIEW_REQUIRED and x.rule_id.startswith("field.") for x in r.results)


@pytest.mark.parametrize(
    "change,reason",
    [
        ("missing_row", "reference_row_missing"),
        ("duplicate_row", "reference_row_duplicate"),
        ("issuer", "reference_issuer_mismatch"),
        ("missing_key", "reference_key_missing"),
    ],
)
def test_table_pairing_failures_stop_rules(tmp_path, change, reason):
    s = Spec()
    pdf = build_pdf(tmp_path / f"{s.product_code}_TS.pdf", s)
    rows, headers = [reference_row(s)], None
    if change == "missing_row":
        rows = [reference_row(s, **{"TDCC Code": "325199990002"})]
    elif change == "duplicate_row":
        rows = [reference_row(s), reference_row(s)]
    elif change == "issuer":
        rows = [reference_row(s, 發行機構="BARC")]
    else:
        headers = ["不存在的欄位" if h == "TDCC Code" else h for h in REFERENCE_HEADERS]
    r = run_check(pdf, build_reference_sheet(tmp_path / "order.xlsx", rows, headers))
    assert any(x.reason_code == reason for x in r.results), [(x.reason_code, x.message) for x in r.results]
    assert not any(x.rule_id.startswith("field.") for x in r.results)


@pytest.mark.parametrize(
    "spec",
    [
        Spec(strike=Decimal("65")),
        Spec(ko=Decimal("105"), ki_pct=Decimal("55")),
        Spec(tenor=12),
        Spec(trade_date=dt.date(2030, 3, 9), issue_date=dt.date(2030, 3, 16)),
        Spec(ko_obs="P", memory=False, ki="none", first_callable=3),
        Spec(currency_zh="人民幣"),
    ],
    ids=["strike", "ko-ki", "tenor", "trade-date", "non-call", "currency"],
)
def test_spec_values_are_drawn_into_both_documents(tmp_path, spec):
    """說明書與投資人須知照規格畫：改了規格值，兩份文件仍與同一規格的參考條件表列一致（Issue #166）。"""
    pdf = build_pdf(tmp_path / f"{spec.product_code}_TS.pdf", spec)
    items = check_all(sheet_for(tmp_path, spec), [pdf], CONFIG).items
    assert [i.report.status for i in items] == [S.PASS, S.PASS], [i.problem_messages for i in items]


def test_spec_value_differing_from_the_reference_row_is_reported(tmp_path):
    s = Spec(strike=Decimal("65"))
    r = run_check(build_pdf(tmp_path / f"{s.product_code}_TS.pdf", s), sheet_for(tmp_path, Spec()))
    assert only(r, "field.strike_pct").status == S.MISMATCH


def test_price_table_currency_is_the_trade_currency_on_the_sheet(tmp_path):
    """第 12 條價格表每列的幣別格是承作幣別：逐列和參考條件表「承作幣別」比（Issue #167）。"""
    r = check(tmp_path, edits=[Edit("USD", "JPY", "ts.art12.table")])  # 例：連結日股、美元計價卻寫成標的的幣別
    rows = {x.field: x for x in r.results if x.rule_id == "field.currency"}
    assert rows["currency"].status == S.PASS, "封面計價幣別正確"
    assert {f: x.status for f, x in rows.items() if f.startswith("price_table_currency")} == {
        "price_table_currency_1": S.MISMATCH,
        "price_table_currency_2": S.MISMATCH,
    }
    assert all(x.status == S.PASS for f, x in rows.items() if not f.startswith("price_table_currency")), "其他出處沒改"
    bad = rows["price_table_currency_1"]
    assert (bad.expected, bad.actual) == ("USD", "JPY") and bad.document_evidence
    assert bad.item.name == "ZZ1 UW 幣別"


def test_product_currency_need_not_match_the_underlying_currency(tmp_path):
    """商品幣別不必等於標的幣別：人民幣計價、連結美股，每一處幣別都 = 參考條件表承作幣別（Issue #170）。"""
    r = check(tmp_path, Spec(currency_zh="人民幣"))
    rows = [x for x in r.results if x.rule_id == "field.currency"]
    assert sum(x.field.startswith("price_table_currency") for x in rows) == 2
    assert all(x.status == S.PASS and x.expected == "CNH" for x in rows), [(x.field, x.status) for x in rows]


def test_blank_trade_currency_on_the_sheet_is_reported_once(tmp_path):
    r = check(tmp_path, overrides={"currency": None})
    [x] = [x for x in r.results if x.rule_id == "field.currency"]
    assert (x.status, x.reason_code) == (S.REVIEW_REQUIRED, "order_missing"), "價格表各列不重複報同一格"


def test_wrong_prefix_and_unknown_template_require_review(tmp_path):
    r = check(tmp_path, Spec(product_code="029199990001"))
    assert any(x.reason_code == "issuer_prefix_mismatch" for x in r.results)
    other = tmp_path / "other"
    other.mkdir()
    r = check(other, edits=[Edit("中文產品說明書(最終版)", "未知說明書", "ts")])
    assert r.status == S.REVIEW_REQUIRED and r.template is None


def test_partial_coupon_and_unrounded_total(tmp_path):
    # 情境分析照年利率畫：月配息率 0.9992%、每期 99.92，合計以年利率算後才四捨五入（599.50、199.83）
    s = Spec(annual=Decimal("11.99"))
    assert check(tmp_path, s).status == S.PASS


def test_cli_and_panel_select_hsbc(tmp_path, monkeypatch):
    s = Spec()
    pdf = build_pdf(tmp_path / f"{s.product_code}_TS.pdf", s)
    excel = sheet_for(tmp_path, s)
    cli_root(tmp_path, monkeypatch)
    assert main([str(excel), *map(str, with_iis([pdf])), "--out", str(tmp_path / "reports")]) == 0
    assert load_record(tmp_path)["items"][0]["template"] == HSBC.template_id
    assert "HSBC" in [x.code for x in REGISTRY]
    session = PanelSession(REVIEW_STANDARD, ROOT / "config")
    session.select(excel, with_iis([pdf]))
    session.load_preview()
    assert session.start_check().batch.status == S.PASS


def test_ambiguous_registry_and_damaged_pdf(tmp_path):
    s = Spec()
    pdf = build_pdf(tmp_path / f"{s.product_code}_TS.pdf", s)
    excel = sheet_for(tmp_path, s)
    from fcn_checker.issuers import Issuer
    from fcn_checker.parsers import hsbc as parser

    fake = Issuer(
        code="FAKE",
        template_id="fake-zh-pd",
        label="FAKE 測試範本",
        parser_version="test",
        not_covered=(),
        detect=parser.detect,
        read=parser.read,
        rules=lambda ctx: [],
    )
    r = run_check(pdf, excel, registry=(HSBC, fake))
    assert any(x.reason_code == "template_ambiguous" for x in r.results)
    pdf.write_bytes(b"broken pdf")
    assert run_check(pdf, excel).status == S.ERROR


def test_encrypted_pdf_is_not_checked(tmp_path):
    s = Spec()
    plain = build_pdf(tmp_path / f"{s.product_code}_TS.pdf", s)
    excel = sheet_for(tmp_path, s)
    (tmp_path / "locked").mkdir()
    encrypted = tmp_path / "locked" / f"{s.product_code}_TS.pdf"
    with fitz.open(plain) as d:
        d.save(encrypted, encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="fake-owner", user_pw="fake-user")
    assert run_check(encrypted, excel).status == S.ERROR


@pytest.mark.parametrize(
    "field,new,rule",
    [
        ("chairman", "不同負責人", "standard.chairman"),
        ("name", "不同銷售機構", "standard.distributor"),
        ("address", "不同地址", "standard.distributor"),
    ],
)
def test_distributor_standard_difference(tmp_path, field, new, rule):

    old = STANDARD["distributor"][field]
    # Change all places containing this value, including longer labels.
    if field == "chairman":
        oldline = "(c)營業所在地：" + STANDARD["distributor"]["address"] + "(d)負責人姓名：" + old
    elif field == "name":
        oldline = "受託或銷售機構之名稱、電話及地址：" + old
    else:
        oldline = old
    r = check(tmp_path, edits=[Edit(oldline, oldline.replace(old, new), "ts")])
    assert any(x.rule_id == rule and x.status == S.MISMATCH for x in r.results)


def test_warning_risk_wording_and_name_rules(tmp_path):

    warning = STANDARD["risk"]["fixed_warning_by_issuer"]["hsbc"]
    r = check(tmp_path, edits=[Edit(warning, warning.replace("RR4", "RR3"), "ts")])
    assert any(x.rule_id == "standard.fixed_warning" and x.status == S.MISMATCH for x in r.results)
    assert any(x.rule_id == "standard.risk_level" and x.status == S.MISMATCH for x in r.results)
    other = tmp_path / "wording"
    other.mkdir()
    r = check(other, edits=[Edit("其他說明", "受託投資", "ts")])
    assert any(x.rule_id == "standard.forbidden_wording" and x.status == S.MISMATCH for x in r.results)
    other = tmp_path / "name"
    other.mkdir()
    en = Spec().en
    r = check(
        other,
        edits=[
            Edit(
                f"商品英文名稱：{en}", "商品英文名稱：Autocallable Fixed Coupon Notes (Non Guaranteed, Unsecured)", "ts"
            )
        ],
    )
    assert any(x.rule_id == "standard.product_name" and x.status == S.MISMATCH for x in r.results)
    assert any(x.rule_id == "doc.name_consistency" and x.status == S.MISMATCH for x in r.results)


@pytest.mark.parametrize(
    "old,new,rule",
    [
        ("70.0000", "71.0000", "doc.scenario_table"),
        ("執行價(即期初股價的70%)", "執行價(即期初股價的71%)", "doc.scenario_header_pct"),
    ],
)
def test_scenario_table_is_independently_compared(tmp_path, old, new, rule):
    r = check(tmp_path, edits=[Edit(old, new, "ts.art18.table")])
    assert any(x.rule_id == rule and x.status == S.MISMATCH for x in r.results)


@pytest.mark.parametrize(
    "old,new,rule",
    [
        ("固定配息率=12%×1/12", "固定配息率=12%×1/12固定配息率=13%×1/12", "field.coupon_pa_pct"),
        ("固定配息率=12%×1/12", "", "field.coupon_pa_pct"),
        ("每單位面額：美元10,000元", "每單位面額：美元20,000元", "doc.denomination"),
        ("註1：假設日期", "未知表格結尾", "schedule.coupon_dates"),
        ("固定配息金額=美元10,000×1.0000%=美元100.00", "未知算式", "doc.scenario_calculations"),
    ],
)
def test_missing_ambiguous_or_unknown_document_requires_review(tmp_path, old, new, rule):
    r = check(tmp_path, edits=[Edit(old, new, "ts")])
    assert any(x.rule_id == rule and x.status == S.REVIEW_REQUIRED for x in r.results)


@pytest.mark.parametrize(
    "mode,rule",
    [
        ("unknown", "order.unknown_column"),
        ("missing", "order.missing_column"),
        ("duplicate", "order.duplicate_column"),
        ("no_header", None),
    ],
)
def test_header_validation(tmp_path, mode, rule):
    s = Spec()
    pdf = build_pdf(tmp_path / f"{s.product_code}_TS.pdf", s)
    excel = sheet_for(tmp_path, s)
    wb = openpyxl.load_workbook(excel)
    ws = wb.active
    if mode == "missing":
        ws.cell(3, 2).value = None
    else:
        col = ws.max_column + 1
        ws.cell(3, col, {"unknown": "未知新欄位", "duplicate": "ISIN Code", "no_header": None}[mode])
        ws.cell(4, col, "額外資料")
    wb.save(excel)
    wb.close()
    r = run_check(pdf, excel)
    if rule:
        assert any(x.rule_id == rule and x.status == S.REVIEW_REQUIRED for x in r.results)
    else:
        assert r.status == S.PASS


def test_panel_preview_selects_the_matching_table_row(tmp_path):
    s = Spec()
    pdf = build_pdf(tmp_path / f"{s.product_code}_TS.pdf", s)
    other = Spec(product_code="325199990002")  # 前面多一列別的商品，這份對到第 5 列
    excel = build_reference_sheet(tmp_path / "order.xlsx", [reference_row(other), reference_row(s)])
    session = PanelSession(REVIEW_STANDARD, ROOT / "config")
    session.select(excel, with_iis([pdf]))
    preview = session.load_preview()
    assert preview.rows[0].product_code == s.product_code
    assert preview.rows[0].reference_row == 5
    assert session.start_check().batch.status == S.PASS


@pytest.mark.parametrize(
    "field,value,status",
    [
        # 說明書顯示到整數（12%、100%、70%）：表上值依顯示位數 half-up 後比對（BARC 比法）
        ("coupon_pa_pct", "12.4999", S.PASS),
        ("coupon_pa_pct", "12.5", S.MISMATCH),
        ("ko_pct", "100.004", S.PASS),
        ("ko_pct", "100.5", S.MISMATCH),
        ("strike_pct", "69.5", S.PASS),
        ("strike_pct", "70.5", S.MISMATCH),
    ],
)
def test_percentages_are_rounded_to_the_displayed_digits(tmp_path, field, value, status):
    r = check(tmp_path, overrides={field: value})
    x = only(r, "field." + field)
    assert x.status == status
    assert x.tolerance == "依說明書顯示位數四捨五入後比對"


@pytest.mark.parametrize(
    "rule,fields",
    [
        ("field.min_amounts", {"minimum_trade", "minimum_subscription", "minimum_additional"}),
        ("doc.subscription_start_date", {"subscription_start", "subscription_end"}),
        ("doc.print_date", {"print_date_review", "print_date_final"}),
    ],
)
def test_each_review_standard_occurrence_is_checked_with_the_shared_rule(tmp_path, rule, fields):
    r = check(tmp_path)
    assert {x.field for x in r.results if x.rule_id == rule and x.status == S.PASS} == fields


@pytest.mark.parametrize(
    "old,new,rule,field",
    [
        (
            "商品申購結束受理日：2030年1月7日",
            "商品申購結束受理日：2030年1月8日",
            "doc.subscription_start_date",
            "subscription_end",
        ),
        (
            "最低加購金額：美元10,000元",
            "最低加購金額：美元20,000元",
            "field.min_amounts",
            "minimum_additional",
        ),
        (
            "(參考性審閱版)內容，刊印日期：2030年1月7日",
            "(參考性審閱版)內容，刊印日期：2030年1月20日",
            "doc.print_date",
            "print_date_review",
        ),
    ],
)
def test_only_the_differing_occurrence_mismatches(tmp_path, old, new, rule, field):
    r = check(tmp_path, edits=[Edit(old, new, "ts")])
    bad = {x.field for x in r.results if x.rule_id == rule and x.status == S.MISMATCH}
    assert bad == {field}


def test_non_default_denomination_requires_review_like_barc(tmp_path):
    r = check(tmp_path, edits=[Edit("每單位面額：美元10,000元", "每單位面額：美元20,000元", "ts")])
    x = only(r, "doc.denomination")
    assert (x.status, x.reason_code) == (S.REVIEW_REQUIRED, "denomination_non_default")


def test_ki_pct_must_be_empty_when_the_document_has_no_ki(tmp_path):
    r = check(tmp_path, Spec(ki="none"), overrides={"ki_pct": 60})
    x = only(r, "field.ki_pct")
    assert (x.status, x.reason_code) == (S.MISMATCH, "value_mismatch")


def test_non_integer_denomination_requires_review(tmp_path):
    r = check(tmp_path, overrides={"denomination": "10000.5"})
    x = only(r, "field.denomination")
    assert (x.status, x.reason_code) == (S.REVIEW_REQUIRED, "order_invalid")


def test_reference_fields_use_the_shared_rule_ids(tmp_path):
    r = check(tmp_path, Spec(ki="none"))
    ids = {x.rule_id for x in r.results}
    assert {"field.underlyings", "field.underlying_prices", "field.currency", "field.ki_pct"} <= ids
    assert not ids & {"field.prices", "field.isin", "field.autocall_dates"}
    for rid in ("field.first_callable_period", "backfill.isin", "backfill.compare_dates"):
        assert sum(x.rule_id == rid for x in r.results) == 1, rid
    ki_price = only(r, "field.underlying_prices", "ZZ1 UW 下限價")
    assert ki_price.status == S.NOT_APPLICABLE and ki_price.document_evidence


@pytest.mark.parametrize(
    "old,new,field",
    [
        ("到期贖回金額為美元10,000×100%=美元10,000.00", "", "principal_formula"),
        ("到期贖回金額為美元10,000×100%=美元10,000.00", "到期贖回金額為未知算式", "principal_formula"),
        ("假設標的（虛構標的1），執行價美元70.0000", "假設標的（虛構標的1），執行價未知", "reference_strike"),
        ("觸及不保本價格美元60.0000", "觸及不保本價格未知", "reference_ki"),
    ],
)
def test_each_required_scenario_formula_or_reference_must_be_present(tmp_path, old, new, field):
    r = check(tmp_path, edits=[Edit(old, new, "ts")])
    assert any(x.status == S.REVIEW_REQUIRED and field in x.field for x in r.results)


def test_daily_schedule_end_must_not_precede_start(tmp_path):
    # Third period's start follows period 2, but its end is moved before that start.
    r = check(tmp_path, edits=[Edit("2030 年4 月7 日", "2030 年3 月6 日", "ts")])
    assert any(x.rule_id == "schedule.autocall_dates" and x.status == S.MISMATCH for x in r.results)


@pytest.mark.parametrize(
    "change,reason",
    [
        ({"投資人應注意": "投資人注意事項"}, "document_missing"),  # 定期KO表沒有結束錨點
        ({"2030 年5 月7 日": "2030 年5 月XX日"}, "document_invalid"),  # 第 3 列決定日不是日期；配息表只印付款日
    ],
)
def test_periodic_ko_table_unreadable_requires_review_not_unexpected(tmp_path, change, reason):
    """定期觀察的提前出場表讀不到：schedule.autocall_dates 轉人工覆核，整份不能記成非預期錯誤。"""
    edits = [Edit(old, new, "ts") for old, new in change.items()]
    r = check(tmp_path, Spec(ko_obs="P", memory=False, ki="none", count=1), edits=edits)
    assert not any(x.rule_id == "batch.unexpected" for x in r.results), [(x.rule_id, x.message) for x in r.results]
    autocall = only(r, "schedule.autocall_dates")
    assert (autocall.status, autocall.reason_code) == (S.REVIEW_REQUIRED, reason)
    assert autocall.message.startswith("說明書")
    assert only(r, "field.ko_pct").status == S.PASS
    assert only(r, "field.ko_observation").status == S.PASS


@pytest.mark.parametrize("amount,status", [("25.00", S.PASS), ("26.00", S.MISMATCH)])
def test_partial_period_coupon_arithmetic(tmp_path, amount, status):
    s = Spec(partial_coupon=True)
    old = "第3個計息期間配息金額=美元10,000×1.0000%×5/20=美元25.00"
    r = check(tmp_path, s, edits=[Edit(old, old.replace("美元25.00", f"美元{amount}"), "ts")])
    assert r.status == status
    assert any(
        x.rule_id == "doc.scenario_calculations"
        and "coupon_amount" in x.field
        and x.actual == Decimal(amount)
        and x.status == status
        for x in r.results
    )


@pytest.mark.parametrize(
    "items,total,status",
    [
        ("+美元25.00", "225.01", S.PASS),
        ("+美元25.00", "224.99", S.PASS),
        ("+美元25.00", "225.02", S.MISMATCH),
        ("+美元25.01", "225.01", S.MISMATCH),
    ],
)
def test_profit_total_allows_rounding_difference_only(tmp_path, items, total, status):
    s = Spec(partial_coupon=True)
    old = "損益=美元10,000.00+美元200.00+美元25.00-美元10,000.00=美元225.00"
    r = check(tmp_path, s, edits=[Edit(old, f"損益=美元10,000.00+美元200.00{items}-美元10,000.00=美元{total}", "ts")])
    [profit] = [x for x in r.results if x.field.startswith("s1.profit.")]
    assert profit.rule_id == "doc.scenario_calculations" and profit.tolerance
    assert profit.status == status and profit.actual == Decimal(total)
    assert r.status == status


@pytest.mark.parametrize("obs", ["D", "P"])
def test_hsbc_batch_backfills_shared_reference_sheet(tmp_path, obs):
    from fcn_checker.saving import run_batch

    s = Spec(ko_obs=obs, ki="none")
    pdf = build_pdf(tmp_path / f"{s.product_code}_TS.pdf", s)
    excel = sheet_for(tmp_path, s)  # 回填欄位（ISIN、發行日、比價日）預設空白
    r, receipt = run_batch(CONFIG, excel, with_iis([pdf]), tmp_path / "reports", root=tmp_path)
    assert receipt.status == S.PASS, [
        (x.rule_id, x.field, x.status, x.reason_code)
        for x in r.items[0].report.results
        if x.status not in (S.PASS, S.NOT_APPLICABLE)
    ]
    assert receipt.filled(r.items[0])
    wb = openpyxl.load_workbook(receipt.output)
    ws = wb["回填後"]
    vals = {c.value: ws.cell(4, c.column).value for c in ws[3]}
    assert vals["ISIN Code"] == "XS1999900001"
    assert vals["發行日"] == dt.datetime(2030, 1, 14)
    assert vals["比價日_2"].date() == s.ends[1]
    assert vals["比價日_6"] == dt.datetime.combine(s.ends[-1], dt.time()), "D 型也填最後一期"
    assert vals["比價日_3"] == ("-" if obs == "D" else dt.datetime.combine(s.ends[2], dt.time()))
    assert vals["比價日_1"] == "-"
    wb.close()
    wb = openpyxl.load_workbook(excel)
    ws = wb.active
    assert all(
        ws.cell(4, c.column).value is None
        for c in ws[3]
        if c.value in ("ISIN Code", "發行日") or str(c.value).startswith("比價日_")
    )
    wb.close()


def test_prefilled_backfill_columns_that_match_pass_and_stay(tmp_path):
    """回填欄位（ISIN、發行日、比價日）事先填好且與說明書相同：全部 MATCH，整份通過。"""
    s = Spec()
    dates = {f"autocall_date_{i}": s.ends[i - 1] if i in (2, 6) else "-" for i in range(1, 13)}
    r = check(tmp_path, s, overrides={"isin": "XS1999900001", "issue_date": dt.date(2030, 1, 14), **dates})

    assert r.status == S.PASS, [(x.rule_id, x.field, x.reason_code) for x in r.results if x.status != S.PASS]
    assert {d.action for d in r.backfill if d.column not in ("TS", "IIS")} == {BackfillAction.MATCH}, "TS、IIS 打勾另填"


@pytest.mark.parametrize(
    "sheet_value,action,status,kept",
    [
        (None, BackfillAction.FILL, S.PASS, dt.datetime(2030, 1, 14)),
        (dt.date(2030, 1, 14), BackfillAction.MATCH, S.PASS, dt.datetime(2030, 1, 14)),
        (dt.date(2030, 1, 15), BackfillAction.MISMATCH, S.MISMATCH, dt.datetime(2030, 1, 15)),
    ],
)
def test_hsbc_issue_date_is_a_backfill_column(tmp_path, sheet_value, action, status, kept):
    from fcn_checker.saving import run_batch

    s = Spec()
    r, receipt = run_batch(
        CONFIG,
        sheet_for(tmp_path, s, {"issue_date": sheet_value}),
        with_iis([build_pdf(tmp_path / f"{s.product_code}_TS.pdf", s)]),
        tmp_path / "reports",
        root=tmp_path,
    )
    report = r.items[0].report
    assert only(report, "backfill.issue_date").status == status
    assert [d.action for d in report.backfill if d.column == "發行日"] == [action]
    assert report.status == status
    assert not any(x.rule_id == "field.issue_date" for x in report.results)
    ws = openpyxl.load_workbook(receipt.output)["回填後"]
    if status == S.PASS:
        assert {c.value: ws.cell(4, c.column).value for c in ws[3]}["發行日"] == kept
    else:
        assert ws.max_row == 3, "沒通過的列不出現在「回填後」"


# ---------------------------------------------------------------- 價格推算（各上手共用規則，語意同 BARC）


class _Replaced:
    """HSBC 讀出結果，但把部分標準欄位換掉（模擬價格表讀出與標的數或 KI 型態不一致）。"""

    def __init__(self, ts, **fields):
        self._ts, self._fields = ts, fields
        self.full_text = ts.full_text

    def f(self, name):
        return self._fields[name](self._ts.f(name)) if name in self._fields else self._ts.f(name)

    def __getattr__(self, attr):
        return getattr(self._ts, attr)


def _check_replaced(tmp_path, **fields):
    import dataclasses

    adapter = dataclasses.replace(HSBC, read=lambda lines: _Replaced(HSBC.read(lines), **fields))
    s = Spec()
    return run_check(build_pdf(tmp_path / f"{s.product_code}_TS.pdf", s), sheet_for(tmp_path, s), (adapter,))


def test_hsbc_price_rows_fewer_than_underlyings_require_review(tmp_path):
    import dataclasses

    r = _check_replaced(tmp_path, underlying_prices=lambda pf: dataclasses.replace(pf, value=pf.value[:1]))
    x = only(r, "derive.prices")
    assert (x.status, x.reason_code, x.expected, x.actual) == (S.REVIEW_REQUIRED, "price_table_row_count", 2, 1)


def test_hsbc_price_table_without_ki_column_for_a_ki_product_requires_review(tmp_path):
    import dataclasses

    from fcn_checker.standard_fields import PriceRow

    def drop_ki(pf):
        rows = tuple(PriceRow(r.ticker, {k: v for k, v in r.prices.items() if k != "ki"}, r.evidence) for r in pf.value)
        return dataclasses.replace(pf, value=rows)

    r = _check_replaced(tmp_path, underlying_prices=drop_ki)
    x = only(r, "derive.prices")
    assert (x.status, x.reason_code) == (S.REVIEW_REQUIRED, "price_table_ki_column")


# ---------------------------------------------------------------- 金額旁的幣別（Issue #170）


@pytest.mark.parametrize(
    ("section", "old", "new", "field"),
    [
        ("art6", "每單位面額：美元10,000元", "每單位面額：日幣10,000元", "denomination_currency"),
        ("art7", "最低交易金額：美元10,000元", "最低交易金額：日幣10,000元", "minimum_trade_currency"),
        ("ch四", "最低申購金額：美元10,000元", "最低申購金額：日幣10,000元", "minimum_subscription_currency"),
        ("ch四", "最低加購金額：美元10,000元", "最低加購金額：日幣10,000元", "minimum_additional_currency"),
        ("art18", "每單位面額為美元10,000元", "每單位面額為日幣10,000元", "scenario_assumption_currency"),
        ("art18", "-美元10,000.00=美元200.00", "-美元10,000.00=日幣200.00", "scenario_1_currency"),
    ],
    ids=lambda x: (
        (x if isinstance(x, str) else "+".join(x))
        if isinstance(x, (str, tuple)) and str(x).endswith(("_currency", "')"))
        else ""
    ),
)
def test_wrong_currency_next_to_an_amount_is_reported_at_that_place(tmp_path, section, old, new, field):
    fields = field if isinstance(field, tuple) else (field,)
    r = check(tmp_path, edits=[Edit(old, new, f"ts.{section}")])
    bad = [(x.rule_id, x.field, x.status) for x in r.results if x.status in (S.MISMATCH, S.REVIEW_REQUIRED)]
    assert bad == [("field.currency", f, S.MISMATCH) for f in fields]
    for x in [x for x in r.results if x.rule_id == "field.currency" and x.field in fields]:
        assert x.expected == "USD" and x.document_evidence
        assert x.actual == "JPY" or "日幣" in str(x.actual), "單一出處顯示 ISO 代碼，情境列出不符的金額"


def test_every_scenario_has_its_own_currency_result(tmp_path):
    r = check(tmp_path)
    fields = [x.field for x in r.results if x.rule_id == "field.currency" and x.field.startswith("scenario_")]
    assert fields == [
        "scenario_assumption_currency",
        "scenario_1_currency",
        "scenario_2_currency",
        "scenario_3_currency",
        "scenario_4_currency",
    ]


def test_strike_and_ki_prices_in_the_scenarios_follow_the_underlying(tmp_path):
    """執行價、觸及不保本價格跟著標的走，不是承作幣別的出處（2026-10-10 使用者確認）；數字仍和價格表比。"""
    edits = [
        Edit("執行價美元70.0000", "執行價日幣70.0000", "ts.art18"),
        Edit("觸及不保本價格美元60.0000", "觸及不保本價格日幣60.0000", "ts.art18"),
    ]
    r = check(tmp_path, edits=edits)
    assert r.status == S.PASS, [(x.rule_id, x.field) for x in r.results if x.status not in (S.PASS, S.NOT_APPLICABLE)]
    r = check(tmp_path, edits=[Edit("執行價美元70.0000", "執行價日幣71.0000", "ts.art18")])
    assert any(x.rule_id == "doc.scenario_parameters" and x.status == S.MISMATCH for x in r.results)
