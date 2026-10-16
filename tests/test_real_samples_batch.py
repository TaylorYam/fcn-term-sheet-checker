"""本機真實樣本測試（批量）：只在本機 data/ 有說明書與參考條件表時執行，CI 自動略過。

測試碼不含任何真實商品代號或數值；只驗證樣本內（IS）整批核對的結果（Issue #43）。
說明書與投資人須知一起核對（ADR 0007，Issue #124）：兩份都通過或放行的商品才回填。
參考條件表的回填欄位空白，回填值等於作業人員確認值的驗證在 OOS（`orders/OOS/`），不在這裡。
"""

from __future__ import annotations

import shutil
from collections import Counter
from pathlib import Path

import openpyxl
import pytest

from fcn_checker.batch import DocKind
from fcn_checker.saving import run_batch
from fcn_checker.schema import CheckStatus
from harness import CONFIG
from real_data import CHECK_FORMAT, TERM_SHEETS, pdfs, requires

PDFS = pdfs()  # 說明書與投資人須知
PROBLEMS = (CheckStatus.MISMATCH, CheckStatus.REVIEW_REQUIRED, CheckStatus.ERROR)

pytestmark = [
    pytest.mark.real_samples,
    requires(TERM_SHEETS / "*.pdf", CHECK_FORMAT),
]


@pytest.fixture(scope="module")
def saved(tmp_path_factory):
    """整批核對並儲存一次：(批量核對結果, 儲存收據)。"""
    tmp = tmp_path_factory.mktemp("real_batch")
    sheet = shutil.copy(CHECK_FORMAT, tmp / CHECK_FORMAT.name)  # 不在 data/ 產生檔案
    return run_batch(CONFIG, Path(sheet), PDFS, tmp / "reports", root=tmp)


def problems(item) -> Counter[tuple[str, str]]:
    return Counter((r.rule_id, r.reason_code) for r in item.report.results if r.status in PROBLEMS)


@pytest.fixture(scope="module")
def outcome(saved):
    return saved[0]


def term_sheets(outcome):
    return [i for i in outcome.items if i.kind == DocKind.TERM_SHEET]


def test_every_issuer_is_supported_and_only_pairs_that_both_pass_are_filled(saved):
    outcome, receipt = saved
    assert not [i.term_sheet.name for i in outcome.items if i.unsupported]
    assert len({i.issuer for i in outcome.items}) == 14
    for item in term_sheets(outcome):
        assert receipt.filled(item) == item.fills_sheet, item.term_sheet.name
        if item.fills_sheet:
            assert item.report.status == CheckStatus.PASS and item.partner.report.status == CheckStatus.PASS


def test_barc_rows_match_every_prefilled_field(outcome):
    paired = [
        i
        for i in term_sheets(outcome)
        if i.issuer == "BARC"
        and any(r.rule_id == "batch.pairing" and r.status == CheckStatus.PASS for r in i.report.results)
    ]
    assert len(paired) == 6, "參考條件表有 6 列 BARC"
    allowed = {
        ("standard.product_name", "value_mismatch"),  # 較早的中文名稱沒有「（不保本）」
    }
    for item in paired:
        assert set(problems(item)) <= allowed, item.term_sheet.name


def test_every_compare_date_on_the_sheet_matches_the_fill_rule(outcome):
    # D 型填 Non-Call 那期與最後一期、P 型從 Non-Call 那期起每期都填（Issue #69）
    assert not [
        (i.term_sheet.name, d.column) for i in outcome.items for d in i.report.backfill if d.action == "mismatch"
    ]


def test_passing_barc_rows_are_marked_filled(saved):
    outcome, receipt = saved
    passed = [i for i in term_sheets(outcome) if i.issuer == "BARC" and i.report.status == CheckStatus.PASS]
    assert len(passed) == 3, "中文名稱是新版的 3 份"
    assert all(receipt.filled(i) for i in passed), "同商品投資人須知也都通過"
    # 表上空白的回填欄位照說明書填入；表上已有的值都相同；期初定價 VWAP 的價格欄一律以說明書覆寫（Issue #122）
    assert all(
        d.action in ("match", "fill") or (d.column.startswith("UL_") and d.action == "overwrite")
        for i in passed
        for d in i.report.backfill
    )


def vwap_codes(path: Path) -> set[str]:
    return {code for code, row in rows_by_code(path).items() if row.get("期初定價") == "VWAP"}


def test_vwap_rows_skip_price_comparison_and_take_the_document_prices(outcome):
    vwap = [i for i in term_sheets(outcome) if i.product_code in vwap_codes(CHECK_FORMAT)]
    assert len(vwap) == 12, "表上有 12 列期初定價 VWAP"
    # 1 份 Nomura 另有其他差異（見 test_real_nomura.py），其餘都通過
    assert Counter(i.report.status for i in vwap) == Counter({CheckStatus.PASS: 11, CheckStatus.MISMATCH: 1})
    for item in vwap:
        assert not [r for r in item.report.results if r.rule_id == "field.underlying_prices"]
        prices = [d for d in item.report.backfill if d.column.startswith("UL_")]
        assert prices and all(d.action in ("match", "fill", "overwrite") for d in prices)


def rows_by_code(path: Path) -> dict[str, dict]:
    """TDCC Code → 該列；第一張工作表，表頭是前 10 列裡有 `TDCC Code` 的那一列。"""
    rows = list(openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True))
    header = next(r for r in rows[:10] if "TDCC Code" in r)
    records = (dict(zip(header, r, strict=False)) for r in rows[rows.index(header) + 1 :])
    return {str(d["TDCC Code"]): d for d in records if d.get("TDCC Code")}


def test_the_result_is_the_import_format(saved):
    """核對後「回填後」是匯入格式（第 2 列表頭，Issue #203），不是參考條件表的欄位。"""
    imp = CONFIG.reference_format.import_format
    out = [c.value for c in openpyxl.load_workbook(saved[1].output)["回填後"][imp.header_row]]
    assert out == list(imp.columns)


# ---------------------------------------------------------------- 投資人須知（Issue #124）


def investor_sheets(outcome, prefix: str = ""):
    return [i for i in outcome.items if i.kind == DocKind.IIS and i.term_sheet.name.startswith(prefix)]


def test_barc_investor_sheets_all_pass_with_their_term_sheets(outcome):
    sheets = investor_sheets(outcome, "029")
    assert len(sheets) == 6
    for item in sheets:
        assert item.report.status == CheckStatus.PASS, (item.term_sheet.name, problems(item))
        assert item.partner is not None and item.report.metadata["inputs"]["term_sheet"]["pages"] == 4


def test_hsbc_investor_sheets_only_differ_where_their_term_sheets_do(outcome):
    sheets = investor_sheets(outcome, "325")
    assert len(sheets) == 7
    # 第七節「受託投資本商品」：交易日在審查標準 wording.exempt_until 以前，特例通過（Issue #247）
    wording = [r for i in sheets for r in i.report.results if r.rule_id == "standard.forbidden_wording"]
    assert len(wording) == 7 and all(r.status == CheckStatus.PASS and "特例通過" in r.tolerance for r in wording)
    # 同商品說明書有相同的標的代碼差異；依使用者確認仍逐字比對尾碼（test_real_hsbc.py，Issue #310）
    assert sum((problems(i) for i in sheets), Counter()) == Counter({("field.underlyings", "value_mismatch"): 1})


def test_every_local_investor_sheet_has_an_iis_template(outcome):
    """樣本內 14 家上手的投資人須知都有範本；未支援上手的投資人須知另見合成測試 test_iis.py。"""
    sheets = investor_sheets(outcome)
    assert len(sheets) == len(term_sheets(outcome)) and not [i for i in sheets if i.unsupported]
    # 只支援新版範本的上手，舊版辨識失敗、轉人工覆核：UBS 舊版 1 份（test_real_ubs.py）
    unknown = [i for i in sheets if not i.report.template]
    assert [i.issuer for i in unknown] == ["UBS"]
