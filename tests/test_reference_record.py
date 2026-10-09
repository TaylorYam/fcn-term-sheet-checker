"""參考條件表的一列（`OrderRecord`）可以不經 Excel 直接以標準欄位值建立，規則與回填規則只收 `Context`。

測試切點：`OrderRecord.offline` → `Context` → 欄位核對表（`rules/reference.FIELD_CHECKS`）與回填規則（`backfill.*`）。
離線建的列有格式設定的每一欄，儲存格依 `[columns]` 的順序從 A 欄起編號。
"""

from __future__ import annotations

import datetime as dt
from decimal import Decimal

from openpyxl.utils import get_column_letter

from fcn_checker import backfill
from fcn_checker.backfill import BackfillAction
from fcn_checker.orders.reference import OrderRecord
from fcn_checker.rules import reference
from fcn_checker.rules.kit import Context
from fcn_checker.schema import ParsedField
from fcn_checker.standard_fields import lookup
from harness import CONFIG, MISMATCH, PASS

FMT = CONFIG.reference_format


def cell_of(header: str, row: int) -> str:
    """離線建的列裡某 Excel 欄名的儲存格：依格式設定 `[columns]` 的順序從 A 欄起編號。"""
    return f"{get_column_letter(list(FMT.columns).index(header) + 1)}{row}"


class _TermSheet:
    """只交出指定標準欄位的讀出結果。"""

    full_text = None

    def __init__(self, **values):
        self.fields = {name: ParsedField.present(name, value, []) for name, value in values.items()}

    def f(self, name: str) -> ParsedField:
        return lookup(self.fields, name)


def _context(ts: _TermSheet, **values) -> Context:
    record = OrderRecord.offline(FMT, {"product_code": "123456789012", **values}, row=7)
    return Context(ts, record, CONFIG.review_standard, FMT, "BARC")


def test_a_row_built_from_a_dict_drives_the_field_checks():
    ts = _TermSheet(trade_date=dt.date(2030, 1, 10), strike_pct=Decimal("80.00"))
    ctx = _context(ts, trade_date=dt.date(2030, 1, 10), strike_pct=Decimal("80"))

    trade = reference.FIELD_CHECKS["trade_date"].check(ctx)
    assert (trade.status, trade.item.name, trade.item.columns) == (PASS, "交易日", ("交易日",))
    assert trade.order_source == [f"{FMT.sheet}!{cell_of('交易日', 7)}"]

    strike = reference.FIELD_CHECKS["strike_pct"].check(ctx)
    assert (strike.status, strike.item.name, strike.tolerance) == (PASS, "K(%)", reference.PCT_TOLERANCE)

    ctx = _context(ts, trade_date=dt.date(2030, 1, 11))
    assert reference.FIELD_CHECKS["trade_date"].check(ctx).status == MISMATCH


def test_backfill_rules_take_only_the_context():
    ts = _TermSheet(isin="XS1999900002")
    r, [d] = backfill.isin(_context(ts))
    assert r.status == PASS
    assert (d.column, d.cell, d.sheet_value, d.expected) == ("ISIN Code", cell_of("ISIN Code", 7), None, "XS1999900002")
    assert d.action == BackfillAction.FILL

    r, [d] = backfill.isin(_context(ts, isin="XS0000000000"))
    assert (r.status, d.action) == (MISMATCH, BackfillAction.MISMATCH)


def test_typed_reads_build_the_underlying_and_autocall_keys_in_one_place():
    record = OrderRecord.offline(
        FMT,
        {"underlying_2": "AAPL UW", "underlying_2_ko_price": Decimal("123.4"), "autocall_date_3": dt.date(2030, 4, 7)},
    )
    assert (record.underlying(2).value, record.underlying(2).column) == ("AAPL UW", "UL_2")
    assert (record.price(2, "ko").value, record.price(2, "ko").column) == (Decimal("123.4"), "UL_2_KO價")
    assert (record.autocall_date(3).value, record.autocall_date(3).column) == (dt.date(2030, 4, 7), "比價日_3")
    assert record.underlying(1).value is None and record.price(1, "ki").value is None, "沒給值的欄位是空白格"
    assert (record.header("issue_date"), record.cell("issue_date")) == ("發行日", cell_of("發行日", 4))
    assert record.underlying_slots == 5
