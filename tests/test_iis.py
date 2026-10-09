"""投資人須知（IIS）核對：說明書與投資人須知成對核對，三方一致才回填（ADR 0007、docs/rules/iis-check-rules.md）。

測試切點：批量入口（預覽＋核對＋儲存）與 PANEL 工作階段；合成 PDF 由各上手合成器產生，數值皆虛構。
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

import pytest

import hsbc_synth
from fcn_checker.batch import DocKind
from fcn_checker.ingestion import IngestionError
from fcn_checker.issuers import BARC
from fcn_checker.panel_workflow import PanelSession
from fcn_checker.saving import run_batch
from fcn_checker.schema import CheckStatus
from harness import CONFIG, REVIEW_STANDARD, ROOT, check_all, cli_root, iis_path, with_iis
from reference_synth import build_reference_sheet
from synth import Spec, barc_adapter, build_iis_pdf, build_pdf, reference_row

PASS, MISMATCH, REVIEW = CheckStatus.PASS, CheckStatus.MISMATCH, CheckStatus.REVIEW_REQUIRED
NOW = dt.datetime(2030, 2, 3, 4, 5, 6)


def pair(tmp_path: Path, spec: Spec | None = None, *, ts_spec: Spec | None = None, **iis) -> tuple[Path, Path]:
    """同商品的說明書與投資人須知；`iis` 交給投資人須知合成器（replace、pages）製造錯誤。"""
    spec = spec or Spec()
    ts = build_pdf(tmp_path / f"{spec.product_code}_TS.pdf", ts_spec or spec, iis=False)
    return ts, build_iis_pdf(iis_path(ts), spec, **iis)


def run(tmp_path: Path, pdfs: list[Path], rows: list[dict], config=CONFIG):
    sheet = build_reference_sheet(tmp_path / "FCN參考條件.xlsx", rows)
    return run_batch(config, sheet, pdfs, tmp_path / "reports", root=tmp_path, now=NOW)


def problems(item) -> list[str]:
    return list(item.problem_messages)


def error_names(receipt) -> list[str]:
    import openpyxl

    ws = openpyxl.load_workbook(receipt.output)["錯誤清單"]
    return [r[1] for r in ws.iter_rows(min_row=2, values_only=True)]


def filled_codes(receipt) -> list[str]:
    import openpyxl

    ws = openpyxl.load_workbook(receipt.output)["回填後"]
    headers = [c.value for c in ws[3]]
    return [dict(zip(headers, r, strict=False))["TDCC Code"] for r in ws.iter_rows(min_row=4, values_only=True)]


# ---------------------------------------------------------------- 一致時兩份都通過，說明書回填


def test_consistent_term_sheet_and_iis_pass_and_the_term_sheet_is_filled(tmp_path):
    spec = Spec()
    ts, iis = pair(tmp_path, spec)
    outcome, receipt = run(tmp_path, [ts, iis], [reference_row(spec)])

    item, sheet = outcome.items
    assert (item.kind, sheet.kind) == (DocKind.TERM_SHEET, DocKind.IIS)
    assert item.partner is sheet and sheet.partner is item
    assert sheet.report.status == PASS, problems(sheet)
    assert sheet.report.template == "barc-zh-iis" and sheet.product_code == spec.product_code
    rule_ids = {r.rule_id for r in sheet.report.results}
    assert {
        "iis.pages",
        "iis.product_code",
        "field.currency",
        "field.underlyings",
        "iis.underlying_prices",
        "iis.monthly_coupon",
        "iis.isin",
        "iis.name_zh",
        "iis.issue_date",
        "standard.fixed_warning",
        "standard.distributor",
        "standard.fees",
    } <= rule_ids
    assert not any(r.rule_id.startswith("backfill.") for r in sheet.report.results), "投資人須知不回填"
    assert item.fills_sheet and receipt.filled(item) and not receipt.filled(sheet)
    assert filled_codes(receipt) == [spec.product_code] and error_names(receipt) == []
    assert sheet.release_problem == "已經通過，不需要人工放行"


def test_hsbc_term_sheet_and_iis_pass_together(tmp_path):
    s = hsbc_synth.Spec()
    ts = hsbc_synth.build_pdf(tmp_path / f"{s.code}_TS.pdf", s)
    sheet = hsbc_synth.build_inquiry(tmp_path / "order.xlsx", s)
    outcome, receipt = run_batch(CONFIG, sheet, with_iis([ts]), tmp_path / "reports", root=tmp_path, now=NOW)

    item, iis = outcome.items
    assert iis.report.status == PASS, problems(iis)
    assert iis.report.template == "hsbc-zh-iis"
    rule_ids = {r.rule_id for r in iis.report.results}
    assert {"iis.page_totals", "doc.print_date", "iis.underlying_names", "standard.issuer_name"} <= rule_ids
    assert "iis.product_code" not in rule_ids and "iis.isin" not in rule_ids, "HSBC 範本沒有商品代號與 ISIN，不核對"
    assert receipt.filled(item)


# ---------------------------------------------------------------- 頁數


@pytest.mark.parametrize("pages", [3, 5])
def test_iis_must_have_exactly_four_pages(tmp_path, pages):
    spec = Spec()
    ts, iis = pair(tmp_path, spec, pages=pages)
    outcome, receipt = run(tmp_path, [ts, iis], [reference_row(spec)])

    item, sheet = outcome.items
    assert sheet.report.status == MISMATCH
    assert problems(sheet) == [f"投資人須知頁數：預期 4 頁／實際 {pages} 頁"]
    assert item.report.status == PASS and not receipt.filled(item), "投資人須知沒過，說明書不回填"
    assert item.not_filled_reason == "同商品的投資人須知尚未通過或人工放行，不回填。"
    assert error_names(receipt) == [iis.name] and filled_codes(receipt) == []


def test_hsbc_page_header_total_must_match_the_pages(tmp_path):
    s = hsbc_synth.Spec()
    ts = hsbc_synth.build_pdf(tmp_path / f"{s.code}_TS.pdf", s, iis=False)
    iis = hsbc_synth.build_iis_pdf(iis_path(ts), s, page_total=5)
    sheet = hsbc_synth.build_inquiry(tmp_path / "order.xlsx", s)
    _, item = check_all(sheet, [ts, iis]).items
    assert problems(item) == ["頁首總頁數：頁首寫「共 5 頁」，實際 4 頁"]


# ---------------------------------------------------------------- 成對


def test_term_sheet_without_iis_requires_review_and_cannot_be_released(tmp_path):
    spec = Spec()
    ts = build_pdf(tmp_path / f"{spec.product_code}_TS.pdf", spec, iis=False)
    outcome, receipt = run(tmp_path, [ts], [reference_row(spec)])

    (item,) = outcome.items
    assert item.status_label == "這批缺投資人須知" and item.report.status == REVIEW
    assert any(r.rule_id.startswith("field.") for r in item.report.results), "這份的規則照常執行"
    assert "這批缺同商品的投資人須知" in item.release_problem
    with pytest.raises(IngestionError, match="投資人須知"):
        outcome.release(item)
    assert not receipt.filled(item) and error_names(receipt) == [ts.name]


def test_iis_without_term_sheet_requires_review_and_compares_nothing_to_the_term_sheet(tmp_path):
    spec = Spec()
    _, iis = pair(tmp_path, spec)
    (item,) = check_all(build_reference_sheet(tmp_path / "ref.xlsx", [reference_row(spec)]), [iis]).items
    assert item.status_label == "這批缺說明書"
    isin = next(r for r in item.report.results if r.rule_id == "iis.isin")
    assert (isin.status, isin.reason_code) == (REVIEW, "term_sheet_unavailable")
    assert "這批缺同商品的說明書" in item.release_problem


def test_iis_of_an_issuer_without_an_iis_template_is_unsupported(tmp_path):
    spec = Spec()
    ts, iis = pair(tmp_path, spec)
    config = CONFIG.with_registry((barc_adapter(iis=None),))
    item, sheet = check_all(
        build_reference_sheet(tmp_path / "ref.xlsx", [reference_row(spec)]), [ts, iis], config
    ).items
    assert sheet.unsupported and sheet.status_label == "未支援上手"
    assert "還沒有投資人須知範本" in problems(sheet)[0]
    assert item.status_label == "通過", "投資人須知在這批裡（只是未支援），說明書不算缺投資人須知"
    assert item.partner is sheet and not item.fills_sheet
    assert item.not_filled_reason == "同商品的投資人須知尚未通過或人工放行，不回填。"


def test_term_sheet_named_as_iis_is_not_a_known_iis_template(tmp_path):
    spec = Spec()
    fake = build_pdf(tmp_path / f"{spec.product_code}_IIS.pdf", spec, iis=False)  # 其實是說明書
    (item,) = check_all(build_reference_sheet(tmp_path / "ref.xlsx", [reference_row(spec)]), [fake]).items
    r = next(r for r in item.report.results if r.rule_id == "template.detect")
    assert (r.status, r.reason_code) == (REVIEW, "template_unknown")
    assert r.message.startswith("不是已支援的投資人須知範本")


# ---------------------------------------------------------------- 三方一致：參考條件表、說明書


@pytest.mark.parametrize(
    ("replace", "message"),
    [
        ({"2030 年7 月11 日": "2030 年7 月12 日"}, "到期日對不起來：參考條件表 2030-07-11／投資人須知 2030-07-12"),
        ({"86.4150": "86.4151"}, "UL_1 執行價對不起來：參考條件表 86.4150／投資人須知 86.4151"),
        ({"每單位商品面額為10,000": "每單位商品面額為20,000"}, "單位面額對不起來：參考條件表 10000／投資人須知 20000"),
        ({"商品年期：6 個月": "商品年期：7 個月"}, "天期(月)對不起來：參考條件表 6／投資人須知 7"),
        ({"ZZA UN Equity": "ZZB UN Equity"}, "標的對不起來：參考條件表 ZZA UN、ZQH UW、DWE UW／投資人須知 ZZB UN"),
        ({"即年利率為12.00%": "即年利率為12.50%"}, "Coupon p.a. (%)對不起來：參考條件表 12.00／投資人須知 12.50"),
        (
            {"每月之配息率（為1.0000%": "每月之配息率（為1.0100%"},
            "月配息率：須等於年利率 ÷ 12（參考條件表 1.0000／投資人須知 1.0100）",
        ),
    ],
)
def test_iis_values_are_compared_with_the_reference_sheet(tmp_path, replace, message):
    spec = Spec()
    ts, iis = pair(tmp_path, spec, replace=replace)
    item, sheet = check_all(build_reference_sheet(tmp_path / "ref.xlsx", [reference_row(spec)]), [ts, iis]).items
    assert item.report.status == PASS
    assert sheet.report.status == MISMATCH
    assert [m for m in problems(sheet) if m.startswith(message)], problems(sheet)


@pytest.mark.parametrize(
    ("replace", "message"),
    [
        ({"2030 年1 月14 日": "2030 年1 月15 日"}, "發行日對不起來：說明書 2030-01-14／投資人須知 2030-01-15"),
        ({"ISIN:XS0000000000": "ISIN:XS0000000001"}, "ISIN對不起來：說明書 XS0000000000／投資人須知 XS0000000001"),
        ({"最低申購金額為10,000": "最低申購金額為20,000"}, "最低申購金額對不起來：說明書 10000／投資人須知 20000"),
        ({"英商巴克萊銀行6個月": "英商巴克萊銀行7個月"}, "中文商品名稱對不起來：說明書 "),
        ({"丁戊電子公司": "丁戊電機公司"}, "標的中文名稱對不起來：說明書 "),
    ],
)
def test_iis_values_missing_from_the_reference_sheet_are_compared_with_the_term_sheet(tmp_path, replace, message):
    spec = Spec()
    ts, iis = pair(tmp_path, spec, replace=replace)
    _, sheet = check_all(build_reference_sheet(tmp_path / "ref.xlsx", [reference_row(spec)]), [ts, iis]).items
    assert [m for m in problems(sheet) if m.startswith(message)], problems(sheet)


def test_a_term_sheet_error_is_reported_only_on_the_term_sheet(tmp_path):
    spec = Spec()
    ts, iis = pair(tmp_path, spec, ts_spec=spec.with_(maturity_date=dt.date(2030, 7, 12)))
    item, sheet = check_all(build_reference_sheet(tmp_path / "ref.xlsx", [reference_row(spec)]), [ts, iis]).items
    assert any(m.startswith("到期日對不起來：參考條件表 2030-07-11／說明書 2030-07-12") for m in problems(item))
    assert sheet.report.status == PASS, "投資人須知和參考條件表一致，不重複報說明書的錯"


def test_iis_cover_product_code_must_match_the_file_name(tmp_path):
    spec = Spec()
    ts, iis = pair(
        tmp_path, spec, replace={"受託或銷售機構商品代號:029199990001": "受託或銷售機構商品代號:029199990002"}
    )
    _, sheet = check_all(build_reference_sheet(tmp_path / "ref.xlsx", [reference_row(spec)]), [ts, iis]).items
    assert problems(sheet) == [
        "封面受託或銷售機構商品代號：封面「受託或銷售機構商品代號」與檔名前 12 碼不同，可能放錯檔案"
        "（預期 029199990001／投資人須知 029199990002）"
    ]


def test_vwap_iis_prices_are_compared_with_the_term_sheet(tmp_path):
    spec = Spec()
    blank = {f"UL_{i}_{c}": None for i in range(1, 4) for c in ("進場價", "執行價", "下限價", "KO價")}
    rows = [reference_row(spec, **{"期初定價": "VWAP", **blank})]
    ts, iis = pair(tmp_path, spec)
    item, sheet = check_all(build_reference_sheet(tmp_path / "ok.xlsx", rows), [ts, iis]).items
    assert sheet.report.status == PASS, problems(sheet)

    (tmp_path / "bad").mkdir()
    ts, iis = pair(tmp_path / "bad", spec, replace={"86.4150": "86.4151"})
    _, sheet = check_all(build_reference_sheet(tmp_path / "bad" / "ref.xlsx", rows), [ts, iis]).items
    assert problems(sheet) == ["UL_1 執行價對不起來：說明書 86.4150／投資人須知 86.4151"]


# ---------------------------------------------------------------- 審查標準


@pytest.mark.parametrize(
    ("replace", "message"),
    [
        (
            {"受託投資或受託買賣之投資標的": "受託投資之投資標的"},
            "禁用語「受託投資」：允許片語以外出現「受託投資」1 處（SOP：須改為「受託買賣」）",
        ),
        (
            {"而非由玉山綜合證券股份有限公司保證": "而非由玉山證券股份有限公司保證"},
            "受託機構名稱（警語 5）：警語第 5 點的受託或銷售機構名稱與審查標準不同",
        ),
        (
            {"為【證券交易所交易證券】": "為【證券交易所之交易證券】"},
            "固定風險警語逐字相符 0 次，應為 1 次",
        ),
        ({"3.本商品風險程度：RR4": "3.本商品風險程度：【RR5】"}, "風險等級："),
    ],
)
def test_iis_review_standard_items(tmp_path, replace, message):
    spec = Spec()
    ts, iis = pair(tmp_path, spec, replace=replace)
    _, sheet = check_all(build_reference_sheet(tmp_path / "ref.xlsx", [reference_row(spec)]), [ts, iis]).items
    assert [m for m in problems(sheet) if m.startswith(message)], problems(sheet)


def test_iis_fee_differs_from_the_review_standard(tmp_path):
    spec = Spec()
    ts, iis = pair(tmp_path, spec.with_(fees={"分銷費用": "0%~3%"}))
    _, sheet = check_all(build_reference_sheet(tmp_path / "ref.xlsx", [reference_row(spec)]), [ts, iis]).items
    assert [m for m in problems(sheet) if m.startswith("分銷費用：")], problems(sheet)


def test_hsbc_iis_forbidden_wording_and_print_date(tmp_path):
    s = hsbc_synth.Spec()
    ts = hsbc_synth.build_pdf(tmp_path / f"{s.code}_TS.pdf", s, iis=False)
    replace = {
        "四、本商品雖經": "四、受託或銷售機構將為投資人受託投資本商品。本商品雖經",
        "2030 年1 月7 日": "2030 年1 月9 日",
    }
    iis = hsbc_synth.build_iis_pdf(iis_path(ts), s, replace=replace)
    sheet = hsbc_synth.build_inquiry(tmp_path / "order.xlsx", s)
    _, item = check_all(sheet, [ts, iis]).items
    messages = problems(item)
    assert any(m.startswith("禁用語「受託投資」：允許片語以外出現「受託投資」1 處") for m in messages), messages
    assert any(m.startswith("刊印日期為交易日 +2 天") for m in messages), messages


# ---------------------------------------------------------------- 人工放行


def test_each_pdf_is_released_on_its_own_and_both_are_needed_to_fill(tmp_path):
    spec = Spec()
    ts, iis = pair(tmp_path, spec, pages=5)
    session = PanelSession(REVIEW_STANDARD, ROOT / "config", install_root=tmp_path)
    session.select(build_reference_sheet(tmp_path / "ref.xlsx", [reference_row(spec)]), [ts, iis])
    session.load_preview()
    item, sheet = session.start_check().batch.items
    assert session.release_state(sheet) == (True, "")
    assert session.release_state(item) == (False, ""), "通過的說明書不能放行，也不必說明"

    session.release(sheet)
    assert sheet.status_label == "人工放行（原：不一致）" and item.fills_sheet
    receipt = session.save(tmp_path / "reports", now=NOW)
    assert receipt.filled(item) and filled_codes(receipt) == [spec.product_code] and error_names(receipt) == []


def test_pdf_preview_marks_the_document_kind(tmp_path):
    spec = Spec()
    ts, iis = pair(tmp_path, spec)
    session = PanelSession(REVIEW_STANDARD, ROOT / "config", install_root=tmp_path)
    session.select(build_reference_sheet(tmp_path / "ref.xlsx", [reference_row(spec)]), [iis, ts])
    first, second = session.load_preview().rows
    assert (first.kind, second.kind) == (DocKind.IIS, DocKind.TERM_SHEET)
    assert first.reference_row == second.reference_row == 4 and first.problem == second.problem == ""


def test_iis_records_its_kind_and_partner(tmp_path):
    import json

    spec = Spec()
    ts, iis = pair(tmp_path, spec)
    _, receipt = run(tmp_path, [ts, iis], [reference_row(spec)])
    items = json.loads(receipt.record.read_text(encoding="utf-8"))["items"]
    assert [(i["pdf"], i["document"], i["partner"], i["filled"]) for i in items] == [
        (ts.name, "說明書", iis.name, True),
        (iis.name, "投資人須知", ts.name, False),
    ]
    assert items[1]["template"] == BARC.iis.template_id


@pytest.mark.parametrize(
    ("spec", "checked", "skipped"),
    [
        (Spec(memory=False), {"field.ko_pct"}, {"field.ki_pct"}),
        (Spec(memory=True), set(), {"field.ko_pct", "field.ki_pct"}),  # 記憶式 KO 欄頭沒有單一百分比
        (Spec(memory=False, ki="AM"), {"field.ko_pct", "field.ki_pct"}, set()),
    ],
    ids=["non_memory", "memory", "knock_in"],
)
def test_barc_price_table_variants(tmp_path, spec, checked, skipped):
    ts, iis = pair(tmp_path, spec)
    _, sheet = check_all(build_reference_sheet(tmp_path / "ref.xlsx", [reference_row(spec)]), [ts, iis]).items
    assert sheet.report.status == PASS, problems(sheet)
    rule_ids = {r.rule_id for r in sheet.report.results}
    assert checked <= rule_ids and not (skipped & rule_ids)
    prices = [r for r in sheet.report.results if r.rule_id == "iis.underlying_prices"]
    assert len(prices) == len(spec.underlyings) * (4 if spec.ki != "none" else 3)


def test_barc_knock_in_price_on_the_iis_is_compared_with_the_sheet(tmp_path):
    spec = Spec(memory=False, ki="AM")
    ts, iis = pair(tmp_path, spec, replace={"74.0700": "74.0701"})  # 123.4500 × 60%
    _, sheet = check_all(build_reference_sheet(tmp_path / "ref.xlsx", [reference_row(spec)]), [ts, iis]).items
    assert problems(sheet) == ["UL_1 下限價對不起來：參考條件表 74.0700／投資人須知 74.0701"]


def test_unreadable_iis_price_table_requires_review_naming_the_iis(tmp_path):
    spec = Spec()
    ts, iis = pair(tmp_path, spec, replace={"最初價格": "期初價格"})
    _, sheet = check_all(build_reference_sheet(tmp_path / "ref.xlsx", [reference_row(spec)]), [ts, iis]).items
    assert sheet.report.status == REVIEW
    assert "價格表：投資人須知的值無法辨識：價格表欄頭找不到或有多個「最初價格」" in problems(sheet)


# ---------------------------------------------------------------- Issue #124 審查後補充


def test_term_sheet_file_name_must_carry_its_cover_product_code(tmp_path):
    spec, other = Spec(), Spec(product_code="029199990002")
    ts = build_pdf(tmp_path / f"{other.product_code}_TS.pdf", spec, iis=False)  # 檔名是 other，封面是 spec
    iis = build_iis_pdf(tmp_path / f"{other.product_code}_IIS.pdf", other)
    item, sheet = check_all(
        build_reference_sheet(tmp_path / "ref.xlsx", [reference_row(spec), reference_row(other)]), [ts, iis]
    ).items
    assert item.status_label == "檔名商品代號不符"
    assert "可能放錯檔案" in item.release_problem
    assert sheet.partner is item and sheet.report.status != PASS, "說明書沒配對成功，投資人須知無法和它比"
    isin = next(r for r in sheet.report.results if r.rule_id == "iis.isin")
    assert isin.reason_code == "term_sheet_unavailable"


def test_iis_failing_its_own_pairing_does_not_make_the_term_sheet_miss_it(tmp_path):
    spec = Spec()
    ts = build_pdf(tmp_path / f"{spec.product_code}_TS.pdf", spec, iis=False)
    fake = build_pdf(tmp_path / f"{spec.product_code}_IIS.pdf", spec, iis=False)  # 範本不符的投資人須知
    item, sheet = check_all(build_reference_sheet(tmp_path / "ref.xlsx", [reference_row(spec)]), [ts, fake]).items
    assert sheet.status_label == "需人工覆核"
    assert item.status_label == "通過" and item.release_problem == "已經通過，不需要人工放行"
    assert not item.fills_sheet


def test_issue_date_on_the_sheet_is_compared_with_the_sheet_and_a_term_sheet_error_is_not_repeated(tmp_path):
    spec = Spec()
    ts, iis = pair(tmp_path, spec, ts_spec=spec.with_(issue_date=dt.date(2030, 1, 15)))
    rows = [reference_row(spec, 發行日=dt.datetime(2030, 1, 14))]
    item, sheet = check_all(build_reference_sheet(tmp_path / "ref.xlsx", rows), [ts, iis]).items
    assert any(m.startswith("發行日對不起來：參考條件表 2030-01-14／說明書 2030-01-15") for m in problems(item))
    assert sheet.report.status == PASS, problems(sheet)

    (tmp_path / "iis").mkdir()
    ts, iis = pair(tmp_path / "iis", spec, replace={"2030 年1 月14 日": "2030 年1 月16 日"})
    _, sheet = check_all(build_reference_sheet(tmp_path / "iis" / "ref.xlsx", rows), [ts, iis]).items
    assert problems(sheet) == ["發行日對不起來：參考條件表 2030-01-14／投資人須知 2030-01-16"]


def test_barc_risk_level_in_the_product_summary_is_checked(tmp_path):
    spec = Spec()
    ts, iis = pair(tmp_path, spec, replace={"3.本商品風險程度：RR4": "3.本商品風險程度：RR5"})
    _, sheet = check_all(build_reference_sheet(tmp_path / "ref.xlsx", [reference_row(spec)]), [ts, iis]).items
    assert problems(sheet) == ["風險等級（商品簡介）：兩邊的值不同（審查標準 RR4／投資人須知 RR5）"]


def test_term_sheet_error_messages_never_show_field_codes(tmp_path):
    spec = Spec()
    _, iis = pair(tmp_path, spec)
    (item,) = check_all(build_reference_sheet(tmp_path / "ref.xlsx", [reference_row(spec)]), [iis]).items
    messages = " ".join(problems(item))
    assert "沒有可比對的同商品說明書" in messages
    assert not any(code in messages for code in ("name_zh", "underlying_prices", "min_amounts", "isin"))


def test_cli_exit_code_and_not_filled_reason_when_the_iis_fails(tmp_path, monkeypatch, capsys):
    from fcn_checker.cli import main

    spec = Spec()
    ts, iis = pair(tmp_path, spec, pages=5)
    sheet = build_reference_sheet(tmp_path / "ref.xlsx", [reference_row(spec)])
    cli_root(tmp_path, monkeypatch)
    assert main([str(sheet), str(ts), str(iis), "--out", str(tmp_path / "reports")]) == 1
    out = capsys.readouterr().out
    assert f"通過  {ts.name}  同商品的投資人須知尚未通過或人工放行，不回填。" in out
    assert f"不一致  {iis.name}" in out


# ---------------------------------------------------------------- 文件身分（Issue #142）


def test_iis_rule_message_naming_the_term_sheet_is_kept_as_written(tmp_path):
    """投資人須知規則描述同商品說明書的值、訊息以「說明書」開頭時，不被改寫成「投資人須知」。"""
    import dataclasses

    from fcn_checker.rules.kit import result
    from fcn_checker.schema import Item

    message = "說明書第 3 期比價日 2030-04-09，投資人須知要和它相同"
    rule = lambda ctx: [  # noqa: E731
        result("iis.fake", "fake", REVIEW, reason="value_mismatch", message=message, item=Item.term_sheet("比價日"))
    ]
    barc = dataclasses.replace(BARC, iis=dataclasses.replace(BARC.iis, rules=rule))
    spec = Spec()
    ts, iis = pair(tmp_path, spec)
    sheet_ref = build_reference_sheet(tmp_path / "ref.xlsx", [reference_row(spec)])
    _, sheet = check_all(sheet_ref, [ts, iis], CONFIG.with_registry((barc,))).items

    (r,) = [r for r in sheet.report.results if r.rule_id == "iis.fake"]
    assert r.message == message and r.document == DocKind.IIS
    assert problems(sheet) == [f"比價日：{message}"]


def test_every_iis_result_records_its_document(tmp_path):
    spec = Spec()
    ts, iis = pair(tmp_path, spec, pages=5)
    item, sheet = check_all(build_reference_sheet(tmp_path / "ref.xlsx", [reference_row(spec)]), [ts, iis]).items

    assert {r.document for r in sheet.report.results} == {DocKind.IIS}
    assert {r.document for r in item.report.results} == {DocKind.TERM_SHEET}


def test_iis_review_standard_messages_name_the_iis_when_fields_are_unreadable():
    """審查標準規則（刊印日期、費用）讀不到投資人須知欄位時，說明以「投資人須知」開頭，不寫成說明書。"""
    from fcn_checker.investor_sheet import IisSheet
    from fcn_checker.orders.reference import OrderRecord
    from fcn_checker.parsers.layout import TextIndex
    from fcn_checker.rules.kit import Context
    from fcn_checker.rules.review_standard import iis_review_standard_rules
    from fcn_checker.schema import ParsedField
    from fcn_checker.standard_fields import fee_field

    std = CONFIG.review_standard
    provided = frozenset({"print_dates", *(fee_field(label) for label in std.fees)})
    sheet = IisSheet({}, TextIndex([]), provided)
    order = OrderRecord.offline(CONFIG.reference_format, {"product_code": "029199990001"})
    ctx = Context(sheet, order, std, CONFIG.reference_format, "BARC", document=DocKind.IIS)

    results = iis_review_standard_rules(ctx, trade=ParsedField.present("trade_date", dt.date(2030, 1, 7), []))

    unreadable = [r for r in results if r.reason_code == "document_missing"]
    assert {r.field for r in unreadable} >= {"print_dates", *std.fees}
    assert all(r.message.startswith("投資人須知") for r in unreadable), [r.message for r in unreadable]


def test_unexpected_error_while_reading_the_iis_names_the_iis(tmp_path):
    import dataclasses

    def broken(lines):
        raise ValueError("讀出失敗")

    barc = dataclasses.replace(BARC, iis=dataclasses.replace(BARC.iis, read=broken))
    spec = Spec()
    ts, iis = pair(tmp_path, spec)
    sheet_ref = build_reference_sheet(tmp_path / "ref.xlsx", [reference_row(spec)])
    _, sheet = check_all(sheet_ref, [ts, iis], CONFIG.with_registry((barc,))).items

    assert problems(sheet) == ["投資人須知：ValueError: 讀出失敗"]
