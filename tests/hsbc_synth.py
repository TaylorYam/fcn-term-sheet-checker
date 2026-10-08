"""HSBC 合成 PDF／整理表；商品、標的與價格皆虛構。"""

from __future__ import annotations

import datetime as dt
import tomllib
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

import openpyxl

from harness import REVIEW_STANDARD
from pdf_writer import FONT, PdfWriter, zh_date
from reference_synth import REFERENCE_FORMAT, issuer_value

ORDER_FORMAT = REFERENCE_FORMAT
STD = tomllib.loads(REVIEW_STANDARD.read_text(encoding="utf-8"))


@dataclass
class Spec:
    obs: str = "D"
    memory: bool = True
    ki: str = "AM"
    count: int = 2
    partial_coupon: bool = False
    annual: Decimal = Decimal("12")
    replacements: dict[str, str] = field(default_factory=dict)
    scenario_replacements: dict[str, str] = field(default_factory=dict)
    code: str = "325199990001"

    @property
    def ends(self):
        return [dt.date(2030, m, 7) for m in range(2, 8)]

    @property
    def payments(self):
        return [d + dt.timedelta(days=3) for d in self.ends]

    @property
    def name(self):
        return STD["product_name"]["hsbc"]["zh"].format(
            tenor=6, ccy_zh="美元", memory_zh="記憶式" if self.memory else ""
        )

    @property
    def en(self):
        return STD["product_name"]["hsbc"]["en"].format(
            maxi_en="Maxi " if self.count > 1 else "",
            daily_en="Daily " if self.obs == "D" else "",
            memory_en="Memory " if self.memory else "",
        )


def build_pdf(path: Path, s: Spec):
    w = PdfWriter()

    def line(t, x=110):
        t = s.replacements.get(t, t)
        if t:
            # Wrap between text runs, never in the middle of a printed number.
            while t:
                end = min(48, len(t))
                while end < len(t) and t[end - 1] in "0123456789,." and t[end] in "0123456789,.":
                    end += 1
                w.line(x, t[:end])
                t = t[end:]

    def chapter(n, name):
        w.new_page()
        w.line(60, f"第{n}章 {name}")

    def article(n, text):
        w.line(60, f"{n}.")
        line(text)

    def sub(n):
        w.line(80, f"({n})")

    def table(scenario=False):
        for t in ["期初股價", "執行價(即期初股價的70%)", "自動提前到期價格(即期初股價的100%)"] + (
            ["觸及不保本價格(即期初股價的60%)"] if s.ki != "none" else []
        ):
            line(s.scenario_replacements.get(t, t) if scenario else t)
        for i in range(s.count):
            for t in [f"ZZ{i + 1} UW", f"虛構標的{i + 1}", "USD", "NASDAQ", "100.0000", "70.0000", "100.0000"] + (
                ["60.0000"] if s.ki != "none" else []
            ):
                line(s.scenario_replacements.get(t, t) if scenario else t)

    short = s.name.replace("（以下簡稱「本商品」）", "")
    line(short, 60)
    line(s.en, 60)
    line("中文產品說明書(最終版)")
    line(f"商品代號/商品中文名稱：{s.code}/{s.name}")
    line(f"商品英文名稱：{s.en}")
    line("商品種類：股權連結商品")
    line("計價幣別：美元")
    line("發行機構：" + STD["issuer_name"]["hsbc"])
    line("電話：+852-0000-0000")
    line("[受託或銷售機構]審查通過之日期：2026年6月11日")
    line("(參考性審閱版)內容，刊印日期：2030年1月7日")
    line("(最終版)刊印日期：2030年1月7日")
    line("受託或銷售機構之名稱、電話及地址：" + STD["distributor"]["name"])
    line("電話：+886-2-5556-1313")
    line(STD["distributor"]["address"])
    line("公會審查")
    warning = STD["risk"]["fixed_warning_by_issuer"]["hsbc"]
    line(warning)
    chapter("一", "商品基本資料")
    for n in range(1, 32):
        article(
            n,
            {
                1: "商品名稱：" + short + s.en,
                2: warning,
                5: "計價幣別：美元",
                6: "每單位面額：美元10,000元",
                7: "最低交易金額：美元10,000元",
                10: "發行價格：100%",
                11: "主要給付項目",
                12: "連結標的資產",
                15: "本商品年期",
                18: "情境分析",
                27: "ISIN：XS1999900001",
            }.get(n, "範本說明"),
        )
        if n == 11:
            sub(1)
            line(f"固定配息率={s.annual}%×1/12")
            line("配息期數=6")
            line("Nt" if s.obs == "D" else "付息日(如未在該計息期間自動提前到期)")
            for i, (end, pay) in enumerate(zip(s.ends, s.payments, strict=True), 1):
                vals = [str(i)]
                if s.obs == "D":
                    start = s.ends[i - 2] + dt.timedelta(days=1) if i > 2 else None
                    if start:
                        while start.weekday() >= 5:
                            start += dt.timedelta(days=1)
                    vals += [zh_date(start) if start else "-", zh_date(end), zh_date(pay), "20" if start else "-"]
                else:
                    vals += [zh_date(pay)]
                for t in vals:
                    line(t)
            line("註1：假設日期")
            sub(2)
            line("自動提前到期價格為期初股價×100%")
            line("所有連結標的皆已成為鎖定股票" if s.memory else "評價等於或大於其自動提前到期價格")
            if s.obs == "D":
                line("自動提前到期決定日為自" + zh_date(s.ends[1]) + "(含)起每個預定交易日")
            else:
                line("自動提前到期決定日")
                line("自動提前到期金額付款日")
                for i in range(1, 6):
                    for t in [str(i), zh_date(s.ends[i]), zh_date(s.payments[i])]:
                        line(t)
                line("投資人應注意")
            sub(3)
            line("執行價為期初股價×70%")
            if s.ki == "none":
                line("期末股價〔等於或大於〕執行價；期末股價〔小於〕執行價")
            else:
                line("觸及不保本價格為期初股價×60%")
                line("觸及不保本事件決定日為" + ("最後評價日" if s.ki == "AM" else "每個預定交易日"))
        if n == 12:
            sub(1)
            table()
            sub(2)
            line("標的說明")
        if n == 15:
            vals = [
                "本商品年期為6個月",
                "發行日：2030年1月14日",
                "到期日目前表定為" + zh_date(s.payments[-1]),
                "日期說明",
                "最後評價日：" + zh_date(s.ends[-1]),
                "交易日：2030年1月7日",
            ]
            for i, t in enumerate(vals, 1):
                sub(i)
                line(t)
        if n == 18:
            line("商品天期為6個月期，每單位面額為美元10,000元")
            line("固定配息率為1.0000%，配息期數=6，且假設")
            table(scenario=True)
            line("*假設天期")
            line("發行價格為100%")
            line("每單位期初投資金額=美元10,000.00(=10,000×100%)")
            for i, zh in enumerate("一二三" if s.ki == "none" else "一二三四"):
                line("情境分析" + zh + ")")
                periods = 2 if i == 0 else 6
                line(f"第1個至第{periods}個計息期間")
                if i > 0:
                    line("於6個月存續期間共6次配息")
                line("固定配息金額=美元10,000×1.0000%=美元100.00")
                if i > 0:
                    line("6個計息期間配息金額共為美元600.00")
                if i == 0 and s.partial_coupon:
                    line("第3個計息期間配息金額=美元10,000×1.0000%×5/20=美元25.00")
                if i >= 2:
                    line("假設標的（虛構標的1），執行價美元70.0000")
                    if s.ki != "none":
                        line("觸及不保本價格美元60.0000")
                if i == (2 if s.ki == "none" else 3):
                    line("交割股數：假設實物給付")
                else:
                    line("到期贖回金額為美元10,000×100%=美元10,000.00")
                    line(
                        f"損益=美元10,000.00+美元{100 * periods}.00"
                        + ("+美元25.00" if i == 0 and s.partial_coupon else "")
                        + f"-美元10,000.00=美元{100 * periods + (25 if i == 0 and s.partial_coupon else 0)}.00"
                    )
                    if i > 0:
                        line("平均年化報酬率(以簡單平均年化報酬率之方式計算)為12.00%")
    chapter("二", "相關機構事業概況")
    line("發行機構：(1)事業名稱：" + STD["issuer_name"]["hsbc"].replace("（", "(").replace("）", ")"))
    line("受託或銷售機構：(a)事業名稱：" + STD["distributor"]["name"] + "(b)電話：02-5556-1313")
    line("(c)營業所在地：" + STD["distributor"]["address"] + "(d)負責人姓名：" + STD["distributor"]["chairman"])
    line("結算機構")
    chapter("三", "商品風險揭露")
    line(warning)
    chapter("四", "一般交易事項")
    line("商品開始受理申購日：2030年1月7日")
    line("商品申購結束受理日：2030年1月7日")
    line("最低申購金額：美元10,000元")
    line("最低加購金額：美元10,000元")
    for t in ["申購費用", "提前贖回費用", "分銷費用"]:
        w.row([(60, t), (180, s.replacements.get("0%~5%", "0%~5%"))])
    chapter("五", "特別記載事項")
    line("其他說明")
    for i, page in enumerate(w.doc):
        page.insert_text((270, 806), f"第{i + 1}頁，共{len(w.doc)}頁", fontname=FONT, fontsize=8)
    return w.save(path)


def build_inquiry(path: Path, s: Spec, overrides=None):
    cfg = tomllib.loads(ORDER_FORMAT.read_text(encoding="utf-8"))
    vals = {
        "product_code": s.code,
        "isin": "XS1999900001",
        "currency": "USD",
        "denomination": 10000,
        "trade_date": dt.date(2030, 1, 7),
        "issue_date": dt.date(2030, 1, 14),
        "final_valuation_date": s.ends[-1],
        "maturity_date": s.payments[-1],
        "ko_pct": 100,
        "strike_pct": 70,
        "ki_pct": 60 if s.ki != "none" else "-",
        "ki_type": s.ki if s.ki != "none" else "-",
        "ko_observation": s.obs,
        "ko_memory": "Y" if s.memory else "N",
        "coupon_pa_pct": float(s.annual),
        "tenor_months": 6,
        "first_callable_period": 2,
        "initial_pricing": "收盤價",
    }
    for i in range(1, 13):
        vals[f"autocall_date_{i}"] = (
            s.ends[i - 1] if (s.obs == "P" and 2 <= i <= 6) or (s.obs == "D" and i in (2, 6)) else "-"
        )
    for i in range(1, 6):
        vals[f"underlying_{i}"] = f"ZZ{i} UW" if i <= s.count else "-"
        for k, v in [("initial", 100), ("strike", 70), ("ko", 100), ("ki", 60)]:
            vals[f"underlying_{i}_{k}_price"] = v if i <= s.count and (k != "ki" or s.ki != "none") else "-"
    vals.update(overrides or {})
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "樣本清單"
    headers = [h for h, v in cfg["columns"].items() if isinstance(v, str)] + ["發行機構"]
    for i, h in enumerate(headers, 1):
        ws.cell(3, i, h)
        ws.cell(4, i, issuer_value("HSBC") if h == "發行機構" else vals[cfg["columns"][h]])
    wb.save(path)
    wb.close()
    return path
