"""PANEL 工作流程入口（PanelSession）：參考條件表＋多份說明書的預覽、核對、儲存與失效檢查。

只用合成檔案，不依賴桌面視窗或 parser 內部。核對規則本身見 test_check_barc*.py、test_batch.py。
"""

import datetime as dt
import json
from dataclasses import FrozenInstanceError
from pathlib import Path

import fitz
import openpyxl
import pytest

from fcn_checker.ingestion import IngestionError
from fcn_checker.panel import parse_args, session_from_args
from fcn_checker.panel_workflow import PanelSession
from harness import ISSUER_PREFIXES, REVIEW_STANDARD, ROOT
from reference_synth import REFERENCE_FORMAT, build_reference_sheet
from synth import SYNTH_ISIN, Spec, build_pdf, reference_row

NOW = dt.datetime(2030, 2, 3, 4, 5, 6)


def inputs(tmp_path: Path, *specs: Spec, rows: list[dict] | None = None):
    specs = specs or (Spec(),)
    pdfs = [build_pdf(tmp_path / f"{s.product_code}_TS.pdf", s) for s in specs]
    sheet = build_reference_sheet(tmp_path / "FCN參考條件.xlsx", rows or [reference_row(s) for s in specs])
    return sheet, pdfs


def session_for(tmp_path: Path, sheet, pdfs, standard=REVIEW_STANDARD, config_dir=None) -> PanelSession:
    session = PanelSession(standard, config_dir or ROOT / "config", install_root=tmp_path)
    session.select(sheet, pdfs)
    return session


# ---------------------------------------------------------------- 預覽


def test_preview_lists_each_pdf_without_writing_anything(tmp_path):
    ok, unknown, other = Spec(), Spec(product_code="029199990002"), Spec(product_code="999199990001")
    sheet, pdfs = inputs(tmp_path, ok, unknown, other, rows=[reference_row(ok)])
    before = set(tmp_path.iterdir())
    session = session_for(tmp_path, sheet, pdfs)

    preview = session.load_preview()
    first, missing, unsupported = preview.rows
    assert (first.issuer, first.product_code, first.reference_row, first.problem) == ("BARC", ok.product_code, 4, "")
    assert first.product_code_evidence[0].page == 1
    assert missing.reference_row is None and "找不到" in missing.problem
    assert unsupported.unsupported and "未支援上手" in unsupported.problem
    assert session.preview == preview
    assert set(tmp_path.iterdir()) == before
    with pytest.raises(FrozenInstanceError):
        first.problem = "modified"


def test_preview_and_result_show_pdfs_sharing_one_reference_row(tmp_path):
    spec = Spec()
    pdfs = [build_pdf(tmp_path / f"{spec.product_code}_{v}.pdf", spec) for v in ("舊版", "新版")]
    sheet = build_reference_sheet(tmp_path / "FCN參考條件.xlsx", [reference_row(spec)])
    session = session_for(tmp_path, sheet, pdfs)
    old, new = session.load_preview().rows
    assert old.reference_row == new.reference_row == 4
    assert "同一批有多份說明書對到同一個 TDCC Code" in old.problem and pdfs[1].name in old.problem
    assert pdfs[0].name in new.problem

    outcome = session.start_check()
    assert "2 份需人工覆核" in outcome.headline
    for item, other in zip(outcome.ordered_items, reversed(pdfs), strict=True):
        assert other.name in outcome.ordered_results(item)[0].message


def test_preview_reports_reference_sheet_column_problems(tmp_path):
    from reference_synth import REFERENCE_HEADERS

    spec = Spec()
    pdf = build_pdf(tmp_path / f"{spec.product_code}_TS.pdf", spec)
    sheet = build_reference_sheet(tmp_path / "ref.xlsx", [reference_row(spec)], [*REFERENCE_HEADERS, "新欄位"])
    preview = session_for(tmp_path, sheet, [pdf]).load_preview()
    assert any("新欄位" in w for w in preview.warnings)


def test_pdf_product_code_is_never_taken_from_the_sheet(tmp_path):
    spec = Spec()
    sheet, pdfs = inputs(tmp_path, spec)
    with fitz.open(pdfs[0]) as doc:
        page = doc[0]
        page.add_redact_annot(page.search_for(spec.product_code)[0])
        page.apply_redactions()
        doc.saveIncr()
    row = session_for(tmp_path, sheet, pdfs).load_preview().rows[0]
    assert row.product_code is None and row.problem


def test_selection_is_required_and_changing_it_clears_preview(tmp_path):
    sheet, pdfs = inputs(tmp_path)
    session = session_for(tmp_path, sheet, pdfs)
    session.load_preview()
    session.select(sheet, [])
    assert session.preview is None
    with pytest.raises(IngestionError, match="請先選取"):
        session.load_preview()


@pytest.mark.parametrize("kind", ["broken_excel", "missing_sheet"])
def test_invalid_reference_sheet_never_leaves_valid_preview(tmp_path, kind):
    sheet, pdfs = inputs(tmp_path)
    session = session_for(tmp_path, sheet, pdfs)
    session.load_preview()
    if kind == "broken_excel":
        sheet.write_bytes(b"not an Excel")
    else:
        wb = openpyxl.load_workbook(sheet)
        wb["樣本清單"].title = "其他"
        wb.save(sheet)
    with pytest.raises(IngestionError):
        session.load_preview()
    assert session.preview is None and session.message


# ---------------------------------------------------------------- 核對


def test_check_requires_preview_and_writes_nothing(tmp_path):
    ok, bad = Spec(), Spec(product_code="029199990002")
    sheet, pdfs = inputs(tmp_path, ok, bad, rows=[reference_row(ok), reference_row(bad, **{"K(%)": 71})])
    session = session_for(tmp_path, sheet, pdfs)
    with pytest.raises(IngestionError, match="預覽"):
        session.start_check()
    session.load_preview()
    before = set(tmp_path.rglob("*"))

    outcome = session.start_check()
    assert set(tmp_path.rglob("*")) == before
    assert session.outcome is outcome
    assert [i.term_sheet for i in outcome.ordered_items] == [pdfs[1], pdfs[0]], "有問題的排前面"
    assert "1 份通過" in outcome.headline and "1 份不一致" in outcome.headline
    assert "按「儲存核對結果」" in outcome.headline and "新檔" not in outcome.headline
    first = outcome.ordered_results(outcome.ordered_items[0])
    assert first[0].status.value == "MISMATCH"
    assert session.receipt is None, "還沒儲存，沒有收據"


@pytest.mark.parametrize("changed", ["selection", "pdf", "sheet", "standard", "config"])
def test_changing_sources_clears_result_and_requires_new_preview(tmp_path, changed):
    sheet, pdfs = inputs(tmp_path)
    standard = tmp_path / "standard.toml"
    standard.write_bytes(REVIEW_STANDARD.read_bytes())
    config = tmp_path / "config"
    config.mkdir()
    for f in (REFERENCE_FORMAT, ISSUER_PREFIXES):
        (config / f.name).write_bytes(f.read_bytes())
    session = session_for(tmp_path, sheet, pdfs, standard, config)
    session.load_preview()
    session.start_check()
    if changed == "selection":
        session.select(sheet, pdfs)
    elif changed == "pdf":
        build_pdf(pdfs[0], Spec(tenor=7))
    elif changed == "sheet":
        build_reference_sheet(sheet, [reference_row(Spec(), **{"K(%)": 71})])
    elif changed == "standard":
        standard.write_text(standard.read_text(encoding="utf-8") + "\n# change\n", encoding="utf-8")
    else:
        prefixes = config / ISSUER_PREFIXES.name
        prefixes.write_text(prefixes.read_text(encoding="utf-8") + "\n# change\n", encoding="utf-8")
    assert not session.check_sources()
    assert session.outcome is None and session.preview is None
    assert "重新載入" in session.message or changed == "selection"
    with pytest.raises(IngestionError, match="預覽"):
        session.start_check()
    assert session.load_preview() is not None


def test_reading_preview_and_outcome_never_touches_the_files(tmp_path):
    spec = Spec()
    sheet, pdfs = inputs(tmp_path, spec, rows=[reference_row(spec, **{"K(%)": 71})])
    session = session_for(tmp_path, sheet, pdfs)
    session.load_preview()
    outcome = session.start_check()
    (item,) = outcome.batch.items
    pdfs[0].unlink()  # 外部刪掉說明書

    assert session.preview is not None and session.outcome is outcome, "讀取屬性不讀檔、不清狀態"
    assert session.release_state(item).allowed, "顯示放行按鈕狀態也不讀檔"
    assert session.check_sources() is False, "明確檢查才失效"
    assert session.preview is None and session.outcome is None
    assert "重新載入" in session.message


@pytest.mark.parametrize("action", ["release", "save"])
def test_release_and_save_check_the_sources_first(tmp_path, action):
    spec = Spec()
    sheet, pdfs = inputs(tmp_path, spec, rows=[reference_row(spec, **{"K(%)": 71})])
    session = session_for(tmp_path, sheet, pdfs)
    session.load_preview()
    (item,) = session.start_check().batch.items
    build_reference_sheet(sheet, [reference_row(spec, **{"K(%)": 72})])  # 外部改了參考條件表

    with pytest.raises(IngestionError, match="重新"):
        if action == "release":
            session.release(item)
        else:
            session.save(tmp_path / "reports", now=NOW)
    assert not item.released and session.outcome is None, "結果失效，畫面要清空"
    assert not (tmp_path / "reports").exists()


def test_pdf_missing_at_preview_is_listed_and_the_rest_still_loads(tmp_path):
    spec = Spec()
    sheet, pdfs = inputs(tmp_path, spec)
    missing = tmp_path / "029199990009_TS.pdf"
    session = session_for(tmp_path, sheet, [*pdfs, missing])
    preview = session.load_preview()
    assert [bool(r.problem) for r in preview.rows] == [False, True]
    assert "找不到說明書" in preview.rows[1].problem
    assert session.check_sources(), "選取時就不存在的說明書不算來源變更"
    session.start_check()
    assert session.save(tmp_path / "reports", now=NOW).output is not None


def test_unchanged_sources_pass_the_check(tmp_path):
    sheet, pdfs = inputs(tmp_path)
    session = session_for(tmp_path, sheet, pdfs)
    assert not session.check_sources(), "還沒有預覽"
    session.load_preview()
    outcome = session.start_check()
    assert session.check_sources() and session.outcome is outcome


# ---------------------------------------------------------------- 儲存


def test_save_writes_result_file_and_record_only_when_asked(tmp_path):
    spec = Spec()
    sheet, pdfs = inputs(tmp_path, spec)
    original = sheet.read_bytes()
    session = session_for(tmp_path, sheet, pdfs)
    with pytest.raises(IngestionError):
        session.save(tmp_path / "reports", now=NOW)
    session.load_preview()
    outcome = session.start_check()

    assert session.save(None).cancelled
    assert not (tmp_path / "reports").exists()
    receipt = session.save(tmp_path / "reports", now=NOW)
    assert receipt.complete
    assert receipt.output == tmp_path / "reports" / "FCN參考條件_核對結果_20300203-040506.xlsx"
    assert f"核對結果檔：{receipt.output}" in receipt.summary
    assert [p.name for p in (tmp_path / "reports").iterdir()] == [receipt.output.name], "選的資料夾只有核對結果檔"
    assert (tmp_path / "runtime" / "核對紀錄" / "20300203-040506.json").is_file()
    ws = openpyxl.load_workbook(receipt.output)["回填後"]
    assert ws["F4"].value == SYNTH_ISIN
    assert sheet.read_bytes() == original
    assert session.outcome is outcome

    again = session.save(tmp_path / "reports", now=NOW + dt.timedelta(seconds=1))
    assert again.complete and again.output != receipt.output


def test_session_keeps_the_latest_receipt_for_the_current_result(tmp_path):
    spec = Spec()
    sheet, pdfs = inputs(tmp_path, spec, rows=[reference_row(spec, **{"K(%)": 71})])
    session = session_for(tmp_path, sheet, pdfs)
    session.load_preview()
    (item,) = session.start_check().batch.items
    session.release(item)
    receipt = session.save(tmp_path / "reports", now=NOW)
    assert session.receipt is receipt and receipt.filled(item), "「已回填」來自最近一次儲存的收據"

    session.cancel_release(item)
    assert session.receipt is None, "放行改了，上一次的收據不再是目前的結果"
    assert receipt.filled(item), "收據本身不變"


def test_save_failure_is_reported_and_result_is_kept(tmp_path):
    sheet, pdfs = inputs(tmp_path)
    session = session_for(tmp_path, sheet, pdfs)
    session.load_preview()
    outcome = session.start_check()
    blocked = tmp_path / "blocked"
    blocked.write_text("keep")

    receipt = session.save(blocked, now=NOW)
    assert not receipt.complete
    assert "無法建立核對結果檔的資料夾" in receipt.summary
    assert blocked.read_text() == "keep"
    assert session.outcome is outcome


def test_record_failure_is_reported_and_the_result_file_is_kept(tmp_path):
    sheet, pdfs = inputs(tmp_path)
    session = session_for(tmp_path, sheet, pdfs)
    session.load_preview()
    session.start_check()
    (tmp_path / "runtime").write_text("不是資料夾", encoding="utf-8")
    receipt = session.save(tmp_path / "reports", now=NOW)
    assert receipt.output is not None and receipt.output.is_file()
    assert not receipt.complete and "核對紀錄" in receipt.summary


def test_save_after_sources_changed_is_refused(tmp_path):
    sheet, pdfs = inputs(tmp_path)
    session = session_for(tmp_path, sheet, pdfs)
    session.load_preview()
    session.start_check()
    build_reference_sheet(sheet, [reference_row(Spec(), **{"K(%)": 71})])
    with pytest.raises(IngestionError):
        session.save(tmp_path / "reports", now=NOW)
    assert not (tmp_path / "reports").exists()


# ---------------------------------------------------------------- 人工放行


def checked(tmp_path: Path, *specs: Spec, rows: list[dict] | None = None):
    sheet, pdfs = inputs(tmp_path, *specs, rows=rows)
    session = session_for(tmp_path, sheet, pdfs)
    session.load_preview()
    return session, session.start_check()


def saved(session: PanelSession, tmp_path: Path):
    receipt = session.save(tmp_path / "reports", now=NOW)
    assert receipt.complete
    wb = openpyxl.load_workbook(receipt.output)
    record = json.loads((tmp_path / "runtime" / "核對紀錄" / "20300203-040506.json").read_text(encoding="utf-8"))
    return wb["回填後"], wb["錯誤清單"], record


@pytest.mark.parametrize(
    ("row", "original"), [({"K(%)": 71}, "不一致"), ({"K(%)": None}, "需人工覆核")], ids=["mismatch", "review"]
)
def test_release_and_cancel_update_the_headline(tmp_path, row, original):
    spec = Spec()
    session, outcome = checked(tmp_path, spec, rows=[reference_row(spec, **row)])
    (item,) = outcome.batch.items
    session.release(item)
    assert item.released
    assert "0 份通過、1 份人工放行" in session.message and f"0 份{original}" in session.message
    assert "再儲存" not in session.message, "還沒儲存過，不必提醒"
    session.cancel_release(item)
    assert not item.released
    assert "0 份人工放行" in session.message and f"1 份{original}" in session.message


def test_changing_a_release_after_saving_says_to_save_again(tmp_path):
    spec = Spec()
    session, outcome = checked(tmp_path, spec, rows=[reference_row(spec, **{"K(%)": 71})])
    (item,) = outcome.batch.items
    session.release(item)
    assert session.save(tmp_path / "reports", now=NOW).complete
    session.cancel_release(item)
    assert "上一次儲存的核對結果檔已不是目前的結果，請再儲存一次" in session.message


def test_release_button_state_hides_the_reason_for_passed_term_sheets(tmp_path):
    ok, bad = Spec(), Spec(product_code="029199990002")
    session, outcome = checked(tmp_path, ok, bad, rows=[reference_row(ok), reference_row(bad, **{"K(%)": 71})])
    passed, mismatch = outcome.batch.items
    assert session.release_state(passed) == (False, ""), "已通過的不能放行，也不必說明原因"
    assert session.release_state(mismatch) == (True, "")
    session.start_check()  # 原結果失效
    allowed, reason = session.release_state(mismatch)
    assert not allowed and "重新核對" in reason


def test_released_items_sort_with_passed_ones(tmp_path):
    ok, bad, other = Spec(), Spec(product_code="029199990002"), Spec(product_code="029199990003")
    rows = [reference_row(ok), reference_row(bad, **{"K(%)": 71}), reference_row(other, **{"K(%)": 71})]
    session, outcome = checked(tmp_path, ok, bad, other, rows=rows)
    released = outcome.batch.items[1]
    session.release(released)
    assert [i.term_sheet.name[:12] for i in outcome.ordered_items] == [
        other.product_code,
        ok.product_code,
        bad.product_code,
    ]
    assert "1 份通過、1 份人工放行、1 份不一致" in outcome.headline


def test_release_selects_the_top_item_and_cancel_keeps_the_same_one(tmp_path):
    ok, bad, other = Spec(), Spec(product_code="029199990002"), Spec(product_code="029199990003")
    rows = [reference_row(ok), reference_row(bad, **{"K(%)": 71}), reference_row(other, **{"K(%)": 71})]
    session, outcome = checked(tmp_path, ok, bad, other, rows=rows)
    item = outcome.ordered_items[1]
    session.release(item)
    assert outcome.ordered_items[0] is not item, "前提：放行那份已排到下方"
    assert outcome.selection_after_release_change(item) is outcome.ordered_items[0], "放行後改選清單第一份"
    session.cancel_release(item)
    assert outcome.selection_after_release_change(item) is item, "取消放行後仍選那份"


@pytest.mark.parametrize("action", ["recheck", "reload", "source_changed"])
def test_release_is_cleared_when_the_result_is_replaced(tmp_path, action):
    spec = Spec()
    session, outcome = checked(tmp_path, spec, rows=[reference_row(spec, **{"K(%)": 71})])
    session.release(outcome.batch.items[0])
    if action == "recheck":
        assert not session.start_check().batch.items[0].released
    elif action == "reload":
        session.load_preview()
        assert session.outcome is None
    else:
        build_reference_sheet(session.reference_sheet, [reference_row(spec, **{"K(%)": 72})])
        assert not session.check_sources() and session.outcome is None
    with pytest.raises(IngestionError):
        session.release(outcome.batch.items[0])


def test_release_only_applies_to_the_current_result(tmp_path):
    spec = Spec()
    session, outcome = checked(tmp_path, spec, rows=[reference_row(spec, **{"K(%)": 71})])
    stale = outcome.batch.items[0]
    session.start_check()
    assert "重新核對" in session.release_problem(stale)
    with pytest.raises(IngestionError, match="重新核對"):
        session.release(stale)


# ---------------------------------------------------------------- 待處理


def test_not_covered_is_grouped_by_issuer_without_overwriting(tmp_path):
    import hsbc_synth
    from fcn_checker.issuers import BARC, HSBC

    barc, hsbc = Spec(), hsbc_synth.Spec()
    pdfs = [
        build_pdf(tmp_path / f"{barc.product_code}_TS.pdf", barc),
        hsbc_synth.build_pdf(tmp_path / f"{hsbc.code}_TS.pdf", hsbc),
        build_pdf(tmp_path / "029199990002_TS.pdf", Spec(product_code="029199990002")),
    ]
    inquiry = openpyxl.load_workbook(hsbc_synth.build_inquiry(tmp_path / "hsbc.xlsx", hsbc))["樣本清單"]
    hsbc_row = {c.value: inquiry.cell(4, c.column).value for c in inquiry[3]}
    rows = [reference_row(barc), hsbc_row, reference_row(Spec(product_code="029199990002"))]
    sheet = build_reference_sheet(tmp_path / "FCN參考條件.xlsx", rows)
    session = session_for(tmp_path, sheet, pdfs)
    session.load_preview()
    outcome = session.start_check()

    groups = {g.issuer: g.descriptions for g in outcome.not_covered}
    assert list(groups) == ["BARC", "HSBC"]
    for issuer in (BARC, HSBC):
        assert groups[issuer.code] == tuple(n["description"] for n in issuer.not_covered), "同一家只列一次"
    shared = {"doc.underlying_names", "field.monthly_ki"}
    for rule_id in shared:  # 兩家都有、說明不同的 rule_id 各自保留
        for issuer in (BARC, HSBC):
            (desc,) = [n["description"] for n in issuer.not_covered if n["rule_id"] == rule_id]
            assert desc in groups[issuer.code]
    assert outcome.not_covered_count == len(BARC.not_covered) + len(HSBC.not_covered)


# ---------------------------------------------------------------- 設定檔與啟動參數


def test_config_falls_back_to_builtin_files_when_root_config_lacks_them(tmp_path):
    old_install = tmp_path / "config"
    old_install.mkdir()
    (old_install / "review_standard.toml").write_bytes(REVIEW_STANDARD.read_bytes())
    session = PanelSession(old_install / "review_standard.toml", old_install)
    assert session.reference_format == REFERENCE_FORMAT.resolve()
    assert session.issuer_prefixes == ISSUER_PREFIXES.resolve()
    assert str(REFERENCE_FORMAT.resolve()) in "\n".join(str(p) for _, p in session.config_paths)

    own = old_install / "issuer_prefixes.toml"
    own.write_bytes(ISSUER_PREFIXES.read_bytes())
    assert PanelSession(old_install / "review_standard.toml", old_install).issuer_prefixes == own.resolve()


def test_panel_accepts_legacy_launcher_arguments(tmp_path):
    args = parse_args(
        [
            "--order-formats-dir",
            str(tmp_path / "config" / "order_formats"),
            "--review-standard",
            str(tmp_path / "config" / "review_standard.toml"),
            "--install-root",
            str(tmp_path),
        ]
    )
    assert args.config_dir == tmp_path / "config", "舊啟動器只給 order_formats 資料夾：取其上一層 config"
    args = parse_args(["--config-dir", str(tmp_path / "cfg")])
    assert args.config_dir == tmp_path / "cfg"


def test_double_click_launch_runs_from_the_root_and_ignores_old_update_pointer(tmp_path, monkeypatch):
    """雙擊入口一律從專案根目錄啟動；舊版自動更新留下的 .local/current.json 不再使用。"""
    import importlib.util

    root = tmp_path / "install"
    (root / "config").mkdir(parents=True)
    for f in (REVIEW_STANDARD, REFERENCE_FORMAT, ISSUER_PREFIXES):
        (root / "config" / f.name).write_bytes(f.read_bytes())
    (root / ".local").mkdir()
    (root / ".local" / "current.json").write_text('{"sha": "' + "2" * 40 + '", "release": ".local/releases/old"}')

    captured = {}
    monkeypatch.setattr("fcn_checker.panel.main", lambda argv: captured.setdefault("argv", argv))
    spec = importlib.util.spec_from_file_location("bootstrap", ROOT / "panel_bootstrap.py")
    bootstrap = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(bootstrap)
    bootstrap.run(str(root / "launch_panel.pyw"))

    args = parse_args(captured["argv"])
    assert args.config_dir == args.builtin_config_dir == root / "config"
    session = session_from_args(args)
    assert session.install_root == root.resolve(), "核對紀錄寫到根目錄的 runtime/"
    assert session.reference_format == (root / "config" / REFERENCE_FORMAT.name).resolve()


# ---------------------------------------------------------------- 圖示


def test_panel_icon_ships_every_windows_size():
    """圖示放在套件內（pip install 後也找得到），且含 100%～250% 縮放下標題列、工作列、桌面會用到的尺寸。

    Windows 小圖示要 16 × 縮放、大圖示要 32 × 縮放；缺尺寸時 Tk 會把小圖放大，工作列就變糊（#111）。
    """
    import struct

    from fcn_checker.panel import PANEL_ICON

    data = PANEL_ICON.read_bytes()
    reserved, kind, count = struct.unpack_from("<HHH", data)
    assert (reserved, kind) == (0, 1), "必須是 Windows .ico"
    sizes = {data[6 + 16 * i] or 256 for i in range(count)}  # 寬度 0 代表 256
    scales = (1, 1.25, 1.5, 1.75, 2, 2.25, 2.5)
    assert sizes >= {round(16 * s) for s in scales} | {round(32 * s) for s in scales} | {48, 96, 128, 256}


def test_window_icon_failure_keeps_the_default_icon():
    """圖示無法載入時不跳錯誤，PANEL 照常開啟。"""
    import tkinter as tk

    from fcn_checker.panel import apply_window_icon

    class Root:
        def iconbitmap(self, default=None):
            raise tk.TclError(f'bitmap "{default}" not defined')

    apply_window_icon(Root())


def test_app_id_is_windows_only(monkeypatch):
    """非 Windows 平台不碰 Windows 專屬 API。"""
    import fcn_checker.panel as panel

    touched = []

    class NoCtypes:
        def __getattr__(self, name):
            touched.append(name)
            raise AssertionError(name)

    monkeypatch.setattr(panel.sys, "platform", "linux")
    monkeypatch.setattr(panel, "ctypes", NoCtypes())
    panel.set_windows_app_id()
    assert touched == []


def test_window_icon_on_non_windows_only_uses_tk(monkeypatch):
    """非 Windows 平台只交給 Tk 設圖示，不碰 Windows 專屬 API。"""
    import fcn_checker.panel as panel

    calls = []

    class Root:
        def iconbitmap(self, default=None):
            calls.append(default)

    monkeypatch.setattr(panel.sys, "platform", "linux")
    monkeypatch.setattr(panel, "_set_window_icons", lambda hwnd: calls.append("win32"))
    panel.apply_window_icon(Root())
    assert calls == [str(panel.PANEL_ICON)]


# ---------------------------------------------------------------- 選檔起始資料夾


def test_file_dialogs_remember_each_kind_across_restarts(tmp_path):
    """參考條件表與說明書各自記住上一次選檔的資料夾；重新建立（重開 PANEL）後仍讀得到。"""
    from fcn_checker.panel import LastFolders

    sheets, pdfs = tmp_path / "條件表", tmp_path / "說明書"
    sheets.mkdir()
    pdfs.mkdir()
    record = tmp_path / ".local" / "panel_folders.json"
    folders = LastFolders(record)
    assert folders.initial("sheet") is None, "還沒選過時交給 Windows 決定"

    folders.remember("sheet", sheets / "FCN參考條件.xlsx")
    folders.remember("pdf", pdfs / "029_TS.pdf")

    reopened = LastFolders(record)
    assert reopened.initial("sheet") == str(sheets)
    assert reopened.initial("pdf") == str(pdfs)


def test_file_dialogs_fall_back_when_folder_or_record_is_unusable(tmp_path):
    """記住的資料夾已刪除、或記錄檔損壞時，不跳錯誤，改由 Windows 決定起始位置。"""
    from fcn_checker.panel import LastFolders

    gone = tmp_path / "已刪除"
    gone.mkdir()
    record = tmp_path / ".local" / "panel_folders.json"
    folders = LastFolders(record)
    folders.remember("sheet", gone / "FCN參考條件.xlsx")
    gone.rmdir()
    assert folders.initial("sheet") is None

    record.write_text("{損壞", encoding="utf-8")
    assert folders.initial("sheet") is None
    folders.remember("pdf", tmp_path / "029_TS.pdf")
    assert folders.initial("pdf") == str(tmp_path), "損壞的記錄檔會被新記錄取代"
