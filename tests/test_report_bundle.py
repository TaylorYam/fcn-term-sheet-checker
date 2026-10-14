"""回報包（Issue #204）：一批核對結果中未通過的商品，匯出成維護者可以原樣重現、作業人員可以逐條標註判斷的資料夾。"""

from __future__ import annotations

import datetime as dt
import json
from pathlib import Path

import openpyxl
import pytest

from fcn_checker import report_bundle
from fcn_checker.ingestion import IngestionError
from fcn_checker.panel_workflow import PanelSession
from fcn_checker.schema import CheckStatus
from harness import CONFIG, ROOT, check_all, with_iis
from reference_synth import REFERENCE_HEADERS, build_reference_sheet, reference_row
from synth import Spec, build_pdf

NOW = dt.datetime(2030, 2, 3, 4, 5, 6)
BUNDLE = "FCN參考條件_回報包_20300203-040506"


def pdf_for(tmp_path: Path, spec: Spec) -> Path:
    return build_pdf(tmp_path / f"{spec.product_code}_TS.pdf", spec)


def checked(tmp_path: Path, specs: list[Spec], rows: list[dict]):
    sheet = build_reference_sheet(tmp_path / "FCN參考條件.xlsx", rows)
    return check_all(sheet, [pdf_for(tmp_path, s) for s in specs]), sheet


def export(outcome, tmp_path: Path, now: dt.datetime = NOW) -> Path:
    return report_bundle.export_bundle(outcome, tmp_path / "out", root=tmp_path, now=now)


def sheet_dicts(path: Path, sheet: str, header_row: int = 1) -> list[dict]:
    rows = list(openpyxl.load_workbook(path)[sheet].iter_rows(min_row=header_row, values_only=True))
    return [dict(zip(rows[0], r, strict=True)) for r in rows[1:]]


def summary(path: Path) -> dict:
    ws = openpyxl.load_workbook(path)["摘要"]
    return {k: v for k, v in ws.iter_rows(values_only=True)}


def failing_pair():
    ok, bad = Spec(product_code="029199990001"), Spec(product_code="029199990002")
    return ok, bad, [reference_row(ok), reference_row(bad, **{"K(%)": 71})]


# ---------------------------------------------------------------- 內容


def test_bundle_holds_only_the_failing_products(tmp_path):
    ok, bad, rows = failing_pair()
    outcome, _ = checked(tmp_path, [ok, bad], rows)
    failing = [i for i in outcome.items if i.product_code == bad.product_code]
    assert {i.report.status for i in failing} == {CheckStatus.MISMATCH}

    out = export(outcome, tmp_path)

    assert out == tmp_path / "out" / BUNDLE
    assert sorted(p.name for p in out.iterdir()) == [
        "PDF",
        "參考條件表_未通過.xlsx",
        "回報總表.xlsx",
        "核對紀錄_未通過.json",
    ]
    assert sorted(p.name for p in (out / "PDF").iterdir()) == [
        f"{bad.product_code}_IIS.pdf",
        f"{bad.product_code}_TS.pdf",
    ]
    assert (out / "PDF" / f"{bad.product_code}_TS.pdf").read_bytes() == failing[0].term_sheet.read_bytes()

    rows_out = sheet_dicts(out / "參考條件表_未通過.xlsx", "樣本清單", header_row=3)
    assert [r["TDCC Code"] for r in rows_out] == [bad.product_code]
    assert rows_out[0]["K(%)"] == 71

    record = json.loads((out / "核對紀錄_未通過.json").read_text(encoding="utf-8"))
    assert [i["pdf"] for i in record["items"]] == [i.term_sheet.name for i in failing]
    assert record["result_file"] is None and record["metadata"]["inputs"]["reference_sheet"]["sha256"]


def test_report_sheet_has_one_row_per_message_with_a_judgement_dropdown(tmp_path):
    ok, bad, rows = failing_pair()
    outcome, _ = checked(tmp_path, [ok, bad], rows)
    failing = [i for i in outcome.items if i.product_code == bad.product_code]

    out = export(outcome, tmp_path)

    path = out / "回報總表.xlsx"
    report = sheet_dicts(path, "回報")
    expected = [(i.term_sheet.name, m) for i in failing for m in i.problem_messages]
    assert [(r["PDF 檔名"], r["錯訊"]) for r in report] == expected
    assert all(r["上手"] == "BARC" and r["TDCC Code"] == bad.product_code for r in report)
    assert all(r["我的判斷"] is None and r["備註"] is None for r in report)
    ws = openpyxl.load_workbook(path)["回報"]
    (dv,) = ws.data_validations.dataValidation
    assert dv.type == "list" and dv.formula1 == '"誤報,真錯,不確定"'
    assert str(dv.sqref) == f"E2:E{len(report) + 1}"


def test_summary_counts_every_pdf_by_category(tmp_path):
    ok, bad, rows = failing_pair()
    other = Spec(product_code="999199990001")  # 未支援上手
    outcome, _ = checked(tmp_path, [ok, bad, other], [*rows, reference_row(other)])

    out = export(outcome, tmp_path)

    counts = summary(out / "回報總表.xlsx")
    assert counts["參考條件表"] == "FCN參考條件.xlsx"
    assert counts["PDF 總數"] == 6
    assert (counts["通過"], counts["不一致"], counts["未支援上手（已知不支援，未列入回報包）"]) == (2, 2, 2)
    assert counts["列入回報包"] == 2


def test_known_unsupported_and_released_documents_are_left_out(tmp_path):
    ok, bad, rows = failing_pair()
    other = Spec(product_code="999199990001")
    outcome, _ = checked(tmp_path, [ok, bad, other], [*rows, reference_row(other)])
    assert [i.term_sheet.name for i in report_bundle.reportable(outcome)] == [
        f"{bad.product_code}_TS.pdf",
        f"{bad.product_code}_IIS.pdf",
    ]

    for item in report_bundle.reportable(outcome):
        outcome.release(item)

    assert report_bundle.reportable(outcome) == ()
    with pytest.raises(IngestionError, match="沒有未通過"):
        export(outcome, tmp_path)
    assert not (tmp_path / "out").exists()


def test_reference_extract_keeps_only_the_known_columns(tmp_path):
    """合成表最右側另有一個無表頭欄（不在參考條件表欄位順序）：匯出時拿掉。"""
    ok, bad, rows = failing_pair()
    rows[1]["庫存狀態"] = "存續"  # 讓最右側的無表頭欄有值
    outcome, sheet = checked(tmp_path, [ok, bad], rows)
    assert openpyxl.load_workbook(sheet)["樣本清單"].max_column > len(REFERENCE_HEADERS)

    out = export(outcome, tmp_path)

    ws = openpyxl.load_workbook(out / "參考條件表_未通過.xlsx")["樣本清單"]
    assert [c.value for c in ws[3]] == REFERENCE_HEADERS
    assert ws.max_column == len(REFERENCE_HEADERS)
    assert ws.cell(4, REFERENCE_HEADERS.index("交易日") + 1).is_date


def test_product_without_a_reference_row_still_brings_its_pdfs(tmp_path):
    ok, bad, _ = failing_pair()
    outcome, _ = checked(tmp_path, [ok, bad], [reference_row(ok)])

    out = export(outcome, tmp_path)

    assert sorted(p.name for p in (out / "PDF").iterdir()) == [
        f"{bad.product_code}_IIS.pdf",
        f"{bad.product_code}_TS.pdf",
    ]
    assert sheet_dicts(out / "參考條件表_未通過.xlsx", "樣本清單", header_row=3) == []
    assert any("找不到" in r["錯訊"] for r in sheet_dicts(out / "回報總表.xlsx", "回報"))


def test_bundle_reproduces_the_same_failures(tmp_path):
    ok, bad, rows = failing_pair()
    outcome, _ = checked(tmp_path, [ok, bad], rows)
    before = [(i.term_sheet.name, i.report.status, i.problem_messages) for i in report_bundle.reportable(outcome)]

    out = export(outcome, tmp_path)

    again = check_all(out / "參考條件表_未通過.xlsx", sorted((out / "PDF").glob("*.pdf")), CONFIG)
    after = [(i.term_sheet.name, i.report.status, i.problem_messages) for i in again.items]
    assert sorted(after) == sorted(before)


# ---------------------------------------------------------------- 不覆蓋、來源變更


def test_existing_bundle_is_never_overwritten(tmp_path):
    ok, bad, rows = failing_pair()
    outcome, _ = checked(tmp_path, [ok, bad], rows)
    first = export(outcome, tmp_path)
    marker = (first / "回報總表.xlsx").read_bytes()

    with pytest.raises(IngestionError, match="已經存在"):
        export(outcome, tmp_path)

    assert (first / "回報總表.xlsx").read_bytes() == marker
    assert [p.name for p in (tmp_path / "out").iterdir()] == [BUNDLE], "沒有留下暫存資料夾"


def test_changed_sources_refuse_the_export(tmp_path):
    ok, bad, rows = failing_pair()
    outcome, sheet = checked(tmp_path, [ok, bad], rows)
    build_reference_sheet(sheet, [reference_row(ok), reference_row(bad, **{"K(%)": 72})])

    with pytest.raises(IngestionError, match="重新") as e:
        export(outcome, tmp_path)

    assert e.value.reason_code == "source_changed"
    assert not (tmp_path / "out").exists()


# ---------------------------------------------------------------- PANEL


def session_for(tmp_path: Path, specs: list[Spec], rows: list[dict]) -> PanelSession:
    sheet = build_reference_sheet(tmp_path / "FCN參考條件.xlsx", rows)
    session = PanelSession(config_dir=ROOT / "config", install_root=tmp_path)
    session.select(sheet, with_iis([pdf_for(tmp_path, s) for s in specs]))
    return session


def test_panel_exports_only_after_a_check_with_failures(tmp_path):
    ok, bad, rows = failing_pair()
    session = session_for(tmp_path, [ok, bad], rows)
    assert not session.can_export_bundle
    session.load_preview()
    assert not session.can_export_bundle, "還沒核對"
    session.start_check()
    assert session.can_export_bundle

    assert session.export_bundle(None) is None, "取消選資料夾：什麼都不寫"
    assert not (tmp_path / "out").exists()
    out = session.export_bundle(tmp_path / "out", now=NOW)
    assert out == tmp_path / "out" / BUNDLE and out.is_dir()
    assert str(out) in session.message
    assert not (tmp_path / "runtime").exists(), "匯出回報包不寫核對紀錄、不需要先儲存"


def test_panel_cannot_export_when_everything_passes(tmp_path):
    ok = Spec()
    session = session_for(tmp_path, [ok], [reference_row(ok)])
    session.load_preview()
    session.start_check()
    assert not session.can_export_bundle
    with pytest.raises(IngestionError, match="沒有未通過"):
        session.export_bundle(tmp_path / "out", now=NOW)


def test_panel_export_with_changed_sources_clears_the_result(tmp_path):
    ok, bad, rows = failing_pair()
    session = session_for(tmp_path, [ok, bad], rows)
    session.load_preview()
    session.start_check()
    build_reference_sheet(session.reference_sheet, [reference_row(ok), reference_row(bad, **{"K(%)": 72})])

    with pytest.raises(IngestionError, match="重新"):
        session.export_bundle(tmp_path / "out", now=NOW)

    assert session.outcome is None and not session.can_export_bundle
    assert not (tmp_path / "out").exists()
