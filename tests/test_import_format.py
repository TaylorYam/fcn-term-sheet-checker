"""核對結果檔「回填後」一律照匯入格式重新產生（Issue #203），不論參考條件表是公司原始待建檔表、作業人員模板或舊 FCN參考條件。

測試切點：批量入口 run_batch（預覽＋核對＋儲存）→ 核對結果檔；以 openpyxl 讀回「回填後」觀察。
匯入格式的欄位清單寫死在這裡（比照公司原始待建檔表第 2 列），不從設定檔推。所有數值皆為虛構。
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation

from fcn_checker.saving import run_batch
from fcn_checker.schema import CheckStatus
from harness import CONFIG, with_iis
from reference_synth import DATE_FORMAT, REFERENCE_HEADERS, build_reference_sheet, reference_row
from synth import SYNTH_ISIN, Spec, build_pdf

NOW = dt.datetime(2030, 2, 3, 4, 5, 6)
IMPORT_HEADERS = [
    "TS/IIS完成檢核", "交易日期", "Product", "TDCC Code", "ISIN Code", "私銀註記", "UF", "計價幣別", "單位面額",
    *(h for n in range(1, 8) for h in (f"分公司_{n}", f"分公司_{n}_承作金額")),
    "交易日", "發行日", *(f"比價日_{n}" for n in range(1, 13)), "下一比價日", "最終比價日", "到期日", "期初定價",
    "KO(%)", "KO(Freq)", "KO(memo)", "K(%)", "KI(%)", "KI(Freq)", "Coupon p.a. (%)", "天期(月)", "Non-Call(月)",
    *(h for n in range(1, 6) for h in (f"UL_{n}", f"UL_{n}_進場價", f"UL_{n}_執行價", f"UL_{n}_下限價", f"UL_{n}_KO價", f"UL_{n}_Memo")),
]  # fmt: skip


def run(tmp_path: Path, sheet: Path, specs: list[Spec]):
    pdfs = [build_pdf(tmp_path / f"{s.product_code}_TS.pdf", s) for s in specs]
    return run_batch(CONFIG, sheet, with_iis(pdfs), tmp_path / "reports", root=tmp_path, now=NOW)


def filled_sheet(receipt):
    """「回填後」與它的資料列（欄名 → 值）。"""
    ws = openpyxl.load_workbook(receipt.output)["回填後"]
    headers = [c.value for c in ws[2]]
    rows = [dict(zip(headers, r, strict=True)) for r in ws.iter_rows(min_row=3, values_only=True)]
    return ws, rows


def _xl(d: dt.date) -> dt.datetime:
    return dt.datetime.combine(d, dt.time())


def test_company_original_sheet_gives_a_clean_import_format(tmp_path):
    """公司原始待建檔表：第 1 列公式、格式化條件、資料驗證、篩選都在，仍產生核對結果檔，「回填後」是乾淨的匯入格式。"""
    spec = Spec()
    row = reference_row(spec)
    row.update(
        {
            "TS/IIS完成檢核": "待檢核",
            "交易日期": row["交易日"],
            "計價幣別": row["承作幣別"],
            "分公司_1": "城中",
            "分公司_1_承作金額": 340000,
            "下一比價日": "-",
        }
    )
    sheet = build_reference_sheet(tmp_path / "待建檔.xlsx", [row], IMPORT_HEADERS, title="ELN待建檔_一般", header_row=2)
    wb = openpyxl.load_workbook(sheet)
    ws = wb.active
    ws["A1"], ws["B1"], ws["F1"] = "最新交易日", "=MAX(B3:B20)", '=IF(D1=1,"V","-")'
    ws.conditional_formatting.add("X3:AN20", FormulaRule(formula=["ISBLANK(X3)"], fill=None))
    dv = DataValidation(type="list", formula1='"D,P,AM"')
    dv.add("AU2")
    ws.add_data_validation(dv)
    ws.auto_filter.ref = "A2:CB20"
    memo = ws.cell(3, IMPORT_HEADERS.index("UL_1_Memo") + 1)
    memo.value, memo.data_type = "=見附件", "s"  # 「=」開頭的文字不是公式
    ws.column_dimensions["D"].width = 14
    wb.save(sheet)

    outcome, receipt = run(tmp_path, sheet, [spec])
    item = outcome.items[0]
    assert item.report.status == CheckStatus.PASS and receipt.filled(item)
    out, rows = filled_sheet(receipt)
    assert [c.value for c in out[1]][:2] == ["最新交易日", _xl(spec.trade_date)]
    assert [c.value for c in out[2]] == IMPORT_HEADERS
    [d] = rows
    assert d["TS/IIS完成檢核"] == "V", "待檢核 → V"
    assert (d["TDCC Code"], d["ISIN Code"], d["計價幣別"]) == (spec.product_code, SYNTH_ISIN, row["承作幣別"])
    assert (d["交易日期"], d["發行日"]) == (_xl(spec.trade_date), _xl(spec.issue_date))
    assert (d["分公司_1"], d["分公司_1_承作金額"], d["下一比價日"]) == ("城中", 340000, "-")
    assert not any(c.data_type == "f" for r in out.iter_rows() for c in r), "不帶公式"
    assert d["UL_1_Memo"] == "=見附件", "「=」開頭的文字照文字寫"
    assert not list(out.conditional_formatting) and not out.data_validations.dataValidation
    assert out.auto_filter.ref is None
    assert out.column_dimensions["D"].width == 14, "欄寬取同名欄（TDCC Code）"


def test_template_layout_gives_the_same_import_format(tmp_path):
    """作業人員模板：表頭第 3 列、上面兩列是資料日期與欄號，多出評價日等欄，沒有 TS/IIS完成檢核、交易日期。"""
    spec = Spec()
    headers = [h for h in REFERENCE_HEADERS if h not in ("TS", "IIS")] + ["評價日", "KI紀錄", "U1_LAST"]
    row = reference_row(spec) | {"評價日": "-", "KI紀錄": "N", "U1_LAST": 123.45}
    sheet = build_reference_sheet(tmp_path / "模板.xlsx", [row], headers, title="工作表1")
    wb = openpyxl.load_workbook(sheet)
    wb.active["A1"], wb.active["B1"] = "資料日期", dt.datetime(2030, 1, 9)
    for k in range(7, len(headers) + 1):
        wb.active.cell(2, k, k)
    wb.save(sheet)

    outcome, receipt = run(tmp_path, sheet, [spec])
    assert receipt.filled(outcome.items[0])
    out, [d] = filled_sheet(receipt)
    assert [c.value for c in out[1]][:2] == ["最新交易日", _xl(spec.trade_date)]
    assert [c.value for c in out[2]] == IMPORT_HEADERS, "多出的欄（評價日、發行機構、庫存狀態…）不帶入"
    assert (d["TS/IIS完成檢核"], d["交易日期"], d["計價幣別"]) == ("V", _xl(spec.trade_date), row["承作幣別"])
    assert (d["分公司_1"], d["下一比價日"]) == (None, None), "表上沒有的欄留空白"


def test_only_passing_rows_in_sheet_order_and_the_latest_kept_trade_date(tmp_path):
    """舊 FCN參考條件：只留通過的列、順序照表、同一商品一列；最新交易日只看留下的列；回填日期沿用表上日期格式。"""
    a, b, c = (Spec(product_code=f"02919999000{n}") for n in (1, 2, 3))
    later = dt.datetime.combine(a.trade_date + dt.timedelta(days=30), dt.time())
    rows = [reference_row(a), reference_row(b, 交易日=later), reference_row(c)]  # b 交易日對不起來、不回填
    sheet = build_reference_sheet(tmp_path / "FCN參考條件.xlsx", rows)

    outcome, receipt = run(tmp_path, sheet, [c, b, a])
    out, kept = filled_sheet(receipt)
    assert [d["TDCC Code"] for d in kept] == [a.product_code, c.product_code]
    assert out["B1"].value == _xl(a.trade_date), "不通過的列不算最新交易日"
    issue = IMPORT_HEADERS.index("發行日") + 1
    assert out.cell(3, issue).number_format == DATE_FORMAT, "回填的日期沿用表上既有日期格式"
