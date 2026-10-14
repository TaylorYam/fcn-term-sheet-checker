"""本機 BNP 真實樣本（Issue #191、#197、#222）；交易資料不進 Git，CI 自動略過。測試碼不含真實商品代號或數值。

參考條件表用公司原始待建檔格式、有 BNP 列的 `核對表格式..xlsx`（`real_data.BNP_REFERENCE`，回填欄位空白）。
斷言與探勘結論一致（docs/rules/bnp-check-rules.md §5）：6 組說明書＋投資人須知都辨識為 BNP、都有參考條件表列；
只有一組有差異（一檔標的的價格欄疑似分割調整，說明書與投資人須知都轉人工覆核；同一份說明書情境一累積配息與每期配息
× 期數不符），這組不回填；其餘 5 組兩份都通過才回填，回填值與說明書一致。
投資人須知（docs/templates/bnp-zh-iis.md）每個範本欄位都讀得到，和同商品說明書讀出的值一致。
"""

import datetime as dt
import shutil
from collections import Counter
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import fitz
import openpyxl
import pytest

from fcn_checker.batch import DocKind
from fcn_checker.extraction import extract_lines
from fcn_checker.issuers import BNP, detect, detect_iis
from fcn_checker.saving import run_batch
from fcn_checker.schema import CheckStatus as S
from fcn_checker.standard_fields import Occurrence
from fcn_checker.text import squash
from harness import CONFIG
from real_data import BNP_REFERENCE, TERM_SHEETS, pdfs, requires

pytestmark = [
    pytest.mark.real_samples,
    requires(TERM_SHEETS / "082*_TS.pdf", BNP_REFERENCE),
]
PROBLEMS = (S.MISMATCH, S.REVIEW_REQUIRED, S.ERROR)
BACKFILL_COLUMNS = ("ISIN Code", "發行日", *(f"比價日_{i}" for i in range(1, 13)))
EMPTY = CONFIG.reference_format.empty_value  # 不需要填的比價日寫空值寫法
SPLIT = Counter({("iis.underlying_prices", "split_suspected"): 3})  # 一檔標的的進場價、執行價、KO 價


@pytest.fixture(scope="module")
def saved(tmp_path_factory):
    """BNP 說明書與投資人須知整批核對並儲存一次：(批量核對結果, 儲存收據)。"""
    tmp = tmp_path_factory.mktemp("real_bnp")
    sheet = shutil.copy(BNP_REFERENCE, tmp / BNP_REFERENCE.name)  # 不在 data/ 產生檔案
    return run_batch(CONFIG, Path(sheet), pdfs("082*.pdf"), tmp / "reports", root=tmp)


@pytest.fixture(scope="module")
def items(saved):
    outcome, _ = saved
    return [i for i in outcome.items if i.kind == DocKind.TERM_SHEET], [
        i for i in outcome.items if i.kind == DocKind.IIS
    ]


def problems(item) -> Counter[tuple[str, str]]:
    return Counter((r.rule_id, r.reason_code) for r in item.report.results if r.status in PROBLEMS)


def test_every_sample_is_recognised_as_bnp(items):
    term_sheets, investor_sheets = items
    assert len(term_sheets) == len(investor_sheets) == 6
    assert all(i.report.template == BNP.template_id for i in term_sheets)
    assert all(i.report.template == BNP.iis.template_id and not i.unsupported for i in investor_sheets)
    assert all(i.report.results and i.partner is not None for i in (*term_sheets, *investor_sheets)), "都有列、都成對"


def test_samples_only_differ_where_the_exploration_said(items):
    term_sheets, _ = items
    bad = sum((problems(i) for i in term_sheets), Counter())
    assert bad == Counter(
        {
            (
                "field.underlying_prices",
                "split_suspected",
            ): 3,  # 一檔標的的進場價、執行價、KO 價都是說明書的同一比例（分割調整）
            (
                "doc.scenario_calculations",
                "value_mismatch",
            ): 1,  # 同一份情境一累積總月配息金額 ≠ 每期配息 × 期數（文件錯誤）
        }
    )
    passed = [i for i in term_sheets if i.report.status == S.PASS]
    assert len(passed) == 5  # 只有 1 份（分割調整＋累積配息）不是全部通過
    # 封面與第二章的受託機構電話、地址都是審查標準的等價寫法
    for field in ("distributor_phone_cover", "distributor_address_cover", "distributor_address_ch2"):
        found = [r for i in term_sheets for r in i.report.results if r.field == field]
        assert len(found) == 6 and all(r.status == S.PASS and "equivalents" in (r.tolerance or "") for r in found), (
            field
        )
    # Non-Call = 天期的 1 份：KO % 與各標的 KO 價照常核對（與 MS 不同）
    assert all(any(r.rule_id == "field.ko_pct" and r.status == S.PASS for r in i.report.results) for i in term_sheets)


def test_bnp_investor_sheets_only_differ_by_the_split_adjusted_prices(items):
    """分割調整那一檔，投資人須知也和說明書一樣轉人工覆核（Issue #222），其餘 5 份全部通過。"""
    term_sheets, investor_sheets = items
    flagged = [i for i in investor_sheets if problems(i)]
    assert len(flagged) == 1 and problems(flagged[0]) == SPLIT
    assert all("分割" in m for m in flagged[0].problem_messages)
    split_ts = next(i for i in term_sheets if problems(i)[("field.underlying_prices", "split_suspected")])
    assert flagged[0].product_code == split_ts.product_code
    assert sum(i.report.status == S.PASS for i in investor_sheets) == 5


def test_only_products_whose_two_documents_pass_are_filled(saved, items):
    outcome, receipt = saved
    term_sheets, _ = items
    filled = [i for i in term_sheets if receipt.filled(i)]
    assert len(filled) == 5
    assert all(i.report.status == S.PASS and i.partner.report.status == S.PASS for i in filled)
    got = rows_by_code(receipt.output)
    assert set(got) == {i.product_code for i in filled}, "「回填後」只有兩份都通過的 5 列"
    for item in filled:
        row, ts = got[item.product_code], BNP.read(_lines(item.term_sheet))
        assert row["TS/IIS完成檢核"] == "V"
        assert row["ISIN Code"] == ts.f("isin").value
        assert as_date(row["發行日"]) == ts.f("issue_date").value
        decisions = {d.column: d for d in item.report.backfill}
        assert all(decisions[c].sheet_value is None and decisions[c].action == "fill" for c in BACKFILL_COLUMNS)
        # 比價日依說明書提前出場排程：D 型填第一個可提前出場期與最後一期、P 型從第一個可提前出場期起每期都填
        sched = ts.f("autocall_schedule").value
        assert sched.dates[sched.periods] == ts.f("final_valuation_date").value
        last = sched.periods + 1
        periods = (
            {sched.first_callable, sched.periods} if sched.observation == "D" else range(sched.first_callable, last)
        )
        want = [sched.dates[n] if n in periods else EMPTY for n in range(1, 13)]
        assert [as_date(row[f"比價日_{n}"]) for n in range(1, 13)] == want, item.term_sheet.name


def test_vwap_rows_take_the_term_sheet_prices(saved, items):
    """期初定價 VWAP 的 1 列：表上進場價空白，回填後各標的價格 = 說明書價格表（Issue #122）。"""
    _, receipt = saved
    term_sheets, _ = items
    got = rows_by_code(receipt.output)
    vwap = [i for i in term_sheets if got.get(i.product_code, {}).get("期初定價") == "VWAP"]
    assert len(vwap) == 1
    for item in vwap:
        row, table = got[item.product_code], BNP.read(_lines(item.term_sheet)).f("underlying_prices").value
        for n, prices in enumerate((r.prices for r in table), 1):
            for col, name in (("initial", "進場價"), ("strike", "執行價"), ("ki", "下限價"), ("ko", "KO價")):
                assert rounded(row[f"UL_{n}_{name}"]) == prices.get(col, EMPTY), (n, name)  # 無 KI：下限價為空值寫法


def rows_by_code(path: Path) -> dict[str, dict]:
    """「回填後」（匯入格式，第 2 列表頭）：TDCC Code → 該列。"""
    ws = openpyxl.load_workbook(path, data_only=True)["回填後"]
    headers = [c.value for c in ws[2]]
    rows = (dict(zip(headers, r, strict=False)) for r in ws.iter_rows(min_row=3, values_only=True))
    return {str(d["TDCC Code"]): d for d in rows if d.get("TDCC Code")}


def as_date(v):
    return v.date() if isinstance(v, dt.datetime) else v


def rounded(v):
    return Decimal(repr(round(v, 9))).quantize(Decimal("0.0001"), ROUND_HALF_UP) if isinstance(v, float | int) else v


def _lines(path):
    doc = fitz.open(path)
    try:
        return extract_lines(doc)
    finally:
        doc.close()


SAME_AS_TERM_SHEET = (  # 投資人須知欄位 → 同商品說明書欄位（探勘結論：6 份全部相同）
    ("isin", "isin"),
    ("currency_zh", "currency_zh"),
    ("denomination", "denomination"),
    ("underlyings", "underlyings"),
    ("strike_pct", "strike_pct"),
    ("ko_pct", "ko_pct"),
    ("tenor_months", "tenor_months"),
    ("trade_date", "trade_date"),
    ("issue_date", "issue_date"),
    ("final_valuation_date", "final_valuation_date"),
    ("maturity_date", "maturity_date"),
    ("monthly_coupon_shown_pct", "monthly_coupon_pct"),
    ("ko_observation", "ko_observation"),
    ("ko_memory", "ko_memory"),
    ("ki_type", "ki_type"),
)


def test_every_iis_field_is_read_and_matches_the_term_sheet():
    pairs = [(p, p.with_name(p.name.replace("_IIS", "_TS"))) for p in pdfs("082*_IIS.pdf")]
    assert len(pairs) == 6
    for iis_pdf, ts_pdf in pairs:
        sheet, ts = BNP.iis.read(_lines(iis_pdf)), BNP.read(_lines(ts_pdf))
        for name in sheet.provided:
            f = sheet.f(name)
            assert f.ok, (iis_pdf.name, name, f.note)
            if isinstance(f.value, tuple) and f.value and isinstance(f.value[0], Occurrence):
                assert all(o.value.ok for o in f.value), (iis_pdf.name, name)
        for mine, theirs in SAME_AS_TERM_SHEET:
            assert sheet.f(mine).value == ts.f(theirs).value, (iis_pdf.name, mine)
        assert squash(sheet.f("name_zh").value) == squash(ts.f("name_zh").value)
        assert squash(sheet.f("name_en").value) == squash(ts.f("name_en").value)
        assert [r.prices for r in sheet.f("underlying_prices").value] == [
            r.prices for r in ts.f("underlying_prices").value
        ]
        if sheet.provides("first_callable_period"):  # 只有 D 型寫出
            assert sheet.f("first_callable_period").value == ts.f("first_callable_period").value
        if sheet.provides("ki_pct"):  # 有 KI 時才有下限價格列
            assert sheet.f("ki_pct").value == ts.f("ki_pct").value


def test_no_other_sample_is_detected_as_a_bnp_investor_sheet():
    """BNP 說明書與其他上手的說明書、投資人須知都不符合 BNP 投資人須知範本。"""
    bnp_iis = {p.name for p in pdfs("082*_IIS.pdf")}
    others = [p for p in pdfs() if p.name not in bnp_iis]
    assert len(others) > 40
    for path in others:
        issuer, _ = detect_iis(_lines(path), (BNP,))
        assert issuer is None, path.name


def test_no_other_sample_is_detected_as_a_bnp_term_sheet(items):
    """BARC、HSBC、MS、SG、UBS 等說明書與所有投資人須知（含 BNP 投資人須知）都不符合 BNP 範本。"""
    term_sheets, _ = items
    bnp = {i.term_sheet.name for i in term_sheets}
    others = [p for p in pdfs() if p.name not in bnp]
    assert len(others) > 40
    for path in others:
        doc = fitz.open(path)
        try:
            lines = extract_lines(doc)
        finally:
            doc.close()
        issuer, _ = detect(lines, (BNP,))
        assert issuer is None, path.name
