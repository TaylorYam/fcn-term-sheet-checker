"""測試切點：批量核對入口 run_batch(核對設定, 參考條件表, 說明書們, 輸出資料夾) → 結果與儲存收據、核對結果檔與核對紀錄；
PANEL 的用法是 preview_batch → check_batch → save_batch。儲存的結果只看寫出的檔案與收據。

只用合成資料（tests/synth.py）；以 openpyxl 讀回產出的 Excel 觀察回填結果。
"""

from __future__ import annotations

import datetime as dt
import json
from pathlib import Path

import openpyxl
import pytest
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation

from fcn_checker import __version__
from fcn_checker.batch import check_batch, preview_batch
from fcn_checker.ingestion import IngestionError, sha256_of
from fcn_checker.issuers import BARC
from fcn_checker.saving import run_batch, save_batch
from fcn_checker.schema import CheckStatus, DetectionResult
from harness import CONFIG, ISSUER_PREFIXES, REVIEW_STANDARD, check_all, load_config, with_iis
from reference_synth import DATE_FORMAT, REFERENCE_FORMAT, REFERENCE_HEADERS, build_reference_sheet, reference_row
from synth import SYNTH_ISIN, Spec, barc_adapter, build_iis_pdf, build_pdf, schedule_rows

PASS, MISMATCH, REVIEW, ERROR = (
    CheckStatus.PASS,
    CheckStatus.MISMATCH,
    CheckStatus.REVIEW_REQUIRED,
    CheckStatus.ERROR,
)
NOW = dt.datetime(2030, 2, 3, 4, 5, 6)


def pdf_for(tmp_path: Path, spec: Spec, name: str | None = None) -> Path:
    return build_pdf(tmp_path / (name or f"{spec.product_code}_TS.pdf"), spec)


def batch(tmp_path: Path, pdfs: list[Path], rows: list[dict], *, sheet: Path | None = None, config=CONFIG):
    sheet = sheet or build_reference_sheet(tmp_path / "FCN參考條件.xlsx", rows)
    outcome, receipt = run_batch(config, sheet, with_iis(pdfs), tmp_path / "reports", root=tmp_path, now=NOW)
    return outcome, receipt, sheet


def sheet_rows(path: Path, sheet: str = "回填後") -> dict[str, dict]:
    """TDCC Code → 該列（第 3 列表頭、第 4 列起資料）；預設讀核對結果檔的「回填後」。"""
    ws = openpyxl.load_workbook(path)[sheet]
    headers = [c.value for c in ws[3]]
    rows = (dict(zip(headers, r, strict=False)) for r in ws.iter_rows(min_row=4, values_only=True))
    return {d["TDCC Code"]: d for d in rows if d.get("TDCC Code")}


def row_of(path: Path, product_code: str) -> dict:
    rows = sheet_rows(path)
    assert product_code in rows, f"「回填後」找不到 {product_code}"
    return rows[product_code]


def error_rows(path: Path) -> list[dict]:
    ws = openpyxl.load_workbook(path)["錯誤清單"]
    rows = list(ws.iter_rows(values_only=True))
    return [dict(zip(rows[0], r, strict=True)) for r in rows[1:]]


def header_cell(ws, header: str, row: int):
    return ws.cell(row, next(c.column for c in ws[3] if c.value == header))


def problems(item) -> list[tuple[str, str, str]]:
    return [(r.rule_id, r.field, r.message) for r in item.report.results if r.status in (MISMATCH, REVIEW, ERROR)]


def only(item, rule_id: str, field: str | None = None):
    out = [r for r in item.report.results if r.rule_id == rule_id and (field is None or r.field == field)]
    assert len(out) == 1, f"{rule_id} {field or ''} 應恰好一筆：{out}"
    return out[0]


def slots(d: dict) -> list:
    return [v.date() if isinstance(v, dt.datetime) else v for v in (d[f"比價日_{i}"] for i in range(1, 13))]


# ---------------------------------------------------------------- 全部一致 → 回填


def test_consistent_daily_term_sheet_passes_and_back_fills_first_and_last_compare_dates(tmp_path):
    spec = Spec()  # D 型、天期 6、第 1 期期末日起可提前出場
    outcome, receipt, sheet = batch(tmp_path, [pdf_for(tmp_path, spec)], [reference_row(spec)])

    item = outcome.items[0]
    assert item.report.status == PASS, problems(item)
    assert receipt.filled(item)
    assert receipt.output == tmp_path / "reports" / "FCN參考條件_核對結果_20300203-040506.xlsx"
    d = row_of(receipt.output, spec.product_code)
    assert d["ISIN Code"] == SYNTH_ISIN
    assert d["發行日"] == dt.datetime.combine(spec.issue_date, dt.time())
    assert sheet_rows(sheet, "樣本清單")[spec.product_code]["發行日"] is None, "原檔不動"
    first = schedule_rows(spec)[0]["valuation"]
    assert slots(d) == [first] + ["-"] * 4 + [spec.final_date] + ["-"] * 6, "D 型：Non-Call 那期與最後一期"


def test_daily_with_non_call_equal_to_tenor_fills_only_the_last_period(tmp_path):
    spec = Spec(guaranteed=6)  # 第 6 期期末日起才可提前出場 → Non-Call = 天期
    outcome, receipt, _ = batch(tmp_path, [pdf_for(tmp_path, spec)], [reference_row(spec)])

    assert outcome.items[0].report.status == PASS, problems(outcome.items[0])
    assert slots(row_of(receipt.output, spec.product_code)) == ["-"] * 5 + [spec.final_date] + ["-"] * 6


def test_latest_compare_date_must_equal_final_valuation_date(tmp_path):
    spec = Spec(ko_overrides={(6, "end"): "2030 年7 月9 日"})  # 最後一期期末日晚於最終評價日 2030-07-08
    outcome, receipt, _ = batch(tmp_path, [pdf_for(tmp_path, spec)], [reference_row(spec)])

    item = outcome.items[0]
    r = only(item, "backfill.compare_dates")
    assert (r.status, r.reason_code) == (REVIEW, "compare_dates_max_mismatch")
    assert "2030-07-09" in r.message and "2030-07-08" in r.message
    assert not receipt.filled(item)
    assert spec.product_code not in sheet_rows(receipt.output)


def test_period_end_back_fills_every_compare_date_from_first_callable_period(tmp_path):
    spec = Spec(ko_obs="P", memory=False, guaranteed=2)  # 前 2 期不可提前出場 → Non-Call = 3
    outcome, receipt, _ = batch(tmp_path, [pdf_for(tmp_path, spec)], [reference_row(spec)])

    item = outcome.items[0]
    assert item.report.status == PASS, problems(item)
    valuation = [r["valuation"] for r in schedule_rows(spec)]
    assert slots(row_of(receipt.output, spec.product_code)) == ["-", "-", *valuation[2:6]] + ["-"] * 6


def test_period_end_with_non_call_equal_to_tenor_fills_only_the_last_period(tmp_path):
    spec = Spec(ko_obs="P", memory=False, guaranteed=5)  # 只有第 6 期可提前出場
    outcome, receipt, _ = batch(tmp_path, [pdf_for(tmp_path, spec)], [reference_row(spec)])

    assert outcome.items[0].report.status == PASS, problems(outcome.items[0])
    last = schedule_rows(spec)[-1]["valuation"]
    assert slots(row_of(receipt.output, spec.product_code)) == ["-"] * 5 + [last] + ["-"] * 6


def test_period_end_memory_uses_autocall_valuation_dates(tmp_path):
    spec = Spec(ko_obs="P", memory=True)  # 自動提前出場評價日表，每期都可提前出場
    outcome, receipt, _ = batch(tmp_path, [pdf_for(tmp_path, spec)], [reference_row(spec)])

    assert outcome.items[0].report.status == PASS, problems(outcome.items[0])
    valuation = [r["valuation"] for r in schedule_rows(spec)]
    assert slots(row_of(receipt.output, spec.product_code)) == valuation + ["-"] * 6


def test_already_filled_matching_values_pass_and_stay_unchanged(tmp_path):
    spec = Spec()
    first = schedule_rows(spec)[0]["valuation"]
    issue = dt.datetime.combine(spec.issue_date, dt.time())
    last = dt.datetime.combine(spec.final_date, dt.time())
    filled = {"TS": "V", "IIS": "V", "ISIN Code": SYNTH_ISIN, "發行日": issue}
    filled |= {"比價日_1": dt.datetime.combine(first, dt.time())}
    filled |= {f"比價日_{i}": "-" for i in range(2, 13)} | {"比價日_6": last}
    outcome, receipt, _ = batch(tmp_path, [pdf_for(tmp_path, spec)], [reference_row(spec, **filled)])

    item = outcome.items[0]
    assert item.report.status == PASS, problems(item)
    assert {d.action for d in item.report.backfill} == {"match"}
    d = row_of(receipt.output, spec.product_code)
    assert (d["TS"], d["IIS"], d["ISIN Code"], d["發行日"]) == ("V", "V", SYNTH_ISIN, issue)
    assert slots(d) == [first] + ["-"] * 4 + [spec.final_date] + ["-"] * 6


def test_wrong_existing_compare_date_is_mismatch_and_nothing_is_back_filled(tmp_path):
    spec = Spec()
    wrong = dt.datetime(2030, 1, 1)
    outcome, receipt, _ = batch(
        tmp_path, [pdf_for(tmp_path, spec)], [reference_row(spec, **{"比價日_1": wrong, "比價日_6": "-"})]
    )

    item = outcome.items[0]
    r = only(item, "backfill.compare_dates")
    assert r.status == MISMATCH
    assert item.report.status == MISMATCH and not receipt.filled(item)
    assert spec.product_code not in sheet_rows(receipt.output), "沒通過就不回填，也不出現在「回填後」"


def test_dash_where_a_compare_date_is_expected_is_mismatch(tmp_path):
    spec = Spec()
    outcome, receipt, _ = batch(tmp_path, [pdf_for(tmp_path, spec)], [reference_row(spec, **{"比價日_1": "-"})])
    assert only(outcome.items[0], "backfill.compare_dates").status == MISMATCH


def test_wrong_isin_is_mismatch_and_keeps_original(tmp_path):
    spec = Spec()
    outcome, receipt, _ = batch(
        tmp_path, [pdf_for(tmp_path, spec)], [reference_row(spec, **{"ISIN Code": "XS9999999999"})]
    )
    r = only(outcome.items[0], "backfill.isin")
    assert (r.status, r.expected, r.actual) == (MISMATCH, "XS9999999999", SYNTH_ISIN)
    assert spec.product_code not in sheet_rows(receipt.output)


def test_wrong_issue_date_is_mismatch_and_keeps_original(tmp_path):
    spec = Spec()
    wrong = dt.datetime(2030, 1, 15)
    outcome, receipt, _ = batch(tmp_path, [pdf_for(tmp_path, spec)], [reference_row(spec, 發行日=wrong)])
    item = outcome.items[0]
    r = only(item, "backfill.issue_date")
    assert (r.status, r.expected, r.actual) == (MISMATCH, wrong.date(), spec.issue_date)
    assert [d.action for d in item.report.backfill if d.column == "發行日"] == ["mismatch"]
    assert item.report.status == MISMATCH and not receipt.filled(item)
    assert spec.product_code not in sheet_rows(receipt.output)


def test_issue_date_missing_from_term_sheet_requires_review_and_is_not_back_filled(tmp_path):
    spec = Spec(omit=frozenset({"issue_date"}))
    outcome, receipt, _ = batch(tmp_path, [pdf_for(tmp_path, spec)], [reference_row(spec)])
    item = outcome.items[0]
    assert only(item, "backfill.issue_date").status == REVIEW
    assert item.status_label == "需人工覆核", "不是配對問題時只顯示中文狀態"
    assert not [d for d in item.report.backfill if d.column == "發行日"]
    assert not receipt.filled(item)
    assert spec.product_code not in sheet_rows(receipt.output)


def test_non_call_must_be_the_first_callable_period(tmp_path):
    spec = Spec()  # 第 1 期可提前出場
    outcome, receipt, _ = batch(tmp_path, [pdf_for(tmp_path, spec)], [reference_row(spec, **{"Non-Call(月)": 0})])
    r = only(outcome.items[0], "field.first_callable_period")
    assert (r.status, r.expected, r.actual) == (MISMATCH, 0, 1)
    assert not receipt.filled(outcome.items[0])


def test_daily_observation_from_first_day_requires_review_without_back_fill(tmp_path):
    spec = Spec(guaranteed=0)  # 第 1 期期始日就有日期（S08、S10 型），比價日填法未確認
    outcome, receipt, _ = batch(tmp_path, [pdf_for(tmp_path, spec)], [reference_row(spec, **{"Non-Call(月)": 1})])

    item = outcome.items[0]
    assert only(item, "backfill.compare_dates").status == REVIEW
    assert only(item, "field.first_callable_period").status == REVIEW
    assert item.report.status == REVIEW and not receipt.filled(item)


def test_more_than_twelve_periods_requires_review(tmp_path):
    spec = Spec(tenor=13, final_date=dt.date(2031, 2, 10), maturity_date=dt.date(2031, 2, 13))
    outcome, receipt, _ = batch(tmp_path, [pdf_for(tmp_path, spec)], [reference_row(spec)])
    r = only(outcome.items[0], "backfill.compare_dates")
    assert (r.status, r.reason_code) == (REVIEW, "compare_dates_unmapped")
    assert not receipt.filled(outcome.items[0])


# ---------------------------------------------------------------- 上手編號與配對


def prefixes_file(tmp_path: Path, extra: str) -> Path:
    p = tmp_path / "prefixes.toml"
    p.write_text(ISSUER_PREFIXES.read_text(encoding="utf-8") + extra, encoding="utf-8")
    return p


def test_prefix_not_in_table_is_unsupported_issuer(tmp_path):
    spec = Spec(product_code="999199990001")
    outcome, receipt, _ = batch(tmp_path, [pdf_for(tmp_path, spec)], [reference_row(spec)])

    item = outcome.items[0]
    r = only(item, "batch.issuer_prefix")
    assert (r.status, r.reason_code) == (REVIEW, "issuer_unsupported")
    assert item.unsupported and item.report.status == REVIEW and item.status_label == "未支援上手"
    assert not any(x.rule_id.startswith("field.") for x in item.report.results)
    errors = error_rows(receipt.output)
    assert [e["PDF 檔名"] for e in errors] == [f"{spec.product_code}_TS.pdf", f"{spec.product_code}_IIS.pdf"]
    assert errors[0]["TDCC Code"] == spec.product_code, "取不到封面商品代號時用檔名前 12 碼"
    assert all("未支援上手" in e["錯訊"] for e in errors)


def test_issuer_in_prefix_table_without_template_is_unsupported(tmp_path):
    spec = Spec(product_code="325199990001")  # 325 = HSBC：對照表有，本測試 registry 刻意不註冊 HSBC
    outcome, receipt, _ = batch(
        tmp_path,
        [pdf_for(tmp_path, spec)],
        [reference_row(spec, 發行機構="HSBC")],
        config=CONFIG.with_registry((BARC,)),
    )
    item = outcome.items[0]
    assert item.unsupported and item.issuer == "HSBC"
    assert "HSBC" in only(item, "batch.issuer_prefix").message


def test_registered_issuer_without_a_template_names_the_issuer(tmp_path):
    """上手編號對照登記了、但還沒有範本的上手（例：037 = SG，Issue #149）：未支援上手，錯訊寫出上手代號。"""
    spec = Spec(product_code="037199990001")
    outcome, _, _ = batch(tmp_path, [pdf_for(tmp_path, spec)], [reference_row(spec, 發行機構="SG")])
    for item in outcome.items:
        r = only(item, "batch.issuer_prefix")
        assert item.unsupported and item.issuer == "SG"
        assert r.message.startswith("上手編號 037 對應 SG，但 SG 還沒有"), r.message


def test_term_sheet_content_of_another_issuer_requires_review(tmp_path):
    other = barc_adapter(code="FAKE", template_id="fake-zh-pd", detect=lambda lines: DetectionResult(False))
    spec = Spec(product_code="777199990001")
    outcome, receipt, _ = batch(
        tmp_path,
        [pdf_for(tmp_path, spec)],
        [reference_row(spec)],
        config=load_config(issuer_prefixes=prefixes_file(tmp_path, '"777" = "FAKE"\n'), registry=(BARC, other)),
    )
    r = only(outcome.items[0], "batch.issuer_prefix")
    assert (r.status, r.reason_code, r.actual) == (REVIEW, "issuer_prefix_mismatch", "BARC")
    assert outcome.items[0].status_label == "檔名上手編號不符"


@pytest.mark.parametrize(
    "name", ["029 任意檔名.pdf", "029199990001_ TS.pdf", "029199990001_ts.pdf", "029199990001.pdf"]
)
def test_file_name_not_ending_with_ts_or_iis_cannot_be_recognized_or_released(tmp_path, name):
    spec = Spec()
    pdf = pdf_for(tmp_path, spec, name=name)
    iis_pdf = build_iis_pdf(tmp_path / f"{spec.product_code}_IIS.pdf", spec)
    outcome, receipt, _ = batch(tmp_path, [pdf, iis_pdf], [reference_row(spec)])

    item, iis = outcome.items
    r = only(item, "batch.file_name")
    assert (r.status, r.reason_code) == (REVIEW, "file_name_unrecognized")
    assert item.kind is None and item.status_label == "檔名無法辨識"
    assert not any(x.rule_id.startswith(("field.", "backfill.")) for x in item.report.results)
    assert "檔名無法辨識" in item.release_problem
    with pytest.raises(IngestionError, match="檔名無法辨識"):
        outcome.release(item)
    assert iis.status_label == "這批缺說明書", "辨識不了的檔案不算同商品的說明書"
    assert [e["PDF 檔名"] for e in error_rows(receipt.output)] == [name, iis.term_sheet.name]
    assert sheet_rows(receipt.output) == {}


def test_product_code_prefix_differs_from_file_name_requires_review(tmp_path):
    spec = Spec(product_code="028199990001")  # 檔名 029、說明書商品代號 028
    outcome, receipt, _ = batch(tmp_path, [pdf_for(tmp_path, spec, name="029199990001_TS.pdf")], [reference_row(spec)])
    r = only(outcome.items[0], "batch.issuer_prefix")
    assert (r.status, r.reason_code, r.actual) == (REVIEW, "issuer_prefix_mismatch", "028199990001")
    assert outcome.items[0].status_label == "檔名上手編號不符"


def test_missing_or_duplicate_reference_row_requires_review(tmp_path):
    spec = Spec()
    pdf = pdf_for(tmp_path, spec)
    outcome, receipt, _ = batch(tmp_path, [pdf], [reference_row(Spec(product_code="029199990002"))])
    assert only(outcome.items[0], "batch.pairing").reason_code == "reference_row_missing"
    assert outcome.items[0].status_label == "條件表找不到這筆", "狀態欄直接寫原因，不顯示英文代碼"

    (tmp_path / "dup").mkdir()
    outcome, receipt, _ = batch(tmp_path / "dup", [pdf], [reference_row(spec), reference_row(spec)])
    r = only(outcome.items[0], "batch.pairing")
    assert (r.status, r.reason_code) == (REVIEW, "reference_row_duplicate")
    assert len(r.order_source) == 2
    assert outcome.items[0].status_label == "條件表有重複列"


def test_pdfs_sharing_one_reference_row_all_require_review_and_others_still_back_fill(tmp_path):
    spec, other = Spec(), Spec(product_code="029199990002")
    (tmp_path / "old").mkdir()
    old = build_pdf(tmp_path / "old" / f"{spec.product_code}_TS.pdf", spec, iis=False)
    new = pdf_for(tmp_path, spec)  # 旁邊有同商品投資人須知
    ok = pdf_for(tmp_path, other)
    outcome, receipt, _ = batch(tmp_path, [old, ok, new], [reference_row(spec), reference_row(other)])

    first, second, third, ok_iis, new_iis = outcome.items
    for item, another in ((first, new), (third, old)):
        r = only(item, "batch.pairing")
        assert (r.status, r.reason_code) == (REVIEW, "reference_row_shared")
        assert "同一批有多份說明書對到同一個 TDCC Code" in r.message and str(another) in r.message
        assert item.report.status == REVIEW and not receipt.filled(item)
        assert item.status_label == "多份對到同一列"
        assert not any(x.rule_id.startswith(("field.", "backfill.")) for x in item.report.results)
    assert new_iis.status_label == "多份對到同一列", "同一列的投資人須知也無法確定配哪一份說明書"
    assert second.report.status == PASS and receipt.filled(second) and second.status_label == "通過"
    assert ok_iis.report.status == PASS and second.partner is ok_iis
    assert list(sheet_rows(receipt.output)) == [other.product_code]
    assert row_of(receipt.output, other.product_code)["ISIN Code"] == SYNTH_ISIN
    errors = error_rows(receipt.output)
    assert [e["PDF 檔名"] for e in errors] == [old.name, new.name, new_iis.term_sheet.name]
    assert str(new) in errors[0]["錯訊"] and str(old) in errors[1]["錯訊"]


def test_same_pdf_selected_twice_says_so_and_same_names_show_full_paths(tmp_path):
    spec = Spec()
    pdf = pdf_for(tmp_path, spec)
    outcome, receipt, _ = batch(tmp_path, [pdf, pdf], [reference_row(spec)])
    for item in outcome.items[:2]:
        assert f"{pdf.name}（同一個檔案重複選取）" in only(item, "batch.pairing").message

    (tmp_path / "v2").mkdir()
    copy = pdf_for(tmp_path / "v2", spec)
    outcome, receipt, _ = batch(tmp_path / "v2", [pdf, copy], [reference_row(spec)])
    assert str(copy) in only(outcome.items[0], "batch.pairing").message
    assert str(pdf) in only(outcome.items[1], "batch.pairing").message


def test_reference_row_of_another_issuer_requires_review(tmp_path):
    spec = Spec()
    outcome, receipt, _ = batch(tmp_path, [pdf_for(tmp_path, spec)], [reference_row(spec, 發行機構="HSBC")])
    r = only(outcome.items[0], "batch.pairing")
    assert (r.status, r.reason_code, r.expected, r.actual) == (REVIEW, "reference_issuer_mismatch", "Barclays", "HSBC")
    assert not any(x.rule_id.startswith("field.") for x in outcome.items[0].report.results)
    assert outcome.items[0].status_label == "條件表發行機構不符"


def test_rows_without_selected_pdf_are_left_out_of_the_result_file(tmp_path):
    spec, other = Spec(), Spec(product_code="029199990002")
    outcome, receipt, sheet = batch(tmp_path, [pdf_for(tmp_path, spec)], [reference_row(spec), reference_row(other)])
    assert receipt.filled(outcome.items[0])
    assert list(sheet_rows(receipt.output)) == [spec.product_code]
    assert error_rows(receipt.output) == []


# ---------------------------------------------------------------- 表上事先填好的欄位


def test_knock_in_term_sheet_checks_ki_columns_and_prices(tmp_path):
    spec = Spec(ki="AM")
    outcome, receipt, _ = batch(tmp_path, [pdf_for(tmp_path, spec)], [reference_row(spec)])
    item = outcome.items[0]
    assert item.report.status == PASS, problems(item)
    prices = [r for r in item.report.results if r.rule_id == "field.underlying_prices"]
    assert len(prices) == 12 and {r.status for r in prices} == {PASS}


def test_prefilled_fields_that_differ_are_mismatches(tmp_path):
    spec = Spec()
    wrong = {
        "KO(Freq)": "P",
        "KO(memo)": "N",
        "KI(%)": 60,
        "單位面額": 5000,
        "UL_2_進場價": 87.25,
        "UL_1_執行價": 86.4151,  # 86.41500 → 4 位小數後仍不同
        "Coupon p.a. (%)": 12.5,
    }
    outcome, receipt, _ = batch(tmp_path, [pdf_for(tmp_path, spec)], [reference_row(spec, **wrong)])
    item = outcome.items[0]
    mismatched = {(r.rule_id, r.field) for r in item.report.results if r.status == MISMATCH}
    assert {
        ("field.ko_observation", "ko_observation"),
        ("field.ko_memory", "ko_memory"),
        ("field.ki_pct", "ki_pct"),
        ("field.denomination", "denomination"),
        ("field.underlying_prices", "ZQH UW 進場價"),
        ("field.underlying_prices", "ZZA UN 執行價"),
        ("field.coupon_pa_pct", "coupon_pa_pct"),
    } <= mismatched
    assert not receipt.filled(item)


def test_prices_with_excel_float_tails_round_to_four_decimals(tmp_path):
    spec = Spec()
    outcome, receipt, _ = batch(
        tmp_path, [pdf_for(tmp_path, spec)], [reference_row(spec, **{"UL_1_執行價": 86.41499999999999})]
    )
    assert outcome.items[0].report.status == PASS, problems(outcome.items[0])


def test_non_default_denomination_still_requires_review_by_standard(tmp_path):
    spec = Spec(denomination=50000)
    outcome, receipt, _ = batch(tmp_path, [pdf_for(tmp_path, spec)], [reference_row(spec)])
    item = outcome.items[0]
    assert only(item, "field.denomination").status == PASS
    assert only(item, "doc.denomination").status == REVIEW
    assert not receipt.filled(item)


def test_unknown_reference_column_requires_review(tmp_path):
    spec = Spec()
    sheet = build_reference_sheet(
        tmp_path / "FCN參考條件.xlsx", [reference_row(spec)], headers=[*REFERENCE_HEADERS, "新欄位"]
    )
    outcome, receipt, _ = batch(tmp_path, [pdf_for(tmp_path, spec)], [], sheet=sheet)
    r = only(outcome.items[0], "order.unknown_column")
    assert r.status == REVIEW and "參考條件表" in r.message and "新欄位" in r.message


# ---------------------------------------------------------------- 核對結果檔


def test_result_file_keeps_only_passing_rows_in_sheet_order_with_the_original_layout(tmp_path):
    a, b, c, d = (Spec(product_code=f"02919999000{n}") for n in (1, 2, 3, 4))
    rows = [reference_row(a), reference_row(b, **{"K(%)": 71}), reference_row(c), reference_row(d)]
    sheet = build_reference_sheet(tmp_path / "FCN參考條件.xlsx", rows)
    wb = openpyxl.load_workbook(sheet)
    ws = wb["樣本清單"]
    ws["A1"] = "FCN 參考條件"
    ws.column_dimensions["E"].width = 17
    for r in range(4, 8):
        ws.row_dimensions[r].height = 20 + r
    wb.save(sheet)
    before = sheet.read_bytes()
    outcome, receipt, _ = batch(tmp_path, [pdf_for(tmp_path, s) for s in (c, b, a)], [], sheet=sheet)  # d 這批沒選到

    assert [receipt.filled(i) for i in outcome.items] == [True, False, True, False, False, False], "投資人須知不回填"
    assert receipt.output == tmp_path / "reports" / "FCN參考條件_核對結果_20300203-040506.xlsx"
    assert sheet.read_bytes() == before, "原檔不動"
    assert list(tmp_path.glob("*_回填_*.xlsx")) == []
    out = openpyxl.load_workbook(receipt.output)
    assert out.sheetnames == ["回填後", "錯誤清單"], "原檔其他工作表不帶入"
    ws, original = out["回填後"], openpyxl.load_workbook(sheet)["樣本清單"]
    assert [[x.value for x in r] for r in ws.iter_rows(max_row=3)] == [
        [x.value for x in r] for r in original.iter_rows(max_row=3)
    ]
    assert ws.column_dimensions["E"].width == 17
    assert list(sheet_rows(receipt.output)) == [a.product_code, c.product_code], "只留通過的列，順序照原表"
    assert ws.max_row == 5
    assert [ws.row_dimensions[r].height for r in (4, 5)] == [24, 26], "列高跟著列走"
    for r in (4, 5):
        assert header_cell(ws, "發行日", r).number_format == DATE_FORMAT, "回填的日期沿用表上既有日期格式"
        assert header_cell(ws, "比價日_2", r).value == "-"


def test_error_list_has_one_row_per_failing_pdf_in_input_order(tmp_path):
    ok, bad, missing = Spec(), Spec(product_code="029199990002"), Spec(product_code="029199990003")
    unsupported = Spec(product_code="999199990001")
    unreadable = tmp_path / "029199990009_TS.pdf"
    unreadable.write_bytes(b"not a pdf")
    nameless = tmp_path / "說明書.pdf"
    nameless.write_bytes(b"not a pdf")
    pdfs = [pdf_for(tmp_path, s) for s in (bad, ok, missing, unsupported)] + [unreadable, nameless]
    rows = [reference_row(ok), reference_row(bad, **{"K(%)": 71, "KO(%)": 101})]
    outcome, receipt, _ = batch(tmp_path, pdfs, rows)

    errors = error_rows(receipt.output)
    assert [(e["TDCC Code"], e["PDF 檔名"]) for e in errors] == [
        (bad.product_code, pdfs[0].name),
        (missing.product_code, pdfs[2].name),
        (unsupported.product_code, pdfs[3].name),
        ("029199990009", unreadable.name),
        (None, nameless.name),
        *((s.product_code, f"{s.product_code}_IIS.pdf") for s in (bad, missing, unsupported)),
    ], "每份沒通過的 PDF 一列，依輸入順序；TDCC Code 取封面商品代號、檔名前 12 碼，都取不到時留白"
    lines = errors[0]["錯訊"].split("\n")
    assert len(lines) == 2 and "K(%)對不起來：參考條件表 71.00／說明書 70.00" in lines, "多條錯訊在同一格，以換行分隔"
    assert all(e["錯訊"] for e in errors)
    assert openpyxl.load_workbook(receipt.output)["錯誤清單"]["C2"].alignment.wrap_text
    assert list(sheet_rows(receipt.output)) == [ok.product_code]


def edit_sheet(sheet: Path, change) -> None:
    wb = openpyxl.load_workbook(sheet)
    change(wb)
    wb.save(sheet)


def test_filter_range_shrinks_with_the_kept_rows_and_the_file_opens_on_the_filled_sheet(tmp_path):
    a, b = Spec(), Spec(product_code="029199990002")
    sheet = build_reference_sheet(tmp_path / "FCN參考條件.xlsx", [reference_row(a), reference_row(b, **{"K(%)": 71})])

    def change(wb):
        wb["樣本清單"].auto_filter.ref = "A3:F5"
        wb.active = wb["詢價表格"]

    edit_sheet(sheet, change)
    outcome, receipt, _ = batch(tmp_path, [pdf_for(tmp_path, s) for s in (a, b)], [], sheet=sheet)
    out = openpyxl.load_workbook(receipt.output)
    assert out["回填後"].auto_filter.ref == "A3:F4"
    assert out.active.title == "回填後"


@pytest.mark.parametrize(
    "feature,change",
    [
        ("合併儲存格", lambda ws: ws.merge_cells("C4:D4")),  # 庫存狀態、當日比價（不回填）
        ("合併儲存格", lambda ws: ws.merge_cells("A4:B4")),  # 蓋到回填欄 IIS：回填前就要擋下
        ("合併儲存格", lambda ws: ws.merge_cells("G4:H4")),  # 蓋到回填欄 ISIN Code
        ("格式化條件", lambda ws: ws.conditional_formatting.add("H4:H9", CellIsRule(operator="equal", formula=["0"]))),
        ("資料驗證", lambda ws: ws.add_data_validation(DataValidation(type="list", formula1='"Y,N"', sqref="T4:T9"))),
        ("公式", lambda ws: ws.cell(1, 1, "=COUNTA(E4:E9)")),
    ],
)
def test_layout_that_cannot_be_trimmed_safely_writes_no_result_file(tmp_path, feature, change):
    spec = Spec()
    sheet = build_reference_sheet(tmp_path / "FCN參考條件.xlsx", [reference_row(spec)])
    edit_sheet(sheet, lambda wb: change(wb["樣本清單"]))
    outcome, receipt, _ = batch(tmp_path, [pdf_for(tmp_path, spec)], [], sheet=sheet)
    assert receipt.output is None and receipt.status == ERROR and not receipt.filled(outcome.items[0])
    e = receipt.errors[0]
    assert (e.rule_id, e.reason_code) == ("output.result_file", "result_layout_unsupported")
    assert feature in e.message
    assert list((tmp_path / "reports").glob("*_核對結果_*.xlsx")) == []
    assert receipt.record is not None and receipt.record.exists()


def test_existing_result_file_is_never_overwritten(tmp_path):
    spec = Spec()
    taken = tmp_path / "reports" / "FCN參考條件_核對結果_20300203-040506.xlsx"
    taken.parent.mkdir()
    taken.write_bytes(b"keep")
    outcome, receipt, _ = batch(tmp_path, [pdf_for(tmp_path, spec)], [reference_row(spec)])
    assert taken.read_bytes() == b"keep"
    assert receipt.output is None and receipt.status == ERROR
    assert (receipt.errors[0].rule_id, receipt.errors[0].reason_code) == ("output.result_file", "output_exists")
    assert not receipt.filled(outcome.items[0])


def test_config_problem_is_reported_when_loading_the_config_and_nothing_is_written(tmp_path):
    with pytest.raises(IngestionError) as raised:
        load_config(reference_format=tmp_path / "missing.toml")
    assert raised.value.reason_code == "config_not_found" and "missing.toml" in str(raised.value)
    assert not (tmp_path / "reports").exists() and not (tmp_path / "runtime").exists()


def test_broken_reference_sheet_is_a_batch_error_and_writes_no_result_file(tmp_path):
    spec = Spec()
    broken = tmp_path / "FCN參考條件.xlsx"
    broken.write_bytes(b"not an Excel")
    outcome, receipt, _ = batch(tmp_path, [pdf_for(tmp_path, spec)], [], sheet=broken)
    assert receipt.status == ERROR and outcome.items == [] and receipt.output is None
    assert outcome.errors[0].rule_id == "input.batch"
    assert not (tmp_path / "reports").exists() and not (tmp_path / "runtime").exists()


def test_unreadable_pdf_does_not_stop_the_batch(tmp_path):
    spec = Spec()
    broken = tmp_path / "029199990009_TS.pdf"
    broken.write_bytes(b"not a pdf")
    outcome, receipt, _ = batch(tmp_path, [broken, pdf_for(tmp_path, spec)], [reference_row(spec)])
    assert outcome.items[0].report.status == ERROR and outcome.items[0].status_label == "執行錯誤"
    assert outcome.items[1].report.status == PASS and receipt.filled(outcome.items[1])
    assert receipt.status == ERROR
    assert [e["PDF 檔名"] for e in error_rows(receipt.output)] == [broken.name]


def test_unexpected_error_in_one_pdf_is_reported_and_the_batch_continues(tmp_path):
    def broken_read(lines):
        raise IndexError("synthetic parser failure")

    broken = barc_adapter(read=broken_read)
    spec = Spec()
    outcome, receipt, _ = batch(
        tmp_path, [pdf_for(tmp_path, spec)], [reference_row(spec)], config=CONFIG.with_registry((broken,))
    )
    item = outcome.items[0]
    assert item.report.status == ERROR
    assert only(item, "batch.unexpected").message == "IndexError: synthetic parser failure"
    assert receipt.output is not None and receipt.record is not None


def test_check_writes_nothing_until_saved(tmp_path):
    spec = Spec()
    pdf = pdf_for(tmp_path, spec)
    sheet = build_reference_sheet(tmp_path / "FCN參考條件.xlsx", [reference_row(spec)])
    before = set(tmp_path.rglob("*"))
    outcome = check_all(sheet, [pdf])
    assert outcome.status == PASS
    assert set(tmp_path.rglob("*")) == before

    receipt = save_batch(outcome, tmp_path / "reports", root=tmp_path, now=NOW)
    assert receipt.filled(outcome.items[0]) and receipt.output.is_file() and receipt.record.is_file()


@pytest.mark.parametrize("changed", ["sheet", "pdf", "standard", "prefixes"])
def test_any_source_changed_after_check_writes_nothing(tmp_path, changed):
    spec = Spec()
    pdf = pdf_for(tmp_path, spec)
    sheet = build_reference_sheet(tmp_path / "FCN參考條件.xlsx", [reference_row(spec)])
    standard, prefixes = tmp_path / "standard.toml", tmp_path / "prefixes.toml"
    standard.write_bytes(REVIEW_STANDARD.read_bytes())
    prefixes.write_bytes(ISSUER_PREFIXES.read_bytes())
    outcome = check_all(sheet, [pdf], load_config(review_standard=standard, issuer_prefixes=prefixes))
    if changed == "sheet":
        build_reference_sheet(sheet, [reference_row(spec, **{"K(%)": 71})])
    elif changed == "pdf":
        build_pdf(pdf, Spec(tenor=7))
    else:
        path = standard if changed == "standard" else prefixes
        path.write_text(path.read_text(encoding="utf-8") + "\n# 核對後被改過\n", encoding="utf-8")
    receipt = save_batch(outcome, tmp_path / "reports", root=tmp_path, now=NOW)
    assert receipt.output is None and receipt.status == ERROR
    assert receipt.errors[0].reason_code == "source_changed"
    assert not receipt.filled(outcome.items[0])
    assert not (tmp_path / "reports").exists(), "不寫核對結果檔"
    assert not (tmp_path / "runtime").exists(), "也不寫核對紀錄"


def test_pdf_missing_since_the_check_does_not_block_saving_but_one_that_appears_does(tmp_path):
    spec = Spec()
    missing = tmp_path / "029199990009_TS.pdf"
    outcome = checked(tmp_path, [pdf_for(tmp_path, spec), missing], [reference_row(spec)])
    receipt = save_batch(outcome, tmp_path / "reports", root=tmp_path, now=NOW)
    assert receipt.output is not None, "核對時就不存在的說明書記成執行錯誤，其他照常儲存"

    pdf_for(tmp_path, Spec(product_code="029199990009"))  # 核對後才出現
    receipt = save_batch(outcome, tmp_path / "reports", root=tmp_path, now=NOW + dt.timedelta(seconds=1))
    assert receipt.output is None and receipt.errors[0].reason_code == "source_changed"


def test_record_hashes_come_from_the_snapshot_taken_before_reading(tmp_path):
    spec = Spec()
    pdf = pdf_for(tmp_path, spec)
    outcome, receipt, sheet = batch(tmp_path, [pdf], [reference_row(spec)])
    record = json.loads(receipt.record.read_text(encoding="utf-8"))
    snapshot = outcome.snapshot
    assert record["metadata"]["inputs"]["reference_sheet"]["sha256"] == snapshot.sha256(sheet)
    assert record["items"][0]["metadata"]["inputs"]["term_sheet"]["sha256"] == snapshot.sha256(pdf)
    assert record["metadata"]["review_standard"]["sha256"] == snapshot.sha256(REVIEW_STANDARD)


# ---------------------------------------------------------------- 預覽沿用到核對


def test_check_reuses_the_preview_when_its_snapshot_is_still_valid(tmp_path):
    ok, missing = Spec(), Spec(product_code="029199990002")
    pdfs = [pdf_for(tmp_path, ok), pdf_for(tmp_path, missing)]
    sheet = build_reference_sheet(tmp_path / "FCN參考條件.xlsx", [reference_row(ok)])
    preview = preview_batch(CONFIG, sheet, with_iis(pdfs))

    outcome = check_batch(preview)
    assert outcome.snapshot is preview.snapshot
    for row, item in zip(preview.rows, outcome.items, strict=True):
        assert (item.term_sheet, item.issuer, item.product_code, item.reference_row) == (
            row.term_sheet,
            row.issuer,
            row.product_code,
            row.reference_row,
        ), "核對結果與預覽的辨識一致"
    assert [i.report.status for i in outcome.items] == [PASS, REVIEW, PASS, REVIEW]


def test_term_sheet_replaced_after_preview_requires_a_new_preview(tmp_path):
    spec = Spec()
    pdf = pdf_for(tmp_path, spec)
    sheet = build_reference_sheet(tmp_path / "FCN參考條件.xlsx", [reference_row(spec)])
    preview = preview_batch(CONFIG, sheet, with_iis([pdf]))
    build_pdf(pdf, Spec(tenor=7))  # 同名、內容不同

    with pytest.raises(IngestionError, match="重新載入預覽") as raised:
        check_batch(preview)
    assert raised.value.reason_code == "source_changed"


def test_config_file_changed_after_loading_the_config_requires_a_new_preview(tmp_path):
    spec = Spec()
    pdf = pdf_for(tmp_path, spec)
    sheet = build_reference_sheet(tmp_path / "FCN參考條件.xlsx", [reference_row(spec)])
    standard = tmp_path / "standard.toml"
    standard.write_bytes(REVIEW_STANDARD.read_bytes())
    config = load_config(review_standard=standard)
    standard.write_text(standard.read_text(encoding="utf-8") + "\n# 載入後被改過\n", encoding="utf-8")

    preview = preview_batch(config, sheet, with_iis([pdf]))  # 快照的設定檔 hash 取自核對設定（載入前取）
    with pytest.raises(IngestionError, match="重新載入預覽"):
        check_batch(preview)


def test_cli_run_with_a_config_file_changed_after_loading_asks_to_check_again(tmp_path):
    spec = Spec()
    standard = tmp_path / "standard.toml"
    standard.write_bytes(REVIEW_STANDARD.read_bytes())
    config = load_config(review_standard=standard)
    standard.write_text(standard.read_text(encoding="utf-8") + "\n# 載入後被改過\n", encoding="utf-8")

    outcome, receipt, _ = batch(tmp_path, [pdf_for(tmp_path, spec)], [reference_row(spec)], config=config)
    assert receipt.status == ERROR and outcome.items == [] and receipt.output is None
    (error,) = outcome.errors
    assert error.reason_code == "source_changed"
    assert "重新核對" in error.message and "預覽" not in error.message, "CLI 沒有預覽可重新載入"


def test_each_term_sheet_is_read_once_across_preview_and_check(tmp_path):
    from fcn_checker.parsers import barc as parser

    reads: list[str] = []

    def counted_read(lines):
        reads.append("read")
        return parser.read(lines)

    config = CONFIG.with_registry((barc_adapter(read=counted_read),))
    specs = [Spec(), Spec(product_code="029199990002")]
    pdfs = [pdf_for(tmp_path, s) for s in specs]
    sheet = build_reference_sheet(tmp_path / "FCN參考條件.xlsx", [reference_row(s) for s in specs])

    preview = preview_batch(config, sheet, with_iis(pdfs))
    assert len(reads) == 2
    outcome = check_batch(preview)
    assert [i.report.status for i in outcome.items] == [PASS, PASS, PASS, PASS]
    assert len(reads) == 2, "核對沿用預覽的讀出，不再讀一次"


def test_record_config_info_comes_from_the_check_config(tmp_path):
    spec = Spec()
    outcome, receipt, _ = batch(tmp_path, [pdf_for(tmp_path, spec)], [reference_row(spec)])
    metadata = json.loads(receipt.record.read_text(encoding="utf-8"))["metadata"]
    for name, info in CONFIG.record().items():
        assert metadata[name] == info
    assert metadata["review_standard"]["version"] == CONFIG.review_standard.version


# ---------------------------------------------------------------- 人工放行


def checked(tmp_path: Path, pdfs: list[Path], rows: list[dict]):
    """只核對不儲存（PANEL 的用法）：之後可放行再 save_batch。"""
    tmp_path.mkdir(exist_ok=True)
    sheet = build_reference_sheet(tmp_path / "FCN參考條件.xlsx", rows)
    return check_all(sheet, pdfs)


def release_both(outcome, item) -> None:
    """放行說明書與同商品投資人須知（兩份都通過或放行才回填，ADR 0007）。"""
    outcome.release(item)
    outcome.release(item.partner)


def saved(outcome, tmp_path: Path, now: dt.datetime = NOW):
    """儲存一次：回傳收據與這次的核對紀錄。"""
    receipt = save_batch(outcome, tmp_path / "reports", root=tmp_path, now=now)
    record = tmp_path / "runtime" / "核對紀錄" / f"{now:%Y%m%d-%H%M%S}.json"
    return receipt, json.loads(record.read_text(encoding="utf-8"))


@pytest.mark.parametrize(
    ("row", "original"),
    [({"K(%)": 71}, "不一致"), ({"K(%)": None}, "需人工覆核")],
    ids=["mismatch", "review"],
)
def test_released_term_sheet_is_filled_like_a_pass(tmp_path, row, original):
    spec = Spec()
    outcome = checked(tmp_path, [pdf_for(tmp_path, spec)], [reference_row(spec, **row)])
    item = outcome.items[0]
    assert outcome.status != PASS and item.release_problem == "" and item.partner.release_problem == ""

    outcome.release(item)
    assert item.released and item.status_label == f"人工放行（原：{original}）"
    assert not item.fills_sheet, "同商品投資人須知也對不起來，還沒放行就不回填"
    assert item.not_filled_reason == "同商品的投資人須知尚未通過或人工放行，不回填。"
    outcome.release(item.partner)
    assert item.fills_sheet and outcome.status == PASS, "整批狀態不必另外重算"

    receipt, record = saved(outcome, tmp_path)
    result = receipt.output
    assert list(sheet_rows(result)) == [spec.product_code], "回填後只有這一列"
    row = row_of(result, spec.product_code)
    assert (row["TS"], row["IIS"], row["ISIN Code"]) == ("V", "V", SYNTH_ISIN), "人工放行的也打勾"
    assert error_rows(result) == [], "人工放行的說明書不列入錯誤清單"
    assert receipt.filled(item)
    entry, iis_entry = record["items"]
    assert entry["manual_release"] is True and entry["filled"] is True
    assert entry["status"] != "PASS", "核對紀錄保留原判定"
    assert (iis_entry["document"], iis_entry["partner"], iis_entry["manual_release"]) == (
        "投資人須知",
        item.term_sheet.name,
        True,
    )
    assert record["status"] == "PASS"


def test_released_flag_cannot_be_set_from_outside(tmp_path):
    import dataclasses

    from fcn_checker.batch import BatchItem

    spec = Spec()
    outcome = checked(tmp_path, [pdf_for(tmp_path, spec)], [reference_row(spec, **{"K(%)": 71})])
    item = outcome.items[0]
    with pytest.raises(AttributeError):
        item.released = True
    with pytest.raises(TypeError):
        BatchItem(item.term_sheet, item.report, _released=True)
    with pytest.raises((TypeError, ValueError)):  # init=False 欄位：3.13 為 TypeError、3.11 為 ValueError
        dataclasses.replace(item, _released=True)


def test_only_items_of_this_outcome_can_be_released(tmp_path):
    spec = Spec()
    rows = [reference_row(spec, **{"K(%)": 71})]
    first = checked(tmp_path / "a", [pdf_for(tmp_path, spec)], rows)
    second = checked(tmp_path / "b", [pdf_for(tmp_path, spec)], rows)
    with pytest.raises(IngestionError, match="重新核對"):
        second.release(first.items[0])
    assert not first.items[0].released and second.status == MISMATCH


def test_cancelled_release_returns_to_the_original_result(tmp_path):
    spec = Spec()
    outcome = checked(tmp_path, [pdf_for(tmp_path, spec)], [reference_row(spec, **{"K(%)": 71})])
    item = outcome.items[0]
    release_both(outcome, item)
    outcome.cancel_release(item)
    assert not item.released and item.status_label == "不一致" and outcome.status == MISMATCH

    receipt, record = saved(outcome, tmp_path)
    result = receipt.output
    assert spec.product_code not in sheet_rows(result)
    assert [e["PDF 檔名"] for e in error_rows(result)] == [item.term_sheet.name]
    assert record["items"][0]["manual_release"] is False and not receipt.filled(item)


def test_changing_a_release_after_saving_needs_a_new_save_and_keeps_the_old_receipt(tmp_path):
    spec = Spec()
    outcome = checked(tmp_path, [pdf_for(tmp_path, spec)], [reference_row(spec, **{"K(%)": 71})])
    item = outcome.items[0]
    release_both(outcome, item)
    first, _ = saved(outcome, tmp_path)
    assert first.filled(item)

    outcome.cancel_release(item)
    assert first.filled(item), "收據記的是那一次儲存，不跟著改"

    second, _ = saved(outcome, tmp_path, NOW + dt.timedelta(seconds=1))
    assert not second.filled(item)
    assert spec.product_code in sheet_rows(first.output)
    assert spec.product_code not in sheet_rows(second.output) and len(error_rows(second.output)) == 1


def test_saving_twice_gives_two_files_and_a_fresh_receipt_without_the_earlier_errors(tmp_path):
    spec = Spec()
    outcome = checked(tmp_path, [pdf_for(tmp_path, spec)], [reference_row(spec)])
    taken = tmp_path / "runtime" / "核對紀錄" / "20300203-040506.json"
    taken.parent.mkdir(parents=True)
    taken.write_text("keep", encoding="utf-8")
    before = repr(outcome)

    first = save_batch(outcome, tmp_path / "reports", root=tmp_path, now=NOW)
    assert first.output.is_file() and first.record is None
    assert [(e.rule_id, e.reason_code) for e in first.errors] == [("output.record", "output_exists")]

    later = NOW + dt.timedelta(seconds=1)
    second = save_batch(outcome, tmp_path / "reports", root=tmp_path, now=later)
    assert second.errors == () and second.status == PASS and second.filled(outcome.items[0])
    assert second.output.is_file() and second.record.is_file() and second.output != first.output
    assert len(list((tmp_path / "reports").glob("*_核對結果_*.xlsx"))) == 2, "兩組檔案"
    assert first.errors and first.status == ERROR, "第一張收據不變"
    assert repr(outcome) == before and outcome.errors == [], "核對結果本身不被儲存改寫"


def test_release_problem_messages_match_the_error_list(tmp_path):
    spec = Spec()
    outcome = checked(tmp_path, [pdf_for(tmp_path, spec)], [reference_row(spec, **{"K(%)": 71, "KO(%)": 99})])
    item, iis = outcome.items
    receipt, _ = saved(outcome, tmp_path)
    error, iis_error = error_rows(receipt.output)
    assert len(item.problem_messages) == 2
    assert error["錯訊"] == "\n".join(item.problem_messages), "放行確認視窗與錯誤清單列出同樣的錯訊"
    assert iis_error["錯訊"] == "\n".join(iis.problem_messages)
    assert "K(%)對不起來：參考條件表 71.00／投資人須知 70.00" in iis.problem_messages


def _unreadable(tmp_path: Path, spec: Spec) -> Path:
    pdf = tmp_path / f"{spec.product_code}_TS.pdf"
    pdf.write_bytes(b"not a pdf")
    return pdf


@pytest.mark.parametrize(
    ("case", "reason"),
    [
        ("pass", "已經通過"),
        ("backfill_unknown", "回填值無法確定"),
        ("backfill_conflict", "請先修正參考條件表"),
        ("shared_row", "同一批有多份文件對到同一列"),
        ("missing_iis", "這批缺同商品的投資人須知"),
        ("missing_row", "沒有對到參考條件表的列"),
        ("unsupported", "未支援上手"),
        ("error", "執行錯誤"),
    ],
)
def test_release_is_refused_when_backfill_is_not_trustworthy(tmp_path, case, reason):
    spec = Spec()
    rows = [reference_row(spec)]
    pdfs = [pdf_for(tmp_path, spec)]
    if case == "backfill_unknown":
        spec = Spec(omit=frozenset({"issue_date"}))
        pdfs = [pdf_for(tmp_path, spec)]
    elif case == "backfill_conflict":
        rows = [reference_row(spec, **{"ISIN Code": "XS9999999999"})]
    elif case == "shared_row":
        (tmp_path / "copy").mkdir()
        pdfs = [pdfs[0], pdf_for(tmp_path / "copy", spec)]
    elif case == "missing_iis":
        (tmp_path / "solo").mkdir()
        pdfs = [build_pdf(tmp_path / "solo" / f"{spec.product_code}_TS.pdf", spec, iis=False)]
    elif case == "missing_row":
        rows = [reference_row(Spec(product_code="029199990009"))]
    elif case == "unsupported":
        pdfs = [pdf_for(tmp_path, Spec(product_code="999199990001"))]
    elif case == "error":
        pdfs = [_unreadable(tmp_path, spec)]
    outcome = checked(tmp_path, pdfs, rows)
    item = outcome.items[0]

    assert reason in item.release_problem
    with pytest.raises(IngestionError, match=reason):
        outcome.release(item)
    assert not item.released


# ---------------------------------------------------------------- 核對紀錄


def test_save_writes_one_audit_record_named_like_the_result_file(tmp_path):
    ok, bad = Spec(), Spec(product_code="029199990002")
    pdfs = [pdf_for(tmp_path, ok), pdf_for(tmp_path, bad)]
    outcome, receipt, sheet = batch(tmp_path, pdfs, [reference_row(ok), reference_row(bad, **{"K(%)": 71})])

    assert receipt.record == tmp_path / "runtime" / "核對紀錄" / "20300203-040506.json"
    assert receipt.output.name == "FCN參考條件_核對結果_20300203-040506.xlsx"
    assert [p.name for p in (tmp_path / "reports").iterdir()] == [receipt.output.name], "輸出資料夾只有核對結果檔"
    data = json.loads(receipt.record.read_text(encoding="utf-8"))
    assert (data["status"], data["result_file"]) == ("MISMATCH", str(receipt.output))
    assert dt.datetime.fromisoformat(data["saved_at"]).replace(tzinfo=None) == NOW, "執行時間（含時區）"
    assert dt.datetime.fromisoformat(data["saved_at"]).tzinfo is not None
    meta = data["metadata"]
    assert meta["program_version"] == __version__ and "program_commit" in meta
    for key, path in (
        ("review_standard", REVIEW_STANDARD),
        ("reference_format", REFERENCE_FORMAT),
        ("issuer_prefixes", ISSUER_PREFIXES),
    ):
        assert (meta[key]["path"], meta[key]["sha256"]) == (str(path.resolve()), sha256_of(path)), key
    assert meta["inputs"]["reference_sheet"]["sha256"] == sha256_of(sheet)

    first, second, *iis = data["items"]
    assert [i["pdf"] for i in data["items"]] == [p.name for p in with_iis(pdfs)]
    assert [(i["document"], i["partner"]) for i in data["items"]] == [
        ("說明書", "029199990001_IIS.pdf"),
        ("說明書", "029199990002_IIS.pdf"),
        ("投資人須知", "029199990001_TS.pdf"),
        ("投資人須知", "029199990002_TS.pdf"),
    ]
    assert first["metadata"]["inputs"]["term_sheet"]["sha256"] == sha256_of(pdfs[0])
    assert (first["status"], first["filled"], first["reference_row"]) == ("PASS", True, 4)
    cells = {d["column"]: d for d in first["backfill"]}
    assert {"cell": "H4", "action": "fill"}.items() <= cells["ISIN Code"].items()
    for column, cell in (("TS", "A4"), ("IIS", "B4")):  # TS、IIS 的回填決策也記下（Issue #127）
        assert {"cell": cell, "sheet_value": None, "expected": "V", "action": "fill"}.items() <= cells[column].items()
    assert (second["status"], second["filled"]) == ("MISMATCH", False)
    [strike] = [r for r in second["results"] if r["rule_id"] == "field.strike_pct"]
    assert strike["status"] == "MISMATCH" and strike["document_evidence"][0]["page"] >= 1
    assert strike["document_evidence"][0]["text"] and strike["order_source"][0].startswith("樣本清單!")
    assert strike["item"] == {"name": "K(%)", "source": "reference", "columns": ["K(%)"]}, "核對紀錄含項目名稱與出處"
    columns = {r["rule_id"]: r["column"] for r in first["results"]}
    assert (columns["field.strike_pct"], columns["field.underlyings"], columns["backfill.compare_dates"]) == (
        "K(%)",
        "標的",
        "比價日",
    ), "既有的 column 欄位值不變：多欄合起來核對時寫項目名稱"
    assert columns["derive.monthly_coupon"] == "Coupon p.a. (%)、天期(月)"


def test_existing_record_is_never_overwritten_and_the_result_file_is_still_written(tmp_path):
    spec = Spec()
    taken = tmp_path / "runtime" / "核對紀錄" / "20300203-040506.json"
    taken.parent.mkdir(parents=True)
    taken.write_text("keep", encoding="utf-8")
    outcome, receipt, _ = batch(tmp_path, [pdf_for(tmp_path, spec)], [reference_row(spec)])
    assert taken.read_text(encoding="utf-8") == "keep"
    assert receipt.output.is_file() and receipt.filled(outcome.items[0])
    assert receipt.record is None and receipt.status == ERROR
    e = receipt.errors[-1]
    assert (e.rule_id, e.reason_code) == ("output.record", "output_exists") and "核對紀錄" in e.message


def test_unwritable_record_folder_is_reported_without_losing_the_result_file(tmp_path):
    spec = Spec()
    (tmp_path / "runtime").write_text("不是資料夾", encoding="utf-8")
    outcome, receipt, _ = batch(tmp_path, [pdf_for(tmp_path, spec)], [reference_row(spec)])
    assert receipt.output.is_file() and receipt.filled(outcome.items[0])
    assert receipt.record is None and receipt.status == ERROR
    assert (receipt.errors[-1].rule_id, receipt.errors[-1].reason_code) == ("output.record", "output_unwritable")
