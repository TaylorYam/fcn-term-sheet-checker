"""合成參考條件表（不分上手）：上手中立的商品規格 `ProductSpec`、與之一致的參考條件表列 `reference_row`（tests/ 只此一份），
以及仿 FCN參考條件 Excel `樣本清單` 的 `build_reference_sheet`（Issue #145）。

各上手的合成器（tests/<上手>_synth.py）以 `ProductSpec` 的子類別加上該上手的版面結構旋鈕，只負責把規格畫成說明書與
投資人須知 PDF（製造錯誤用 `edits`，見 tests/pdf_writer.py 的 `Edit`）；參考條件表列不由上手產生。欄名對照與「發行機構」欄寫法取自參考條件表格式設定
（config/reference_sheet.toml）。所有數值皆為虛構。
"""

from __future__ import annotations

import datetime as dt
import tomllib
from dataclasses import dataclass, replace
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Any, Self

import openpyxl

from harness import CURRENCY_ISO, REFERENCE_FORMAT

_FORMAT = tomllib.loads(REFERENCE_FORMAT.read_text(encoding="utf-8"))
# Excel 欄名 → 標準欄位（[columns] 內的 ignored 清單不是欄名）
REFERENCE_COLUMNS: dict[str, str] = {h: f for h, f in _FORMAT["columns"].items() if isinstance(f, str)}
_COLUMN_OF = {f: h for h, f in REFERENCE_COLUMNS.items()}  # 標準欄位 → Excel 欄名

REFERENCE_HEADERS = [
    "TS",
    "IIS",
    "庫存狀態",
    "當日比價",
    "Product",
    "發行機構",
    "TDCC Code",
    "ISIN Code",
    "私銀註記",
    "單位面額",
    "承作幣別",
    "交易日",
    "發行日",
    *[f"比價日_{i}" for i in range(1, 13)],
    "最終比價日",
    "到期日",
    "期初定價",
    "KO(%)",
    "KO(Freq)",
    "KO(memo)",
    "K(%)",
    "KI(%)",
    "KI(Freq)",
    "UF",
    "Coupon p.a. (%)",
    "天期(月)",
    "Non-Call(月)",
    *[f"UL_{i}{suffix}" for i in range(1, 6) for suffix in ("", "_進場價", "_執行價", "_下限價", "_KO價", "_Memo")],
]
DATE_FORMAT = "mm-dd-yy"
Q4 = Decimal("0.0001")
DEFAULT_DENOMINATION = {"USD": 10000, "JPY": 1000000, "CNH": 100000}  # 幣別預設面額


def price(initial: Decimal, pct: Decimal) -> Decimal:
    """期初價 × 百分比，四位小數（參考條件表各標的價格欄與說明書價格表同一算法）。"""
    return (initial * pct / 100).quantize(Q4, ROUND_HALF_UP)


@dataclass(frozen=True)
class UL:
    """一檔連結標的：名稱、交易所、彭博代碼、期初價。"""

    name: str
    exchange: str
    ticker: str
    initial: Decimal


DEFAULT_ULS = (
    UL("甲乙丙科技股份有限公司ADR", "紐約證券交易所", "ZZA UN", Decimal("123.4500")),
    UL("Zeta Quantum Holdings Inc", "那斯達克證券交易所", "ZQH UW", Decimal("87.2000")),
    UL("丁戊電子公司", "那斯達克證券交易所", "DWE UW", Decimal("1234.5600")),
)


@dataclass
class ProductSpec:
    """一檔 FCN 商品的規格（上手中立）：參考條件表一列的內容，也是各上手合成說明書與投資人須知的依據。

    `issuer` 與 `product_code` 沒有預設值：由各上手合成器的子類別補上，並加該上手的版面結構旋鈕；子類別由其他欄位推得的
    規格（例：BARC 的 Non-Call、MS／HSBC 的標的與日期）改宣告 `init=False`，建構或 `with_` 時給了會直接報錯。
    預設值下說明書、投資人須知與參考條件表列完全一致。
    """

    issuer: str  # 上手代號（參考條件表「發行機構」欄的寫法由它查格式設定；檔名前三碼也依它）
    product_code: str
    currency_zh: str = "美元"
    tenor: int = 6  # 天期（月）
    memory: bool = True  # 記憶式提前出場
    ko_obs: str = "D"  # KO 觀察方式：D 期間每日／P 期末定日
    ki: str = "none"  # KI 型態：none／AM／D，BARC 另有 M、MS 另有 P
    strike: Decimal = Decimal("70.00")  # 執行價 %
    ko: Decimal = Decimal("100.00")  # 自動提前出場價 %
    ki_pct: Decimal = Decimal("60.00")  # 下限價 %（ki 為 none 時不用）
    annual: Decimal = Decimal("12.00")  # 年利率 %
    first_callable: int = 1  # Non-Call(月)：第一個可提前出場期
    trade_date: dt.date = dt.date(2030, 1, 7)
    issue_date: dt.date = dt.date(2030, 1, 14)
    final_date: dt.date = dt.date(2030, 7, 8)  # 最終比價日
    maturity_date: dt.date = dt.date(2030, 7, 11)
    denomination: int | None = None  # None → 幣別預設面額
    underlyings: tuple[UL, ...] = DEFAULT_ULS

    def with_(self, **kw: Any) -> Self:
        return replace(self, **kw)

    @property
    def ccy(self) -> str:
        return CURRENCY_ISO[self.currency_zh]

    @property
    def denom(self) -> int:
        return self.denomination or DEFAULT_DENOMINATION[self.ccy]


def issuer_value(issuer: str) -> str:
    """上手代號 → 該上手在參考條件表「發行機構」欄的寫法。"""
    return _FORMAT["issuer_values"][issuer]


def as_headers(fields: dict[str, Any]) -> dict[str, Any]:
    """標準欄位名 → Excel 欄名，給 `reference_row(spec, **as_headers({...}))` 以標準欄位名覆寫。"""
    return {_COLUMN_OF[k]: v for k, v in fields.items()}


def make_row(issuer: str, fields: dict[str, Any], **overrides: Any) -> dict[str, Any]:
    """一列參考條件表：fields 以標準欄位為鍵（依格式設定對到 Excel 欄名），overrides 以 Excel 欄名覆寫。

    沒給的欄位（例：回填欄位 TS、IIS、ISIN、發行日、比價日）留空；只用來辨識或存續管理的欄位填固定的虛構值。
    """
    row: dict[str, Any] = {
        "庫存狀態": None,
        "當日比價": None,
        "Product": "FCN",
        "發行機構": issuer_value(issuer),
        "私銀註記": "-",
        "UF": 1.5900000000000034,
        "期初定價": "收盤價",
        **{f"UL_{i}_Memo": "-" for i in range(1, 6)},
    }
    row.update({h: fields[f] for h, f in REFERENCE_COLUMNS.items() if f in fields})
    row.update(overrides)
    return row


def _xl_date(d: dt.date) -> dt.datetime:
    return dt.datetime.combine(d, dt.time())


def reference_row(spec: ProductSpec, **overrides: Any) -> dict[str, Any]:
    """與合成說明書一致的參考條件表列（不分上手）；回填欄位（TS、IIS、ISIN、發行日、比價日）預設空白。
    overrides 以 Excel 欄名覆寫。"""
    fields: dict[str, Any] = {
        "product_code": spec.product_code,
        "denomination": spec.denom,
        "currency": spec.ccy,
        "trade_date": _xl_date(spec.trade_date),
        "final_valuation_date": _xl_date(spec.final_date),
        "maturity_date": _xl_date(spec.maturity_date),
        "ko_pct": float(spec.ko),
        "ko_observation": spec.ko_obs,
        "ko_memory": "Y" if spec.memory else "N",
        "strike_pct": float(spec.strike),
        "ki_pct": float(spec.ki_pct) if spec.ki != "none" else "-",
        "ki_type": spec.ki if spec.ki != "none" else "-",
        "coupon_pa_pct": float(spec.annual),
        "tenor_months": spec.tenor,
        "first_callable_period": spec.first_callable,
    }
    for i in range(1, 6):
        u = spec.underlyings[i - 1] if i <= len(spec.underlyings) else None
        fields[f"underlying_{i}"] = u.ticker if u else "-"
        fields[f"underlying_{i}_initial_price"] = float(u.initial) if u else "-"
        fields[f"underlying_{i}_strike_price"] = float(price(u.initial, spec.strike)) if u else "-"
        fields[f"underlying_{i}_ki_price"] = float(price(u.initial, spec.ki_pct)) if u and spec.ki != "none" else "-"
        fields[f"underlying_{i}_ko_price"] = float(price(u.initial, spec.ko)) if u else "-"
    return make_row(spec.issuer, fields, **overrides)


def build_reference_sheet(
    path: Path, rows: list[dict[str, Any]], headers: list[str] | None = None, extra_sheets: tuple[str, ...] = ()
) -> Path:
    """仿 FCN參考條件 的 `樣本清單`：第 3 列表頭、第 4 列起資料，最右側一個無表頭欄；另有一張其他工作表。"""
    headers = headers or REFERENCE_HEADERS
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "樣本清單"
    for k, h in enumerate(headers, start=1):
        ws.cell(row=3, column=k, value=h)
    for r, row in enumerate(rows, start=4):
        for k, h in enumerate(headers, start=1):
            c = ws.cell(row=r, column=k, value=row.get(h))
            if isinstance(c.value, dt.datetime):  # 空白格維持「通用格式」，回填時要沿用表上日期格式
                c.number_format = DATE_FORMAT
        ws.cell(row=r, column=len(headers) + 9, value=1 if row.get("庫存狀態") else None)
    other = wb.create_sheet("詢價表格")
    other["B3"] = "其他工作表（回填時不得改動）"
    for name in extra_sheets:
        wb.create_sheet(name)
    wb.save(path)
    return path
